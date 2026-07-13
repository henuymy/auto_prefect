"""Generate the metric target Excel template with full area hierarchy reference."""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

PROJECT_DIR = Path(__file__).resolve().parents[3]
REF_JSON = PROJECT_DIR / "docs" / "area_indicator_ref.json"
OUTPUT = PROJECT_DIR / "docs" / "指标目标值模板.xlsx"


def main() -> int:
    with open(REF_JSON, encoding="utf-8") as f:
        areas = json.load(f)

    wb = Workbook()
    ws = wb.active
    ws.title = "指标目标值"

    # ── shared styles ──
    header_font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    data_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    note_font = Font(name="微软雅黑", size=9, color="808080")
    note_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    label_font = Font(name="微软雅黑", size=9, color="4472C4")
    label_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")

    # ══════════════════════════════════════════════════════════════════
    # Sheet 1: import template
    # ══════════════════════════════════════════════════════════════════
    columns = [
        ("period_type", "周期类型", 18),
        ("area_code", "区域编码", 18),
        ("indicator_code", "指标编码", 18),
        ("target_value", "目标值(绝对值)", 18),
    ]

    for col_idx, (field, label, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=field)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

        label_cell = ws.cell(row=2, column=col_idx, value=label)
        label_cell.font = label_font
        label_cell.alignment = data_align
        label_cell.border = thin_border
        label_cell.fill = label_fill

    # example rows
    examples = [
        ("DAY_ACC", "AQ", "sgs_ajvwdz", 1000),
        ("DAY_ACC", "AQ701", "sgs_ajvwdz", 800),
        ("MONTH", "AQ", "sgs_ajvwdz", 30000),
        ("REALTIME", "AQ", "sgs_ajvwdz", 950),
    ]
    for row_idx, values in enumerate(examples, 3):
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = data_align
            cell.border = thin_border

    # period_type dropdown validation
    dv = DataValidation(
        type="list",
        formula1='"REALTIME,DAY_ACC,MONTH"',
        allow_blank=False,
        showErrorMessage=True,
        errorTitle="输入错误",
        error="period_type 只支持 REALTIME、DAY_ACC、MONTH",
    )
    dv.sqref = "A3:A10000"
    ws.add_data_validation(dv)

    # notes
    notes = [
        "填写说明：",
        "1. period_type：周期类型，只允许 REALTIME / DAY_ACC / MONTH",
        '2. area_code：区域编码，必须与系统 area 表中已启用的 area_code 一致（参见"参照"sheet）',
        '3. indicator_code：指标编码，必须与系统 indicator 表中已启用的 code 一致（参见"参照"sheet）',
        "4. target_value：目标值（绝对值，非百分数）",
        "   例如：目标完成 1000 户，则填 1000；系统会用 实际值/目标值 计算完成率",
        "5. 同一 (period_type, area_code, indicator_code) 不允许重复",
        "6. 导入后，表中已有但 Excel 中缺失的记录会被自动停用",
    ]
    note_row = 3 + len(examples) + 1
    for i, note in enumerate(notes):
        cell = ws.cell(row=note_row + i, column=1, value=note)
        cell.font = Font(
            name="微软雅黑",
            bold=(i == 0),
            size=9,
            color="808080",
        )
        cell.fill = note_fill
    ws.freeze_panes = "A3"

    # ══════════════════════════════════════════════════════════════════
    # Sheet 2: reference with full hierarchy
    # ══════════════════════════════════════════════════════════════════
    ws2 = wb.create_sheet("参照-区域与指标")
    ws2.sheet_properties.tabColor = "A9D08E"

    ref_header_font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
    ref_fill_area = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    ref_fill_ind = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
    data_font = Font(name="微软雅黑", size=10)
    level_fills = {
        "CITY": PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid"),
        "BRANCH": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
        "GRID": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
        "CHANNEL": PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
    }

    # Area columns: city, branch, grid, channel — each with code + name
    area_cols = [
        ("city_code", "市公司编码", 14),
        ("city_name", "市公司名称", 14),
        ("branch_code", "区公司编码", 14),
        ("branch_name", "区公司名称", 14),
        ("grid_code", "网格编码", 16),
        ("grid_name", "网格名称", 20),
        ("channel_code", "渠道编码", 22),
        ("channel_name", "渠道名称", 28),
    ]

    # Indicator columns (after a gap)
    ind_col_offset = len(area_cols) + 2
    ind_cols = [
        ("indicator_code", "指标编码", 20),
        ("indicator_name", "指标名称", 26),
    ]

    # Write headers row 1 (field names) and row 2 (Chinese labels)
    for col_idx, (key, label, width) in enumerate(area_cols, 1):
        cell = ws2.cell(row=1, column=col_idx, value=key)
        cell.font = ref_header_font
        cell.fill = ref_fill_area
        cell.alignment = header_align
        cell.border = thin_border
        ws2.column_dimensions[get_column_letter(col_idx)].width = width

        label_cell = ws2.cell(row=2, column=col_idx, value=label)
        label_cell.font = label_font
        label_cell.alignment = data_align
        label_cell.border = thin_border
        label_cell.fill = label_fill

    for col_idx, (key, label, width) in enumerate(ind_cols, ind_col_offset):
        cell = ws2.cell(row=1, column=col_idx, value=key)
        cell.font = ref_header_font
        cell.fill = ref_fill_ind
        cell.alignment = header_align
        cell.border = thin_border
        ws2.column_dimensions[get_column_letter(col_idx)].width = width

        label_cell = ws2.cell(row=2, column=col_idx, value=label)
        label_cell.font = Font(name="微软雅黑", size=9, color="548235")
        label_cell.alignment = data_align
        label_cell.border = thin_border
        label_cell.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

    # gap column
    ws2.column_dimensions[get_column_letter(len(area_cols) + 1)].width = 3

    # Fill data rows
    for i, a in enumerate(areas, 3):
        lt = a["level_type"]
        fill = level_fills.get(lt)

        city_code = city_name = None
        branch_code = branch_name = None
        grid_code = grid_name = None
        channel_code = channel_name = None

        if lt == "CITY":
            city_code, city_name = a["area_code"], a["area_name"]
        elif lt == "BRANCH":
            city_code, city_name = a.get("parent_code"), a.get("parent_name")
            branch_code, branch_name = a["area_code"], a["area_name"]
        elif lt == "GRID":
            city_code, city_name = a.get("grandparent_code"), a.get("grandparent_name")
            branch_code, branch_name = a.get("parent_code"), a.get("parent_name")
            grid_code, grid_name = a["area_code"], a["area_name"]
        elif lt == "CHANNEL":
            city_code, city_name = a.get("great_grandparent_code"), a.get("great_grandparent_name")
            branch_code, branch_name = a.get("grandparent_code"), a.get("grandparent_name")
            grid_code, grid_name = a.get("parent_code"), a.get("parent_name")
            channel_code, channel_name = a["area_code"], a["area_name"]

        values = [
            city_code, city_name,
            branch_code, branch_name,
            grid_code, grid_name,
            channel_code, channel_name,
        ]
        for col_idx, val in enumerate(values, 1):
            cell = ws2.cell(row=i, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if fill:
                cell.fill = fill

    # Indicator reference data (only 1 indicator currently)
    indicators = [
        {"code": "sgs_ajvwdz", "name": "审计无违规地址(V宽带)"},
    ]
    for j, ind in enumerate(indicators, 3):
        ws2.cell(row=j, column=ind_col_offset, value=ind["code"]).font = data_font
        ws2.cell(row=j, column=ind_col_offset).border = thin_border
        ws2.cell(row=j, column=ind_col_offset + 1, value=ind["name"]).font = data_font
        ws2.cell(row=j, column=ind_col_offset + 1).border = thin_border

    # Summary note
    summary_row = len(areas) + 3
    ws2.cell(row=summary_row, column=1, value=f"共 {len(areas)} 个区域").font = note_font
    ws2.cell(row=summary_row, column=ind_col_offset, value=f"共 {len(indicators)} 个指标").font = note_font

    # Auto-filter on the area columns so users can filter by city/branch/grid
    ws2.auto_filter.ref = f"A2:H{len(areas) + 2}"

    ws2.freeze_panes = "A3"

    wb.save(OUTPUT)
    print(f"模板已生成: {OUTPUT}")
    print(f"  区域: {len(areas)} (含完整层级)")
    print(f"  指标: {len(indicators)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
