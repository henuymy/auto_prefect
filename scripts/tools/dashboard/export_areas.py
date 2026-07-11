"""Export the city-ops area hierarchy from Zhengzhou root A to Excel.

The script reuses the existing city-ops request configuration and cookie dump.
It stops at CHANNEL and does not request the STAFF level.

Usage:
    python scripts/tools/dashboard/export_areas.py
    python scripts/tools/dashboard/export_areas.py --output runtime/exports/dashboard_areas.xlsx
"""

from __future__ import annotations

import argparse
import copy
import json
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_DIR = Path(__file__).resolve().parents[2]
if str(PROJECT_DIR) not in sys.path:
    sys.path.insert(0, str(PROJECT_DIR))

from services.method_service import download_reports


DEFAULT_CONFIG_PATH = PROJECT_DIR / "config" / "reports" / "家客和存量通报.json"
DEFAULT_OUTPUT_PATH = PROJECT_DIR / "docs" / "数据驾驶舱区域层级.xlsx"
DEFAULT_COOKIE_DUMP_PATH = PROJECT_DIR / "runtime" / "cookies" / "cookie_dump.json"

ROOT_CODE = "A"
ROOT_NAME = "郑州市"
DRILLDOWN_LEVELS = [
    ("区县", "BRANCH", 2),
    ("网格", "GRID", 3),
    ("渠道经理", "CHANNEL_MANAGER", 4),
    ("渠道", "CHANNEL", 5),
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="从地市平台 A/郑州市开始下探，并导出区域层级 Excel（不采集人员）"
    )
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG_PATH, help="报表配置文件")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_PATH, help="输出 Excel 文件")
    parser.add_argument(
        "--cookie-dump",
        type=Path,
        default=DEFAULT_COOKIE_DUMP_PATH,
        help="登录后导出的 Cookie 文件",
    )
    parser.add_argument("--max-workers", type=int, default=16, help="同层最大并行请求数")
    parser.add_argument("--max-requests", type=int, default=5000, help="最大请求数保护阈值")
    return parser.parse_args()


def load_source_report(config_path: Path) -> dict:
    with config_path.resolve().open("r", encoding="utf-8") as file:
        config = json.load(file)

    for item in config.get("downloads", []):
        if (
            item.get("enabled", True)
            and item.get("stage") == "city_ops"
            and item.get("response_mode") == "json_drilldown_to_excel"
        ):
            return copy.deepcopy(item)
    raise RuntimeError(f"配置中未找到可复用的 city_ops 下探请求: {config_path}")


def build_export_report(source_report: dict, max_workers: int, max_requests: int) -> dict:
    report = copy.deepcopy(source_report)
    report["name"] = "数据驾驶舱区域层级原始数据"
    report["enabled"] = True
    report["data"] = copy.deepcopy(report.get("data") or {})
    report["data"].update(
        {
            "areaId": ROOT_CODE,
            "queryDate": datetime.now().strftime("%Y%m%d"),
        }
    )

    # A small valid indicator keeps the hierarchy response compact.
    report["data"]["indCodes"] = "sgs_ajvwdz"
    report["data"]["diyCodes"] = "sgs_ajvwdz"
    report["drilldown"] = {
        "data_path": "result.tableData",
        "request_area_field": "areaId",
        "next_area_field": "areaCode",
        "levels": [item[0] for item in DRILLDOWN_LEVELS],
        "skip_self_row": True,
        "max_requests": max(1, max_requests),
        "max_workers": max(1, min(max_workers, 32)),
    }
    report["excel"] = {
        "sheet_name": "原始层级",
        "columns": [
            {"field": "__level_name", "header": "层级"},
            {"field": "__level_index", "header": "层级索引", "type": "number"},
            {"field": "__request_area_id", "header": "父级编码"},
            {"field": "areaCode", "header": "区域编码"},
            {"field": "areaName", "header": "区域名称"},
        ],
    }
    return report


def read_raw_rows(raw_path: Path) -> list[dict]:
    workbook = load_workbook(raw_path, read_only=True, data_only=True)
    worksheet = workbook.active
    headers = [cell.value for cell in next(worksheet.iter_rows())]
    rows = []
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        if row.get("区域编码") and row.get("区域名称"):
            rows.append(row)
    workbook.close()
    return rows


def normalize_rows(raw_rows: list[dict], collected_at: str) -> list[dict]:
    level_map = {
        label: {"level_type": level_type, "level_no": level_no}
        for label, level_type, level_no in DRILLDOWN_LEVELS
    }
    rows = [
        {
            "area_code": ROOT_CODE,
            "area_name": ROOT_NAME,
            "level_type": "CITY",
            "level_no": 1,
            "parent_code": None,
            "enabled": 1,
            "collected_at": collected_at,
        }
    ]
    seen = {("CITY", ROOT_CODE)}
    for raw_row in raw_rows:
        area_code = str(raw_row.get("区域编码") or "").strip()
        level = level_map.get(str(raw_row.get("层级") or "").strip())
        if not level:
            continue
        identity = (level["level_type"], area_code)
        if not area_code or identity in seen:
            continue
        rows.append(
            {
                "area_code": area_code,
                "area_name": str(raw_row.get("区域名称") or "").strip(),
                "level_type": level["level_type"],
                "level_no": level["level_no"],
                "parent_code": str(raw_row.get("父级编码") or "").strip() or None,
                "enabled": 1,
                "collected_at": collected_at,
            }
        )
        seen.add(identity)
    return rows


def style_worksheet(worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=10, color="000000")

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center")

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 26
    widths = [24, 32, 20, 10, 24, 10, 22]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width


def write_final_workbook(rows: list[dict], output_path: Path, request_count: int) -> None:
    output_path = output_path.resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "区域层级"
    headers = [
        "area_code",
        "area_name",
        "level_type",
        "level_no",
        "parent_code",
        "enabled",
        "collected_at",
    ]
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(field) for field in headers])
    style_worksheet(worksheet)

    summary = workbook.create_sheet("导出说明")
    summary_rows = [
        ("项目", "内容"),
        ("根节点", f"{ROOT_CODE} / {ROOT_NAME}"),
        ("采集层级", "CITY → BRANCH → GRID → CHANNEL_MANAGER → CHANNEL"),
        ("人员层", "不采集"),
        ("区域数量", len(rows)),
        ("平台请求数", request_count),
        ("父子关系", "parent_code 指向父级 area_code"),
        ("用途", "核对组织层级，并作为 dashboard_area 初始化数据参考"),
    ]
    for row in summary_rows:
        summary.append(row)
    style_worksheet(summary)
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 72

    workbook.save(output_path)


def export_areas(args: argparse.Namespace) -> tuple[Path, int, int]:
    config_path = args.config.resolve()
    cookie_dump_path = args.cookie_dump.resolve()
    output_path = args.output.resolve()
    if not config_path.exists():
        raise FileNotFoundError(f"配置文件不存在: {config_path}")
    if not cookie_dump_path.exists():
        raise FileNotFoundError(f"Cookie 文件不存在，请先完成自动登录: {cookie_dump_path}")

    source_report = load_source_report(config_path)
    report = build_export_report(source_report, args.max_workers, args.max_requests)
    with tempfile.TemporaryDirectory(prefix="dashboard-area-export-") as temp_dir_value:
        temp_dir = Path(temp_dir_value)
        manifest = download_reports(
            {
                "cookie_dump_path": str(cookie_dump_path),
                "output_dir": str(temp_dir),
                "manifest_path": str(temp_dir / "manifest.json"),
                "request_timeout_seconds": 120,
                "verify_ssl": False,
                "trust_env": False,
                "proxies": {},
                "reports": [report],
            }
        )
        result = manifest["results"][0]
        raw_rows = read_raw_rows(Path(result["output_path"]))
        collected_at = datetime.now().isoformat(timespec="seconds")
        rows = normalize_rows(raw_rows, collected_at)
        write_final_workbook(rows, output_path, int(result.get("requests") or 1))
        return output_path, len(rows), int(result.get("requests") or 1)


def main() -> int:
    args = parse_args()
    output_path, row_count, request_count = export_areas(args)
    print(f"导出完成: {output_path}")
    print(f"区域数量: {row_count}")
    print(f"平台请求数: {request_count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
