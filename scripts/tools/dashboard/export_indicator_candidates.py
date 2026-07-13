"""Scan report JSON files and export dashboard indicator candidates to Excel."""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


PROJECT_DIR = Path(__file__).resolve().parents[3]
DEFAULT_CONFIG_DIR = PROJECT_DIR / "config" / "reports"
DEFAULT_OUTPUT = PROJECT_DIR / "docs" / "数据驾驶舱指标候选.xlsx"
NON_INDICATOR_FIELDS = {
    "__level_name",
    "__level_index",
    "__parent_area_id",
    "__request_area_id",
    "areaCode",
    "areaName",
}


@dataclass
class Candidate:
    code: str
    names: set[str] = field(default_factory=set)
    source_fields: set[str] = field(default_factory=set)
    reports: set[str] = field(default_factory=set)
    downloads: set[str] = field(default_factory=set)
    occurrence_count: int = 0
    output_configured: bool = False
    value_types: set[str] = field(default_factory=set)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="扫描报表配置并生成指标候选 Excel")
    parser.add_argument("--config-dir", type=Path, default=DEFAULT_CONFIG_DIR)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def split_codes(value: object) -> list[str]:
    if not isinstance(value, str):
        return []
    return [part.strip() for part in value.split(",") if part.strip()]


def configured_columns(download: dict) -> dict[str, dict]:
    columns = (download.get("excel") or {}).get("columns") or []
    return {
        str(column.get("field") or "").strip(): column
        for column in columns
        if str(column.get("field") or "").strip()
    }


def scan_configs(config_dir: Path) -> tuple[dict[str, Candidate], list[dict], list[str]]:
    candidates: dict[str, Candidate] = {}
    occurrences: list[dict] = []
    scanned_files = []

    for config_path in sorted(config_dir.glob("*.json")):
        config = json.loads(config_path.read_text(encoding="utf-8"))
        scanned_files.append(config_path.name)
        report_name = str(config.get("name") or config.get("id") or config_path.stem)
        for index, download in enumerate(config.get("downloads") or []):
            data = download.get("data") or {}
            code_sources: dict[str, set[str]] = defaultdict(set)
            for source_field in ("indCodes", "diyCodes"):
                for code in split_codes(data.get(source_field)):
                    code_sources[code].add(source_field)
            if not code_sources:
                continue

            columns = configured_columns(download)
            download_name = str(download.get("name") or f"downloads[{index}]")
            for code, source_fields in code_sources.items():
                candidate = candidates.setdefault(code, Candidate(code=code))
                candidate.source_fields.update(source_fields)
                candidate.reports.add(report_name)
                candidate.downloads.add(download_name)
                candidate.occurrence_count += 1

                column = columns.get(code)
                header = str((column or {}).get("header") or "").strip()
                value_type = str((column or {}).get("type") or "").strip().upper()
                if header and header != code:
                    candidate.names.add(header)
                if column:
                    candidate.output_configured = True
                if value_type:
                    candidate.value_types.add(value_type)

                occurrences.append(
                    {
                        "配置文件": config_path.name,
                        "报表名称": report_name,
                        "下载项": download_name,
                        "stage": str(download.get("stage") or ""),
                        "指标编码": code,
                        "参数来源": "+".join(sorted(source_fields)),
                        "中文名称": header if header != code else "",
                        "已配置输出列": "是" if column else "否",
                        "输出类型": value_type,
                        "启用下载项": "是" if download.get("enabled", True) else "否",
                    }
                )
    return candidates, occurrences, scanned_files


def candidate_rows(candidates: dict[str, Candidate]) -> list[dict]:
    rows = []
    for sort_order, candidate in enumerate(
        sorted(candidates.values(), key=lambda item: item.code.lower()),
        start=1,
    ):
        names = sorted(candidate.names)
        has_name = bool(names)
        rows.append(
            {
                "indicator_code": candidate.code,
                "indicator_name": " / ".join(names),
                "request_source": "+".join(sorted(candidate.source_fields)),
                "value_type": (
                    "/".join(sorted(candidate.value_types))
                    if candidate.value_types
                    else "NUMBER"
                ),
                "unit": "",
                "suggested_enabled": "是" if has_name and candidate.output_configured else "待确认",
                "sort_order": sort_order * 10,
                "output_configured": "是" if candidate.output_configured else "否",
                "report_count": len(candidate.reports),
                "occurrence_count": candidate.occurrence_count,
                "source_reports": "；".join(sorted(candidate.reports)),
                "source_downloads": "；".join(sorted(candidate.downloads)),
                "review_note": "" if has_name else "配置中未找到中文列名，请确认名称和是否需要采集",
            }
        )
    return rows


def style_sheet(worksheet, widths: list[int]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=10, color="000000")

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 28
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    if worksheet.max_row >= 2 and worksheet.max_column >= 1:
        table = Table(
            displayName=f"Table_{worksheet.title.replace(' ', '_')}",
            ref=worksheet.dimensions,
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)


def write_workbook(
    output_path: Path,
    candidates: list[dict],
    occurrences: list[dict],
    scanned_files: list[str],
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()

    candidate_sheet = workbook.active
    candidate_sheet.title = "指标候选"
    candidate_headers = list(candidates[0].keys()) if candidates else ["indicator_code"]
    candidate_sheet.append(candidate_headers)
    for row in candidates:
        candidate_sheet.append([row.get(header) for header in candidate_headers])
    style_sheet(
        candidate_sheet,
        [30, 32, 20, 16, 12, 16, 14, 16, 14, 16, 35, 45, 48],
    )

    occurrence_sheet = workbook.create_sheet("配置出现明细")
    occurrence_headers = list(occurrences[0].keys()) if occurrences else ["配置文件"]
    occurrence_sheet.append(occurrence_headers)
    for row in occurrences:
        occurrence_sheet.append([row.get(header) for header in occurrence_headers])
    style_sheet(occurrence_sheet, [30, 28, 36, 16, 30, 20, 35, 16, 16, 16])

    notes_sheet = workbook.create_sheet("扫描说明")
    notes = [
        ("项目", "内容"),
        ("扫描时间", datetime.now().isoformat(timespec="seconds")),
        ("扫描目录", str(DEFAULT_CONFIG_DIR.relative_to(PROJECT_DIR))),
        ("扫描文件数", len(scanned_files)),
        ("扫描文件", "；".join(scanned_files)),
        ("候选指标数", len(candidates)),
        ("出现明细数", len(occurrences)),
        ("指标来源", "合并 downloads[].data.indCodes 和 diyCodes，自动去除空编码"),
        ("中文名称", "优先匹配同一下载项 excel.columns 中 field 相同的 header"),
        ("建议启用", "有中文名称且已配置输出列时标记为“是”，否则标记为“待确认”"),
        ("注意", "本文件是候选清单，不会自动写入 indicator 表"),
    ]
    for row in notes:
        notes_sheet.append(row)
    style_sheet(notes_sheet, [24, 100])

    workbook.save(output_path)


def verify_workbook(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        formula_errors = []
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith(
                        ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?")
                    ):
                        formula_errors.append(f"{worksheet.title}!{cell.coordinate}")
        return {
            "sheets": workbook.sheetnames,
            "candidate_rows": workbook["指标候选"].max_row - 1,
            "occurrence_rows": workbook["配置出现明细"].max_row - 1,
            "formula_errors": formula_errors,
        }
    finally:
        workbook.close()


def main() -> int:
    args = parse_args()
    config_dir = args.config_dir.resolve()
    output_path = args.output.resolve()
    candidates, occurrences, scanned_files = scan_configs(config_dir)
    rows = candidate_rows(candidates)
    write_workbook(output_path, rows, occurrences, scanned_files)
    verification = verify_workbook(output_path)
    print(f"导出完成: {output_path}")
    print(f"候选指标: {verification['candidate_rows']}")
    print(f"配置出现明细: {verification['occurrence_rows']}")
    print(f"工作表: {', '.join(verification['sheets'])}")
    print(f"公式错误: {len(verification['formula_errors'])}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
