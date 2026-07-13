"""Export a full raw sample for one dashboard indicator without writing MySQL."""

from __future__ import annotations

import argparse
import copy
import json
import re
import sys
import tempfile
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_DIR = Path(__file__).resolve().parents[3]
if str(PROJECT_DIR) not in sys.path:
    sys.path.insert(0, str(PROJECT_DIR))

from services.method_service import download_reports


DEFAULT_CONFIG_PATH = PROJECT_DIR / "config" / "reports" / "家客和存量通报.json"
DEFAULT_COOKIE_DUMP_PATH = PROJECT_DIR / "runtime" / "cookies" / "cookie_dump.json"
DEFAULT_OUTPUT_PATH = PROJECT_DIR / "docs" / "数据驾驶舱单指标采集样本.xlsx"
DEFAULT_INDICATOR_CODE = "sgs_ajvwdz"
DEFAULT_INDICATOR_NAME = "爱家亲情网(V网版)"
ROOT_CODE = "A"

RESPONSE_LEVELS = ["BRANCH", "GRID", "CHANNEL_MANAGER", "CHANNEL"]
REQUEST_LEVELS = ["CITY", "BRANCH", "GRID", "CHANNEL_MANAGER"]
FORMAL_AREA_LEVELS = {"CITY", "BRANCH", "GRID", "CHANNEL"}
INTEGER_PATTERN = re.compile(r"^[+-]?\d+$")
DECIMAL_PATTERN = re.compile(r"^[+-]?(?:\d+\.\d*|\.\d+)$")
EMPTY_MARKERS = {"", "--", "-", "null", "none", "n/a", "nan"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="完整采集单个指标并导出原始样本，不写入 MySQL"
    )
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG_PATH)
    parser.add_argument("--cookie-dump", type=Path, default=DEFAULT_COOKIE_DUMP_PATH)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_PATH)
    parser.add_argument("--indicator-code", default=DEFAULT_INDICATOR_CODE)
    parser.add_argument("--indicator-name", default=DEFAULT_INDICATOR_NAME)
    parser.add_argument("--query-date", default=datetime.now().strftime("%Y%m%d"))
    parser.add_argument("--max-workers", type=int, default=16)
    parser.add_argument("--max-requests", type=int, default=5000)
    return parser.parse_args()


def load_source_report(config_path: Path) -> dict:
    config = json.loads(config_path.read_text(encoding="utf-8"))
    for item in config.get("downloads") or []:
        if (
            item.get("enabled", True)
            and item.get("stage") == "city_ops"
            and item.get("response_mode") == "json_drilldown_to_excel"
        ):
            return copy.deepcopy(item)
    raise RuntimeError(f"配置中未找到 city_ops 下探请求: {config_path}")


def build_report(args: argparse.Namespace) -> dict:
    report = load_source_report(args.config.resolve())
    report["name"] = "数据驾驶舱单指标原始样本"
    report["data"] = copy.deepcopy(report.get("data") or {})
    report["data"].update(
        {
            "areaId": ROOT_CODE,
            "queryDate": args.query_date,
            "indCodes": args.indicator_code,
            "diyCodes": args.indicator_code,
        }
    )
    report["drilldown"] = {
        "data_path": "result.tableData",
        "request_area_field": "areaId",
        "next_area_field": "areaCode",
        "levels": RESPONSE_LEVELS,
        "skip_self_row": False,
        "max_requests": max(1, args.max_requests),
        "max_workers": max(1, min(args.max_workers, 32)),
    }
    report["excel"] = {
        "sheet_name": "平台原始结果",
        "columns": [
            {"field": "__level_index", "header": "响应层级索引"},
            {"field": "__level_name", "header": "响应层级"},
            {"field": "__parent_area_id", "header": "上级请求编码"},
            {"field": "__request_area_id", "header": "本次请求areaId"},
            {"field": "areaCode", "header": "返回区域编码"},
            {"field": "areaName", "header": "返回区域名称"},
            {"field": args.indicator_code, "header": "原始指标值"},
        ],
    }
    return report


def read_raw_rows(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        headers = list(next(iterator))
        return [dict(zip(headers, values)) for values in iterator]
    finally:
        workbook.close()


def classify_value(value: object) -> tuple[str, Decimal | None]:
    if value is None:
        return "EMPTY", None
    if isinstance(value, bool):
        return "NON_NUMERIC", None
    if isinstance(value, int):
        return "INTEGER", Decimal(value)
    if isinstance(value, float):
        return "INTEGER" if value.is_integer() else "DECIMAL", Decimal(str(value))

    text = str(value).strip()
    if text.lower() in EMPTY_MARKERS:
        return "EMPTY", None
    normalized = text.replace(",", "")
    if INTEGER_PATTERN.fullmatch(normalized):
        return "INTEGER", Decimal(normalized)
    if DECIMAL_PATTERN.fullmatch(normalized):
        try:
            return "DECIMAL", Decimal(normalized)
        except InvalidOperation:
            pass
    return "NON_NUMERIC", None


def enrich_rows(raw_rows: list[dict]) -> list[dict]:
    enriched = []
    for row in raw_rows:
        level_index = int(row.get("响应层级索引") or 0)
        request_code = str(row.get("本次请求areaId") or "").strip()
        area_code = str(row.get("返回区域编码") or "").strip()
        is_self = area_code == request_code
        node_type = (
            REQUEST_LEVELS[level_index]
            if is_self and level_index < len(REQUEST_LEVELS)
            else RESPONSE_LEVELS[level_index]
        )
        value_kind, numeric_value = classify_value(row.get("原始指标值"))
        enriched.append(
            {
                "query_area_id": request_code,
                "area_code": area_code,
                "area_name": str(row.get("返回区域名称") or "").strip(),
                "node_type": node_type,
                "is_self_summary": "是" if is_self else "否",
                "is_dashboard_area": "是" if node_type in FORMAL_AREA_LEVELS else "否",
                "parent_request_code": str(row.get("上级请求编码") or "").strip(),
                "raw_value": row.get("原始指标值"),
                "value_kind": value_kind,
                "numeric_value": float(numeric_value) if numeric_value is not None else None,
            }
        )
    return enriched


def style_sheet(worksheet, widths: list[int]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=10, color="000000")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 26
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width


def write_workbook(
    rows: list[dict],
    output_path: Path,
    args: argparse.Namespace,
    request_count: int,
    elapsed_seconds: float,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()

    raw_sheet = workbook.active
    raw_sheet.title = "原始采集"
    raw_headers = [
        "query_area_id",
        "area_code",
        "area_name",
        "node_type",
        "is_self_summary",
        "is_dashboard_area",
        "parent_request_code",
        "raw_value",
        "value_kind",
        "numeric_value",
    ]
    raw_sheet.append(raw_headers)
    for row in rows:
        raw_sheet.append([row.get(header) for header in raw_headers])
    style_sheet(raw_sheet, [24, 24, 42, 22, 18, 18, 24, 18, 18, 18])

    quality_sheet = workbook.create_sheet("数据质量")
    quality_sheet.append(["分类", "数量", "示例"])
    value_examples: dict[str, list[str]] = defaultdict(list)
    counts = Counter(row["value_kind"] for row in rows)
    for row in rows:
        value = row["raw_value"]
        example = repr(value)
        if example not in value_examples[row["value_kind"]] and len(value_examples[row["value_kind"]]) < 10:
            value_examples[row["value_kind"]].append(example)
    for kind in ("INTEGER", "DECIMAL", "EMPTY", "NON_NUMERIC"):
        quality_sheet.append([kind, counts.get(kind, 0), "；".join(value_examples[kind])])
    style_sheet(quality_sheet, [24, 16, 100])

    level_sheet = workbook.create_sheet("层级汇总")
    level_sheet.append(["节点类型", "总行数", "正式区域行数", "空值数", "非数字数"])
    for level in ("CITY", "BRANCH", "GRID", "CHANNEL_MANAGER", "CHANNEL"):
        level_rows = [row for row in rows if row["node_type"] == level]
        level_sheet.append(
            [
                level,
                len(level_rows),
                sum(row["is_dashboard_area"] == "是" for row in level_rows),
                sum(row["value_kind"] == "EMPTY" for row in level_rows),
                sum(row["value_kind"] == "NON_NUMERIC" for row in level_rows),
            ]
        )
    style_sheet(level_sheet, [28, 18, 20, 16, 16])

    notes_sheet = workbook.create_sheet("执行说明")
    notes = [
        ("项目", "内容"),
        ("指标编码", args.indicator_code),
        ("指标名称", args.indicator_name),
        ("查询日期", args.query_date),
        ("根请求", ROOT_CODE),
        ("平台请求数", request_count),
        ("采集行数", len(rows)),
        ("采集耗时（秒）", elapsed_seconds),
        ("正式区域", "CITY、BRANCH、GRID、CHANNEL"),
        ("渠道经理", "仅用于请求路径分析，不写入 area"),
        ("原始值策略", "不强制转换；raw_value 保留平台返回值"),
        ("数据库写入", "本脚本不连接、不写入驾驶舱 MySQL"),
    ]
    for row in notes:
        notes_sheet.append(row)
    style_sheet(notes_sheet, [28, 100])

    workbook.save(output_path)


def export_sample(args: argparse.Namespace) -> tuple[Path, list[dict], dict]:
    if not args.cookie_dump.resolve().exists():
        raise FileNotFoundError(f"Cookie 文件不存在: {args.cookie_dump.resolve()}")
    report = build_report(args)
    with tempfile.TemporaryDirectory(prefix="dashboard-metric-sample-") as temp_value:
        temp_dir = Path(temp_value)
        manifest = download_reports(
            {
                "cookie_dump_path": str(args.cookie_dump.resolve()),
                "output_dir": str(temp_dir),
                "manifest_path": str(temp_dir / "manifest.json"),
                "request_timeout_seconds": 120,
                "verify_ssl": False,
                "trust_env": False,
                "proxies": {},
                "reports": [report],
            }
        )
        result = manifest["results"][0]
        rows = enrich_rows(read_raw_rows(Path(result["output_path"])))
        output_path = args.output.resolve()
        write_workbook(
            rows,
            output_path,
            args,
            int(result.get("requests") or 1),
            float((result.get("timings") or {}).get("total_seconds") or 0),
        )
        return output_path, rows, result


def main() -> int:
    args = parse_args()
    output_path, rows, result = export_sample(args)
    counts = Counter(row["value_kind"] for row in rows)
    print(f"导出完成: {output_path}")
    print(f"平台请求数: {result.get('requests', 1)}")
    print(f"采集行数: {len(rows)}")
    print(
        "数值分类: "
        + ", ".join(
            f"{kind}={counts.get(kind, 0)}"
            for kind in ("INTEGER", "DECIMAL", "EMPTY", "NON_NUMERIC")
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
