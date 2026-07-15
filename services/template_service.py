"""Template copy and update service."""

from __future__ import annotations

import json
import shutil
from time import perf_counter
from datetime import datetime
from pathlib import Path

from openpyxl.cell.cell import MergedCell
from openpyxl import load_workbook


from infrastructure.excel_client import get_sheet, open_excel
from services.runtime_paths import is_runtime_relative_path, resolve_runtime_relative_path

PROJECT_DIR = Path(__file__).resolve().parents[1]
TEMPLATE_UPDATE_ENGINES = {"hybrid", "com_copy"}


def load_json(path):
    resolved = Path(path).resolve()
    with resolved.open("r", encoding="utf-8") as f:
        return json.load(f), resolved


def resolve_path(value, base_dir=PROJECT_DIR):
    if not value:
        return None
    path = Path(value)
    if is_runtime_relative_path(path):
        return resolve_runtime_relative_path(path)
    if path.is_absolute():
        return path
    return (base_dir / path).resolve()


def load_compare_context(config, base_dir=PROJECT_DIR):
    compare_result_path = resolve_path(config.get("compare_result_path"), base_dir)
    if not compare_result_path:
        raise ValueError("缺少 compare_result_path")
    compare_result, _ = load_json(compare_result_path)

    compare_config = {}
    compare_config_dir = None
    compare_config_path = resolve_path(config.get("compare_config_path"), base_dir)
    if compare_config_path and compare_config_path.exists():
        compare_config, resolved_compare_config_path = load_json(compare_config_path)
        compare_config_dir = resolved_compare_config_path.parent

    source_report_path = resolve_path(config.get("source_report_path"), base_dir)
    template_path = resolve_path(config.get("template_path"), base_dir)
    if not source_report_path and compare_config:
        source_report_path = resolve_path(compare_config.get("new_report_path"), compare_config_dir)
    if not template_path and compare_config:
        template_path = resolve_path(compare_config.get("template_path"), compare_config_dir)

    sheets_have_source_paths = all(
        item.get("source_report_path")
        for item in compare_result.get("sheets", [])
    )
    if not source_report_path and not sheets_have_source_paths:
        raise ValueError("缺少 source_report_path，也无法从 compare_config_path 中读取 new_report_path")
    if not template_path:
        raise ValueError("缺少 template_path，也无法从 compare_config_path 中读取 template_path")
    return compare_result, source_report_path, template_path


def should_write_sheet(sheet_result, write_sheets):
    result = sheet_result.get("result")
    if write_sheets == "all_compared":
        return result in {"same", "changed"}
    if write_sheets == "changed":
        return result == "changed"
    raise ValueError("write_sheets 只支持 changed 或 all_compared")


def output_template_path(template_path, output_dir):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return output_dir / f"{template_path.stem}_updated_{timestamp}{template_path.suffix}"


def clear_target_sheet(sheet):
    used = sheet.UsedRange
    used.Clear()


def copy_sheet_content(source_sheet, target_sheet):
    clear_target_sheet(target_sheet)
    source_sheet.UsedRange.Copy()
    target_sheet.Range("A1").PasteSpecial(Paste=-4104)


def normalize_template_update_engine(engine):
    normalized = str(engine or "hybrid").strip().lower()
    if normalized not in TEMPLATE_UPDATE_ENGINES:
        raise ValueError("template_update.engine 只支持 hybrid 或 com_copy")
    return normalized


def worksheet_value_bounds(sheet):
    max_row = 0
    max_column = 0
    for row in sheet.iter_rows():
        row_has_value = False
        for cell in row:
            if cell.value is not None:
                row_has_value = True
                max_column = max(max_column, cell.column)
        if row_has_value:
            max_row = max(max_row, row[0].row)
    return max_row, max_column


def worksheet_values(sheet):
    max_row, max_column = worksheet_value_bounds(sheet)
    if not max_row or not max_column:
        return []
    return [
        [sheet.cell(row=row_index, column=column_index).value for column_index in range(1, max_column + 1)]
        for row_index in range(1, max_row + 1)
    ]


def clear_target_data_area(sheet, rows, columns):
    for row_index in range(1, rows + 1):
        for column_index in range(1, columns + 1):
            cell = sheet.cell(row=row_index, column=column_index)
            if not isinstance(cell, MergedCell) and cell.data_type != "f":
                cell.value = None


def write_values_to_target_sheet(source_sheet, target_sheet):
    values = worksheet_values(source_sheet)
    source_rows = len(values)
    source_columns = max((len(row) for row in values), default=0)
    target_rows, target_columns = worksheet_value_bounds(target_sheet)
    clear_target_data_area(
        target_sheet,
        max(source_rows, target_rows),
        max(source_columns, target_columns),
    )
    for row_index, row_values in enumerate(values, start=1):
        for column_index, value in enumerate(row_values, start=1):
            cell = target_sheet.cell(row=row_index, column=column_index)
            if not isinstance(cell, MergedCell):
                cell.value = value


def refresh_open_workbook(excel, workbook):
    try:
        workbook.RefreshAll()
    except Exception:
        pass
    try:
        excel.CalculateUntilAsyncQueriesDone()
    except Exception:
        pass
    try:
        excel.CalculateFullRebuild()
    except Exception:
        try:
            workbook.Application.Calculate()
        except Exception:
            pass
    workbook.Save()


def refresh_workbook_with_excel(output_path, visible=False):
    excel = open_excel(visible=visible, manual_calculation=True)
    workbook = None
    try:
        workbook = excel.Workbooks.Open(
            str(output_path),
            UpdateLinks=0,
            ReadOnly=False,
            IgnoreReadOnlyRecommended=True,
        )
        refresh_open_workbook(excel, workbook)
    finally:
        if workbook is not None:
            workbook.Close(SaveChanges=False)
        excel.Quit()


def update_condition_met(sheet_results, update_condition):
    if update_condition == "any_changed":
        return any(item.get("result") == "changed" for item in sheet_results)
    if update_condition == "all_changed":
        return bool(sheet_results) and all(item.get("result") == "changed" for item in sheet_results)
    raise ValueError("update_condition 只支持 any_changed 或 all_changed")


def update_template_copy(source_report_path, template_path, output_path, sheet_results, write_sheets, visible=False):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)

    excel = open_excel(visible=visible)
    source_workbook = None
    target_workbook = None
    updated_sheets = []
    try:
        source_workbook = excel.Workbooks.Open(
            str(source_report_path),
            UpdateLinks=0,
            ReadOnly=True,
            IgnoreReadOnlyRecommended=True,
        )
        target_workbook = excel.Workbooks.Open(
            str(output_path),
            UpdateLinks=0,
            ReadOnly=False,
            IgnoreReadOnlyRecommended=True,
        )

        for sheet_result in sheet_results:
            if not should_write_sheet(sheet_result, write_sheets):
                continue
            source_sheet_name = sheet_result["new_sheet_name"]
            target_sheet_name = sheet_result["template_sheet_name"]
            source_sheet = get_sheet(source_workbook, source_sheet_name)
            target_sheet = get_sheet(target_workbook, target_sheet_name)
            copy_sheet_content(source_sheet, target_sheet)
            updated_sheets.append({
                "name": sheet_result.get("name") or source_sheet_name,
                "source_sheet_name": source_sheet_name,
                "template_sheet_name": target_sheet_name,
                "compare_result": sheet_result.get("result"),
            })

        refresh_open_workbook(excel, target_workbook)
    finally:
        try:
            excel.CutCopyMode = False
        except Exception:
            pass
        if source_workbook is not None:
            source_workbook.Close(SaveChanges=False)
        if target_workbook is not None:
            target_workbook.Close(SaveChanges=False)
        excel.Quit()
    return updated_sheets


def update_template_copy_multi(source_report_path, template_path, output_path, sheet_results, write_sheets, visible=False):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)

    excel = open_excel(visible=visible)
    target_workbook = None
    source_workbooks = {}
    updated_sheets = []
    try:
        target_workbook = excel.Workbooks.Open(
            str(output_path),
            UpdateLinks=0,
            ReadOnly=False,
            IgnoreReadOnlyRecommended=True,
        )

        for sheet_result in sheet_results:
            if not should_write_sheet(sheet_result, write_sheets):
                continue
            source_path = resolve_path(sheet_result.get("source_report_path") or source_report_path)
            if not source_path:
                raise ValueError(f"缺少 sheet 的 source_report_path: {sheet_result}")
            source_key = str(source_path)
            if source_key not in source_workbooks:
                source_workbooks[source_key] = excel.Workbooks.Open(
                    source_key,
                    UpdateLinks=0,
                    ReadOnly=True,
                    IgnoreReadOnlyRecommended=True,
                )
            source_sheet_name = sheet_result["new_sheet_name"]
            target_sheet_name = sheet_result["template_sheet_name"]
            source_sheet = get_sheet(source_workbooks[source_key], source_sheet_name)
            target_sheet = get_sheet(target_workbook, target_sheet_name)
            copy_sheet_content(source_sheet, target_sheet)
            updated_sheets.append({
                "name": sheet_result.get("name") or source_sheet_name,
                "download_name": sheet_result.get("download_name"),
                "source_report_path": source_key,
                "source_sheet_name": source_sheet_name,
                "template_sheet_name": target_sheet_name,
                "compare_result": sheet_result.get("result"),
            })

        target_workbook.Save()
    finally:
        try:
            excel.CutCopyMode = False
        except Exception:
            pass
        for workbook in source_workbooks.values():
            workbook.Close(SaveChanges=False)
        if target_workbook is not None:
            target_workbook.Close(SaveChanges=False)
        excel.Quit()
    return updated_sheets


def update_template_hybrid(source_report_path, template_path, output_path, sheet_results, write_sheets, visible=False):
    timings = {}
    started = perf_counter()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)
    timings["copy_template_seconds"] = round(perf_counter() - started, 3)

    write_started = perf_counter()
    keep_vba = output_path.suffix.lower() == ".xlsm"
    target_workbook = load_workbook(output_path, keep_vba=keep_vba)
    source_workbooks = {}
    updated_sheets = []
    try:
        for sheet_result in sheet_results:
            if not should_write_sheet(sheet_result, write_sheets):
                continue
            source_path = resolve_path(sheet_result.get("source_report_path") or source_report_path)
            if not source_path:
                raise ValueError(f"缺少 sheet 的 source_report_path: {sheet_result}")
            source_key = str(source_path)
            if source_key not in source_workbooks:
                source_workbooks[source_key] = load_workbook(
                    source_key,
                    data_only=True,
                    keep_vba=Path(source_key).suffix.lower() == ".xlsm",
                )
            source_sheet_name = sheet_result["new_sheet_name"]
            target_sheet_name = sheet_result["template_sheet_name"]
            if source_sheet_name not in source_workbooks[source_key].sheetnames:
                raise ValueError(f"源报表缺少 sheet: {source_sheet_name}")
            if target_sheet_name not in target_workbook.sheetnames:
                raise ValueError(f"模板缺少 sheet: {target_sheet_name}")
            write_values_to_target_sheet(
                source_workbooks[source_key][source_sheet_name],
                target_workbook[target_sheet_name],
            )
            updated_sheets.append({
                "name": sheet_result.get("name") or source_sheet_name,
                "download_name": sheet_result.get("download_name"),
                "source_report_path": source_key,
                "source_sheet_name": source_sheet_name,
                "template_sheet_name": target_sheet_name,
                "compare_result": sheet_result.get("result"),
            })

        target_workbook.save(output_path)
    finally:
        target_workbook.close()
        for workbook in source_workbooks.values():
            workbook.close()
    timings["write_values_seconds"] = round(perf_counter() - write_started, 3)

    refresh_started = perf_counter()
    refresh_workbook_with_excel(output_path, visible=visible)
    timings["excel_refresh_seconds"] = round(perf_counter() - refresh_started, 3)
    timings["total_seconds"] = round(sum(timings.values()), 3)
    return updated_sheets, timings


def write_manifest(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


def update_template(config, base_dir=PROJECT_DIR):
    compare_result, source_report_path, template_path = load_compare_context(config, base_dir)

    manifest_path = resolve_path(config.get("manifest_path", "modules/template_updater/output/update_manifest.json"), base_dir)
    output_dir = resolve_path(config.get("output_dir", "modules/template_updater/output/templates"), base_dir)
    engine = normalize_template_update_engine(config.get("engine"))
    write_sheets = config.get("write_sheets", "changed")
    update_condition = config.get("update_condition", "any_changed")
    allow_same_update = bool(config.get("allow_same_update", False))

    result = compare_result.get("result")
    if result == "invalid":
        raise RuntimeError("比对结果为 invalid，停止更新模板")
    if result == "same" and not allow_same_update:
        manifest = {
            "status": "skipped",
            "reason": "compare_result_same",
            "engine": engine,
            "update_condition": update_condition,
            "source_report_path": str(source_report_path),
            "template_path": str(template_path),
            "output_path": None,
            "updated_sheets": [],
            "generated_at": datetime.now().isoformat(),
        }
        write_manifest(manifest_path, manifest)
        return manifest
    if result != "same" and not update_condition_met(compare_result.get("sheets", []), update_condition):
        manifest = {
            "status": "skipped",
            "reason": "update_condition_not_met",
            "engine": engine,
            "update_condition": update_condition,
            "source_report_path": str(source_report_path),
            "template_path": str(template_path),
            "output_path": None,
            "updated_sheets": [],
            "generated_at": datetime.now().isoformat(),
        }
        write_manifest(manifest_path, manifest)
        return manifest
    if result not in {"changed", "same"}:
        raise RuntimeError(f"不支持的比对结果: {result}")

    output_path = output_template_path(template_path, output_dir)
    updater = update_template_copy_multi if engine == "com_copy" else update_template_hybrid
    update_started = perf_counter()
    update_result = updater(
        source_report_path,
        template_path,
        output_path,
        compare_result.get("sheets", []),
        write_sheets,
        visible=bool(config.get("visible", False)),
    )
    if isinstance(update_result, tuple):
        updated_sheets, timings = update_result
    else:
        updated_sheets = update_result
        timings = {"total_seconds": round(perf_counter() - update_started, 3)}
    manifest = {
        "status": "updated",
        "engine": engine,
        "source_report_path": str(source_report_path),
        "template_path": str(template_path),
        "output_path": str(output_path),
        "write_sheets": write_sheets,
        "update_condition": update_condition,
        "updated_sheets": updated_sheets,
        "timings": timings,
        "generated_at": datetime.now().isoformat(),
    }
    write_manifest(manifest_path, manifest)
    return manifest


def update_template_from_config(config_path, base_dir=PROJECT_DIR):
    config_path = resolve_path(config_path, base_dir)
    config, resolved = load_json(config_path)
    return update_template(config, base_dir=resolved.parent)
