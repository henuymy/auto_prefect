"""Read Tencent Docs online sheets and save them as local workbooks."""

from __future__ import annotations

import json
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from time import perf_counter, sleep
from typing import Any
from urllib.parse import quote, urlparse

import requests
from openpyxl import Workbook

from utils.config_loader import load_json_with_runtime_override


PROJECT_DIR = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG_PATH = PROJECT_DIR / "config" / "modules" / "tencent_docs.json"
DOCS_BASE_URL = "https://docs.qq.com"
DOCS_URL_ID_PATTERN = re.compile(r"docs\.qq\.com/(?:doc|sheet|slide|pdf)/([A-Za-z0-9_-]+)", re.I)
CELL_REF_PATTERN = re.compile(r"^([A-Za-z]+)(\d+)$")
RANGE_PATTERN = re.compile(r"^([A-Za-z]+)(\d+):([A-Za-z]+)(\d+)$")


@dataclass(frozen=True)
class TencentDocsCredentials:
    client_id: str
    access_token: str
    open_id: str


class TencentDocsApiError(RuntimeError):
    def __init__(
        self,
        message: str,
        status_code: int | None = None,
        code: Any = None,
        ret: Any = None,
        range_invalid: bool = False,
    ):
        super().__init__(message)
        self.status_code = status_code
        self.code = code
        self.ret = ret
        self.range_invalid = range_invalid


def resolve_path(value: str | Path | None, base_dir: Path = PROJECT_DIR) -> Path:
    if not value:
        return DEFAULT_CONFIG_PATH
    path = Path(value)
    return path if path.is_absolute() else (base_dir / path).resolve()


def load_tencent_docs_config(config_path: str | Path | None = None, base_dir: Path = PROJECT_DIR) -> dict[str, Any]:
    path = resolve_path(config_path, base_dir=base_dir)
    payload, _ = load_json_with_runtime_override(path)
    return payload


def load_credentials(config: dict[str, Any]) -> TencentDocsCredentials:
    credentials = config.get("credentials") or {}
    client_id = str(credentials.get("client_id") or "").strip()
    access_token = str(credentials.get("access_token") or "").strip()
    open_id = str(credentials.get("open_id") or "").strip()
    missing = [
        name
        for name, value in {
            "client_id": client_id,
            "access_token": access_token,
            "open_id": open_id,
        }.items()
        if not value
    ]
    if missing:
        raise ValueError(
            f"腾讯文档凭据缺少字段: {missing}，请检查 "
            "config/runtime.local.json 的 module_overrides.tencent_docs"
        )
    return TencentDocsCredentials(client_id=client_id, access_token=access_token, open_id=open_id)


def auth_headers(credentials: TencentDocsCredentials) -> dict[str, str]:
    return {
        "Access-Token": credentials.access_token,
        "Client-Id": credentials.client_id,
        "Open-Id": credentials.open_id,
    }


def parse_api_json(response: requests.Response) -> dict[str, Any]:
    try:
        payload = response.json()
    except ValueError as exc:
        content_type = response.headers.get("Content-Type", "")
        raise TencentDocsApiError(
            f"腾讯文档接口响应不是 JSON: HTTP {response.status_code}, content_type={content_type}",
            status_code=response.status_code,
        ) from exc
    if response.status_code < 200 or response.status_code >= 300:
        raise TencentDocsApiError(
            f"腾讯文档接口 HTTP 错误: HTTP {response.status_code}",
            status_code=response.status_code,
        )
    ret = payload.get("ret")
    code = payload.get("code")
    if ret not in (None, 0):
        raise TencentDocsApiError(
            f"腾讯文档接口返回失败: ret={ret}",
            status_code=response.status_code,
            ret=ret,
        )
    if code not in (None, 0):
        raise TencentDocsApiError(
            f"腾讯文档接口返回失败: code={code}",
            status_code=response.status_code,
            code=code,
            range_invalid=is_range_invalid_message(payload.get("message")),
        )
    return payload


def retry_settings(config: dict[str, Any], report: dict[str, Any] | None = None) -> dict[str, Any]:
    report = report or {}
    retry = config.get("retry") or {}
    return {
        "attempts": int(report.get("request_retry_attempts") or retry.get("attempts") or 3),
        "backoff_seconds": float(report.get("request_retry_backoff_seconds") or retry.get("backoff_seconds") or 1),
        "max_backoff_seconds": float(report.get("request_retry_max_backoff_seconds") or retry.get("max_backoff_seconds") or 8),
    }


def is_retryable_api_error(exc: TencentDocsApiError) -> bool:
    if exc.status_code in {429, 500, 502, 503, 504}:
        return True
    if exc.code in {429, 500, 502, 503, 504, "429", "500", "502", "503", "504"}:
        return True
    if exc.ret in {429, 500, 502, 503, 504, "429", "500", "502", "503", "504"}:
        return True
    return False


def request_api_json(
    session: requests.Session,
    url: str,
    credentials: TencentDocsCredentials,
    timeout: int,
    retry: dict[str, Any] | None = None,
    params: dict[str, Any] | None = None,
) -> dict[str, Any]:
    retry = retry or {}
    attempts = max(1, int(retry.get("attempts") or 1))
    backoff = max(0, float(retry.get("backoff_seconds") or 0))
    max_backoff = max(backoff, float(retry.get("max_backoff_seconds") or backoff))
    last_error: Exception | None = None
    for attempt in range(1, attempts + 1):
        try:
            response = session.get(url, headers=auth_headers(credentials), params=params, timeout=timeout)
            return parse_api_json(response)
        except requests.RequestException as exc:
            last_error = exc
            retryable = True
        except TencentDocsApiError as exc:
            last_error = exc
            retryable = is_retryable_api_error(exc)
        if attempt >= attempts or not retryable:
            raise last_error
        sleep(min(max_backoff, backoff * (2 ** (attempt - 1))))
    raise RuntimeError("腾讯文档接口请求失败")


def is_range_invalid_error(exc: Exception) -> bool:
    if isinstance(exc, TencentDocsApiError) and exc.range_invalid:
        return True
    text = str(exc)
    return is_range_invalid_message(text)


def is_range_invalid_message(message: Any) -> bool:
    text = str(message or "")
    return "'range' invalid" in text or "RangeSize Validate error" in text


def extract_encoded_id(doc_url: str) -> str:
    value = str(doc_url or "").strip()
    match = DOCS_URL_ID_PATTERN.search(value)
    if match:
        return match.group(1)
    parsed = urlparse(value)
    if parsed.netloc and parsed.path:
        candidate = Path(parsed.path).name
        if candidate:
            return candidate
    return value


def sheet_id_from_doc_url(doc_url: str) -> str:
    value = str(doc_url or "").strip()
    if not value:
        return ""
    try:
        from urllib.parse import parse_qs

        query = parse_qs(urlparse(value).query)
        return (query.get("tab") or [""])[0]
    except Exception:
        match = re.search(r"[?&]tab=([^&#]+)", value)
        return match.group(1) if match else ""


def column_to_index(column: str) -> int:
    result = 0
    for char in column.upper():
        if not ("A" <= char <= "Z"):
            raise ValueError(f"非法列名: {column}")
        result = result * 26 + (ord(char) - ord("A") + 1)
    return result


def index_to_column(index: int) -> str:
    if index < 1:
        raise ValueError("列序号必须从 1 开始")
    chars = []
    while index:
        index, remainder = divmod(index - 1, 26)
        chars.append(chr(ord("A") + remainder))
    return "".join(reversed(chars))


def parse_a1_range(value: str) -> tuple[int, int, int, int]:
    text = str(value or "").strip()
    match = RANGE_PATTERN.match(text)
    if not match:
        single = CELL_REF_PATTERN.match(text)
        if not single:
            raise ValueError(f"读取范围必须是 A1:B2 格式: {value}")
        start_col = end_col = column_to_index(single.group(1))
        start_row = end_row = int(single.group(2))
    else:
        start_col = column_to_index(match.group(1))
        start_row = int(match.group(2))
        end_col = column_to_index(match.group(3))
        end_row = int(match.group(4))
    if start_row < 1 or end_row < 1:
        raise ValueError(f"读取范围行号必须从 1 开始: {value}")
    if end_row < start_row or end_col < start_col:
        raise ValueError(f"读取范围终点不能小于起点: {value}")
    return start_col, start_row, end_col, end_row


def format_a1_range(start_col: int, start_row: int, end_col: int, end_row: int) -> str:
    return f"{index_to_column(start_col)}{start_row}:{index_to_column(end_col)}{end_row}"


def split_a1_range_groups(value: str, limits: dict[str, Any] | None = None) -> list[list[str]]:
    limits = limits or {}
    max_rows = int(limits.get("max_rows_per_request") or 1000)
    max_columns = int(limits.get("max_columns_per_request") or 200)
    max_cells = int(limits.get("max_cells_per_request") or 10000)
    start_col, start_row, end_col, end_row = parse_a1_range(value)
    groups = []
    col_start = start_col
    while col_start <= end_col:
        col_end = min(end_col, col_start + max_columns - 1)
        col_count = col_end - col_start + 1
        rows_per_chunk = max(1, min(max_rows, max_cells // col_count))
        chunks = []
        row_start = start_row
        while row_start <= end_row:
            row_end = min(end_row, row_start + rows_per_chunk - 1)
            chunks.append(format_a1_range(col_start, row_start, col_end, row_end))
            row_start = row_end + 1
        groups.append(chunks)
        col_start = col_end + 1
    return groups


def split_a1_range(value: str, limits: dict[str, Any] | None = None) -> list[str]:
    chunks = []
    for group in split_a1_range_groups(value, limits):
        chunks.extend(group)
    return chunks


def sheet_bounds_from_property(sheet_property: dict[str, Any] | None) -> tuple[int, int] | None:
    if not sheet_property:
        return None
    for row_key, column_key in (("rowCount", "columnCount"), ("rowTotal", "columnTotal"), ("rows", "columns")):
        try:
            rows = int(sheet_property.get(row_key))
            columns = int(sheet_property.get(column_key))
        except (TypeError, ValueError):
            continue
        if rows >= 1 and columns >= 1:
            return rows, columns
    return None


def find_sheet_property(spreadsheet: dict[str, Any], sheet_id: str) -> dict[str, Any] | None:
    target = str(sheet_id)
    for item in spreadsheet.get("properties") or []:
        found = item.get("sheetId") or item.get("sheetID") or item.get("sheet_id")
        if str(found) == target:
            return item
    return None


def resolve_sheet_range(range_address: str, sheet_property: dict[str, Any] | None) -> tuple[str, bool]:
    text = str(range_address or "").strip()
    bounds = sheet_bounds_from_property(sheet_property)
    if text.lower() in {"", "auto", "used", "used_range"}:
        if not bounds:
            raise ValueError("Sheet 读取范围为空，且无法从腾讯文档元数据自动探测行列范围")
        rows, columns = bounds
        return format_a1_range(1, 1, columns, rows), True

    start_col, start_row, end_col, end_row = parse_a1_range(text)
    return format_a1_range(start_col, start_row, end_col, end_row), False


def safe_sheet_name(raw_name: str, used_names: set[str]) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", "_", str(raw_name or "")).strip("'").strip() or "Sheet"
    cleaned = cleaned[:31]
    candidate = cleaned
    index = 2
    while candidate in used_names:
        suffix = f"_{index}"
        candidate = f"{cleaned[:31 - len(suffix)]}{suffix}"
        index += 1
    used_names.add(candidate)
    return candidate


def safe_filename(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value or "")).strip()
    return cleaned or f"tencent-sheet-{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"


def resolve_file_id(session: requests.Session, report: dict[str, Any], credentials: TencentDocsCredentials, timeout: int, retry: dict[str, Any] | None = None) -> str:
    file_id = str(report.get("file_id") or "").strip()
    if file_id:
        return file_id
    doc_url = str(report.get("doc_url") or report.get("url") or "").strip()
    encoded_id = str(report.get("encoded_id") or extract_encoded_id(doc_url)).strip()
    if not encoded_id:
        raise ValueError("腾讯文档抓取缺少 file_id 或 doc_url")
    payload = request_api_json(
        session,
        f"{DOCS_BASE_URL}/openapi/drive/v2/util/converter",
        credentials,
        timeout,
        retry=retry,
        params={"type": 2, "value": encoded_id},
    )
    data = payload.get("data") or {}
    converted = data.get("fileID") or data.get("fileId") or data.get("file_id")
    if not converted:
        raise RuntimeError(f"腾讯文档 fileID 转换响应缺少 fileID: {payload}")
    return str(converted)


def get_spreadsheet(session: requests.Session, file_id: str, credentials: TencentDocsCredentials, timeout: int, concise: bool = True, retry: dict[str, Any] | None = None) -> dict[str, Any]:
    payload = request_api_json(
        session,
        f"{DOCS_BASE_URL}/openapi/spreadsheet/v3/files/{quote(file_id, safe='$')}",
        credentials,
        timeout,
        retry=retry,
        params={"concise": 1 if concise else 0},
    )
    return payload.get("data") or payload


def find_sheet_id(spreadsheet: dict[str, Any], sheet_config: dict[str, Any]) -> str:
    sheet_id = str(sheet_config.get("sheet_id") or "").strip()
    if sheet_id:
        return sheet_id
    target_name = str(sheet_config.get("sheet_name") or "").strip()
    if not target_name:
        raise ValueError("Sheet 必须填写 sheet_id 或 sheet_name")
    for item in spreadsheet.get("properties") or []:
        title = str(item.get("title") or item.get("name") or "").strip()
        if title == target_name:
            found = item.get("sheetId") or item.get("sheetID") or item.get("sheet_id")
            if found:
                return str(found)
    raise RuntimeError(f"腾讯文档中找不到 Sheet: {target_name}")


def get_range(session: requests.Session, file_id: str, sheet_id: str, range_address: str, credentials: TencentDocsCredentials, timeout: int, retry: dict[str, Any] | None = None) -> dict[str, Any]:
    payload = request_api_json(
        session,
        f"{DOCS_BASE_URL}/openapi/spreadsheet/v3/files/{quote(file_id, safe='$')}/{quote(sheet_id, safe='')}/{quote(range_address, safe=':')}",
        credentials,
        timeout,
        retry=retry,
    )
    data = payload.get("data") or payload
    return data.get("gridData") or {}


def cell_to_value(cell: dict[str, Any] | None):
    if not cell:
        return None
    value = cell.get("cellValue")
    if value is None:
        return None
    if not isinstance(value, dict):
        return value
    if "text" in value:
        return text_to_value(value.get("text"))
    if "number" in value:
        number = value.get("number")
        try:
            numeric = float(number)
        except (TypeError, ValueError):
            return number
        return int(numeric) if numeric.is_integer() else numeric
    if "time" in value:
        return time_to_text(value.get("time"))
    if "location" in value:
        location = value.get("location") or {}
        return location.get("name") or json.dumps(location, ensure_ascii=False, separators=(",", ":"))
    if "link" in value:
        link = value.get("link") or {}
        return link.get("text") or link.get("url") or json.dumps(link, ensure_ascii=False, separators=(",", ":"))
    if "select" in value:
        select = value.get("select")
        if isinstance(select, dict):
            return select.get("text") or select.get("name") or select_value_to_text(select)
        return str(select)
    if "value" in value and "options" in value:
        return select_value_to_text(value)
    if "boolValue" in value:
        return bool(value.get("boolValue"))
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def time_to_text(value: Any) -> Any:
    if not isinstance(value, dict):
        return value
    year = value.get("year")
    month = value.get("month")
    day = value.get("day")
    hour = value.get("hour", 0)
    minute = value.get("minute", 0)
    second = value.get("second", 0)
    if year is None or month is None or day is None:
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if int(year) == 1899 and int(month) == 12 and int(day) == 30 and not (hour or minute or second):
        return ""
    date_part = f"{int(year):04d}-{int(month):02d}-{int(day):02d}"
    if hour or minute or second:
        return f"{date_part} {int(hour):02d}:{int(minute):02d}:{int(second):02d}"
    return date_part


def select_value_to_text(value: dict[str, Any]) -> str:
    selected = value.get("value")
    if not isinstance(selected, list):
        selected = [selected] if selected not in (None, "") else []
    option_text_by_id = {
        str(option.get("id")): str(option.get("text") or option.get("name") or "")
        for option in value.get("options") or []
        if isinstance(option, dict)
    }
    texts = []
    for selected_id in selected:
        text = option_text_by_id.get(str(selected_id), str(selected_id or ""))
        if text:
            texts.append(text)
    return ", ".join(texts)


def text_to_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    text = value.strip()
    if not (text.startswith("{") and text.endswith("}")):
        return value
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        return value
    if isinstance(parsed, dict) and "value" in parsed and "options" in parsed:
        return select_value_to_text(parsed)
    return value


def write_grid_to_worksheet(worksheet, grid_data: dict[str, Any], fallback_range: str):
    fallback_col, fallback_row, _, _ = parse_a1_range(fallback_range)
    start_row = int(grid_data.get("startRow") if grid_data.get("startRow") is not None else fallback_row - 1) + 1
    start_col = int(grid_data.get("startColumn") if grid_data.get("startColumn") is not None else fallback_col - 1) + 1
    written = 0
    for row_offset, row in enumerate(grid_data.get("rows") or []):
        for col_offset, cell in enumerate(row.get("values") or []):
            value = cell_to_value(cell)
            worksheet.cell(row=start_row + row_offset, column=start_col + col_offset).value = value
            written += 1
    return written


def write_sheet_ranges(
    session: requests.Session,
    workbook: Workbook,
    file_id: str,
    spreadsheet: dict[str, Any],
    sheet_config: dict[str, Any],
    credentials: TencentDocsCredentials,
    timeout: int,
    limits: dict[str, Any],
    used_sheet_names: set[str],
    retry: dict[str, Any] | None = None,
    chunk_delay: float = 0,
) -> dict[str, Any]:
    sheet_id = find_sheet_id(spreadsheet, sheet_config)
    sheet_property = find_sheet_property(spreadsheet, sheet_id)
    configured_range = str(sheet_config.get("range") or "").strip()
    range_address, range_auto_adjusted = resolve_sheet_range(configured_range, sheet_property)
    output_name = (
        sheet_config.get("output_sheet_name")
        or sheet_config.get("sheet_name")
        or sheet_id
    )
    worksheet = workbook.create_sheet(safe_sheet_name(str(output_name), used_sheet_names))
    chunk_groups = split_a1_range_groups(range_address, limits)
    planned_chunks = sum(len(group) for group in chunk_groups)
    fetched_chunks = 0
    written_cells = 0
    stopped_reason = ""
    for group in chunk_groups:
        group_written = 0
        for chunk in group:
            try:
                grid_data = get_range(session, file_id, sheet_id, chunk, credentials, timeout, retry=retry)
            except RuntimeError as exc:
                if (written_cells or group_written) and is_range_invalid_error(exc):
                    stopped_reason = "range_invalid_after_data"
                    break
                raise
            fetched_chunks += 1
            written = write_grid_to_worksheet(worksheet, grid_data, chunk)
            written_cells += written
            group_written += written
            if chunk_delay > 0:
                sleep(chunk_delay)
    return {
        "sheet_id": sheet_id,
        "sheet_name": sheet_config.get("sheet_name"),
        "output_sheet_name": worksheet.title,
        "configured_range": configured_range,
        "range": range_address,
        "range_auto_adjusted": range_auto_adjusted,
        "planned_chunks": planned_chunks,
        "fetched_chunks": fetched_chunks,
        "chunks": fetched_chunks,
        "cells": written_cells,
        "stopped_reason": stopped_reason,
    }


def download_tencent_sheet_report(
    report: dict[str, Any],
    output_dir: str | Path,
    base_dir: Path = PROJECT_DIR,
    session: requests.Session | None = None,
) -> dict[str, Any]:
    started = perf_counter()
    module_config = load_tencent_docs_config(report.get("tencent_config_path"), base_dir=base_dir)
    credentials = load_credentials(module_config)
    timeout = int(report.get("request_timeout_seconds") or module_config.get("request_timeout_seconds") or 60)
    limits = module_config.get("read_limits") or {}
    retry = retry_settings(module_config, report)
    chunk_delay = float(report.get("chunk_request_interval_seconds") or module_config.get("chunk_request_interval_seconds") or module_config.get("request_interval_seconds") or 0)
    session = session or requests.Session()
    file_id = resolve_file_id(session, report, credentials, timeout, retry=retry)
    spreadsheet = get_spreadsheet(session, file_id, credentials, timeout, concise=False, retry=retry)

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    used_sheet_names: set[str] = set()
    sheet_results = []
    for sheet_config in report.get("sheets") or []:
        sheet_results.append(
            write_sheet_ranges(
                session,
                workbook,
                file_id,
                spreadsheet,
                sheet_config,
                credentials,
                timeout,
                limits,
                used_sheet_names,
                retry=retry,
                chunk_delay=chunk_delay,
            )
        )
        delay = float(report.get("request_interval_seconds") or module_config.get("request_interval_seconds") or 0)
        if delay > 0:
            sleep(delay)
    if not sheet_results:
        raise ValueError("腾讯文档抓取至少需要一个 Sheet 范围")

    output_dir = resolve_path(output_dir, base_dir=base_dir)
    if not output_dir.is_absolute():
        output_dir = base_dir / output_dir
    output_dir.mkdir(parents=True, exist_ok=True)
    filename = safe_filename(report.get("output_filename") or f"{report.get('name') or '腾讯文档'}.xlsx")
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        filename = f"{filename}.xlsx"
    output_path = output_dir / filename
    try:
        workbook.save(output_path)
    except PermissionError as exc:
        raise PermissionError(f"腾讯文档下载结果无法保存，文件可能正被 Excel 打开，请关闭后重试: {output_path}") from exc
    finally:
        workbook.close()
    return {
        "name": report.get("name"),
        "source": "tencent_sheet",
        "transport": "tencent_docs_openapi",
        "file_id": file_id,
        "bytes": output_path.stat().st_size,
        "output_path": str(output_path),
        "sheets": sheet_results,
        "downloaded_at": datetime.now().isoformat(),
        "timings": {"total_seconds": round(perf_counter() - started, 3)},
    }
