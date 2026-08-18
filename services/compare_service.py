"""Compare downloaded report sheets with template sheets."""

from __future__ import annotations

import json
import hashlib
from concurrent.futures import ThreadPoolExecutor
from zipfile import BadZipFile
from dataclasses import asdict
from time import perf_counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from models.compare_result import CompareResult, WorkbookCompareResult
from infrastructure.excel_client import get_sheet, open_excel, sheet_names as workbook_sheet_names
from services.runtime_paths import (
    is_runtime_relative_path,
    resolve_runtime_relative_path,
)


PROJECT_DIR = Path(__file__).resolve().parents[1]
COMPARE_ENGINES = {"openpyxl", "com"}


def load_json(path):
    resolved = Path(path).resolve()
    with resolved.open("r", encoding="utf-8") as f:
        return json.load(f), resolved


def resolve_path(value, base_dir=PROJECT_DIR):
    if not value:
        return None
    path = Path(value)
    if path.parts and path.parts[0].lower() == "runtime":
        raise ValueError(f"运行根相对路径不得以 runtime/ 开头: {value}")
    if is_runtime_relative_path(path):
        return resolve_runtime_relative_path(path)
    if path.is_absolute():
        return path.resolve()
    return (base_dir / path).resolve()


def normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(Decimal(str(value)).normalize())
    if isinstance(value, int):
        return str(value)
    text = str(value).replace("\u3000", " ").strip()
    if text == "":
        return ""
    try:
        decimal_value = Decimal(text.replace(",", ""))
    except (InvalidOperation, ValueError):
        return " ".join(text.split())
    if decimal_value == decimal_value.to_integral():
        return str(decimal_value.to_integral())
    return str(decimal_value.normalize())


def normalize_matrix(values):
    return [[normalize_cell(cell) for cell in row] for row in values]


def is_empty_row(row):
    return all(cell == "" for cell in row)


def trim_empty_edges(rows):
    rows = [list(row) for row in rows if not is_empty_row(row)]
    if not rows:
        return []

    max_width = max(len(row) for row in rows)
    padded = [row + [""] * (max_width - len(row)) for row in rows]
    non_empty_cols = [
        index
        for index in range(max_width)
        if any(row[index] != "" for row in padded)
    ]
    if not non_empty_cols:
        return []
    return [row[non_empty_cols[0]:non_empty_cols[-1] + 1] for row in padded]


def values_to_matrix(values):
    if values is None:
        return []
    if not isinstance(values, tuple):
        return [[values]]
    if values and not isinstance(values[0], tuple):
        return [list(values)]
    return [list(row) for row in values]


def read_sheet_table_com(workbook, sheet_name, range_address=None):
    sheet = get_sheet(workbook, sheet_name)
    rng = sheet.Range(range_address) if range_address else sheet.UsedRange
    return trim_empty_edges(normalize_matrix(values_to_matrix(rng.Value)))


def worksheet_range_values(sheet, range_address=None):
    if range_address:
        cells = sheet[range_address]
        if not isinstance(cells, tuple):
            return [[cells.value]]
        if cells and not isinstance(cells[0], tuple):
            return [[cell.value for cell in cells]]
        return [[cell.value for cell in row] for row in cells]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def iter_worksheet_rows(sheet, range_address=None):
    if range_address:
        yield from worksheet_range_values(sheet, range_address)
    else:
        for row in sheet.iter_rows(values_only=True):
            yield list(row)


def stream_sheet_hash_openpyxl(workbook, sheet_name, range_address=None):
    if sheet_name not in workbook.sheetnames:
        raise KeyError(f"工作表不存在: {sheet_name}")
    digest = hashlib.sha256()
    non_empty_rows = 0
    for row in iter_worksheet_rows(workbook[sheet_name], range_address):
        normalized = [normalize_cell(cell) for cell in row]
        if not is_empty_row(normalized):
            non_empty_rows += 1
        digest.update(str(len(normalized)).encode("ascii"))
        digest.update(b"\x1e")
        for cell in normalized:
            encoded = cell.encode("utf-8")
            digest.update(str(len(encoded)).encode("ascii"))
            digest.update(b"\x1f")
            digest.update(encoded)
            digest.update(b"\x1d")
        digest.update(b"\x1c")
    return {
        "hash": digest.hexdigest(),
        "non_empty_rows": non_empty_rows,
    }


def read_sheet_table_openpyxl(workbook, sheet_name, range_address=None):
    if sheet_name not in workbook.sheetnames:
        raise KeyError(f"工作表不存在: {sheet_name}")
    return trim_empty_edges(normalize_matrix(worksheet_range_values(workbook[sheet_name], range_address)))


def split_header_rows(rows, header_row=1):
    if header_row < 0:
        raise ValueError("header_row 必须大于等于 0")
    if not rows:
        return [], []
    if header_row == 0:
        max_width = max(len(row) for row in rows)
        return [f"列{index}" for index in range(1, max_width + 1)], [row for row in rows if not is_empty_row(row)]
    header_index = header_row - 1
    if header_index >= len(rows):
        raise ValueError(f"header_row={header_row} 超出数据行数 {len(rows)}")
    return rows[header_index], [row for row in rows[header_index + 1:] if not is_empty_row(row)]


def unique_headers(headers):
    result = []
    seen = {}
    for index, header in enumerate(headers, start=1):
        name = header or f"列{index}"
        count = seen.get(name, 0)
        seen[name] = count + 1
        if count:
            name = f"{name}_{count + 1}"
        result.append(name)
    return result


def filter_columns(headers, rows, ignore_columns=None):
    ignored = set(ignore_columns or [])
    indexes = [index for index, header in enumerate(headers) if header not in ignored]
    return (
        [headers[index] for index in indexes],
        [[row[index] if index < len(row) else "" for index in indexes] for row in rows],
    )


def rows_to_dicts(headers, rows):
    return [
        {header: row[index] if index < len(row) else "" for index, header in enumerate(headers)}
        for row in rows
    ]


def row_key(row, key_columns):
    return tuple(row.get(column, "") for column in key_columns)


def compare_without_key(new_rows, old_rows, sample_limit):
    changed_rows = []
    max_rows = max(len(new_rows), len(old_rows))
    changed_count = 0
    for index in range(max_rows):
        new_row = new_rows[index] if index < len(new_rows) else None
        old_row = old_rows[index] if index < len(old_rows) else None
        if new_row != old_row:
            changed_count += 1
            if len(changed_rows) < sample_limit:
                changed_rows.append({"row_number": index + 1, "old": old_row, "new": new_row})
    return {
        "added_rows": max(len(new_rows) - len(old_rows), 0),
        "removed_rows": max(len(old_rows) - len(new_rows), 0),
        "changed_rows": changed_count,
        "sample_diffs": changed_rows,
    }


def compare_with_key(headers, new_rows, old_rows, key_columns, sample_limit):
    missing_keys = [column for column in key_columns if column not in headers]
    if missing_keys:
        return {"invalid_reason": f"主键字段不存在: {missing_keys}"}

    new_items = rows_to_dicts(headers, new_rows)
    old_items = rows_to_dicts(headers, old_rows)
    new_map = {row_key(row, key_columns): row for row in new_items}
    old_map = {row_key(row, key_columns): row for row in old_items}
    if len(new_map) != len(new_items) or len(old_map) != len(old_items):
        return {"invalid_reason": "主键字段存在重复值，无法按主键比对"}

    added_keys = sorted(set(new_map) - set(old_map))
    removed_keys = sorted(set(old_map) - set(new_map))
    changed = []
    for key in sorted(set(new_map) & set(old_map)):
        changed_columns = [
            header for header in headers
            if new_map[key].get(header, "") != old_map[key].get(header, "")
        ]
        if changed_columns:
            changed.append({
                "key": key,
                "changed_columns": changed_columns,
                "old": {column: old_map[key].get(column, "") for column in changed_columns},
                "new": {column: new_map[key].get(column, "") for column in changed_columns},
            })
    return {
        "added_rows": len(added_keys),
        "removed_rows": len(removed_keys),
        "changed_rows": len(changed),
        "sample_added_keys": added_keys[:sample_limit],
        "sample_removed_keys": removed_keys[:sample_limit],
        "sample_diffs": changed[:sample_limit],
    }


def table_hash(headers, rows):
    payload = json.dumps([headers, rows], ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def compare_tables(new_table, old_table, header_row=1, ignore_columns=None, key_columns=None, sample_limit=10):
    new_rows = trim_empty_edges(normalize_matrix(new_table))
    old_rows = trim_empty_edges(normalize_matrix(old_table))
    if not new_rows:
        return CompareResult("invalid", message="新报表数据为空")
    if not old_rows:
        return CompareResult("invalid", new_row_count=len(new_rows), message="模板旧数据为空")

    new_headers, new_data = split_header_rows(new_rows, header_row=header_row)
    old_headers, old_data = split_header_rows(old_rows, header_row=header_row)
    new_headers = unique_headers(new_headers)
    old_headers = unique_headers(old_headers)
    new_headers, new_data = filter_columns(new_headers, new_data, ignore_columns)
    old_headers, old_data = filter_columns(old_headers, old_data, ignore_columns)

    if new_headers != old_headers:
        return CompareResult(
            "invalid",
            diff_summary={
                "new_headers": new_headers,
                "old_headers": old_headers,
                "missing_in_new": [header for header in old_headers if header not in new_headers],
                "missing_in_old": [header for header in new_headers if header not in old_headers],
            },
            new_row_count=len(new_data),
            old_row_count=len(old_data),
            message="新报表和模板旧数据字段不一致",
        )

    new_hash = table_hash(new_headers, new_data)
    old_hash = table_hash(old_headers, old_data)
    if new_hash == old_hash:
        return CompareResult(
            "same",
            diff_summary={
                "added_rows": 0,
                "removed_rows": 0,
                "changed_rows": 0,
                "content_hash": new_hash,
                "fast_path": "hash_equal",
            },
            new_row_count=len(new_data),
            old_row_count=len(old_data),
            message="数据一致",
        )

    if key_columns:
        diff = compare_with_key(new_headers, new_data, old_data, key_columns, sample_limit)
        if diff.get("invalid_reason"):
            return CompareResult(
                "invalid",
                diff_summary=diff,
                new_row_count=len(new_data),
                old_row_count=len(old_data),
                message=diff["invalid_reason"],
            )
    else:
        diff = compare_without_key(new_data, old_data, sample_limit)

    changed = any(diff.get(name, 0) for name in ("added_rows", "removed_rows", "changed_rows"))
    return CompareResult(
        "changed" if changed else "same",
        diff_summary=diff,
        new_row_count=len(new_data),
        old_row_count=len(old_data),
        message="数据有变化" if changed else "数据一致",
    )


def normalize_sheet_mapping(item):
    if isinstance(item, str):
        return {"name": item, "new_sheet_name": item, "template_sheet_name": item}
    if not isinstance(item, dict):
        raise ValueError(f"sheet_mappings 只支持字符串或对象: {item!r}")
    new_sheet_name = item.get("new_sheet_name") or item.get("new_sheet") or item.get("sheet")
    template_sheet_name = item.get("template_sheet_name") or item.get("template_sheet") or new_sheet_name
    if not new_sheet_name:
        raise ValueError(f"sheet_mappings 缺少 new_sheet_name: {item!r}")
    return {
        "name": item.get("name") or new_sheet_name,
        "new_sheet_name": new_sheet_name,
        "template_sheet_name": template_sheet_name,
        "new_range": item.get("new_range"),
        "template_range": item.get("template_range"),
        "header_row": item.get("header_row"),
        "ignore_columns": item.get("ignore_columns"),
        "key_columns": item.get("key_columns"),
    }


def build_sheet_mappings_from_names(config, new_sheet_names, template_sheet_names):
    if config.get("sheet_mappings"):
        return [normalize_sheet_mapping(item) for item in config["sheet_mappings"]]

    if config.get("new_sheet_name") or config.get("template_sheet_name"):
        new_sheet_name = config.get("new_sheet_name")
        template_sheet_name = config.get("template_sheet_name") or new_sheet_name
        if not new_sheet_name:
            raise ValueError("配置了 template_sheet_name 时，也必须配置 new_sheet_name")
        return [{
            "name": new_sheet_name,
            "new_sheet_name": new_sheet_name,
            "template_sheet_name": template_sheet_name,
            "new_range": config.get("new_range"),
            "template_range": config.get("template_range"),
        }]

    skip_sheets = set(config.get("skip_sheets") or [])
    template_sheet_names = set(template_sheet_names)
    mappings = []
    for sheet_name in new_sheet_names:
        if sheet_name in skip_sheets:
            continue
        mappings.append({
            "name": sheet_name,
            "new_sheet_name": sheet_name,
            "template_sheet_name": sheet_name,
            "missing_template_sheet": sheet_name not in template_sheet_names,
        })
    return mappings


def build_sheet_mappings(config, new_workbook, template_workbook):
    return build_sheet_mappings_from_names(
        config,
        workbook_sheet_names(new_workbook),
        workbook_sheet_names(template_workbook),
    )


def merge_sheet_config(config, mapping):
    header_row = mapping.get("header_row")
    if header_row is None:
        header_row = config.get("header_row", 1)
    return {
        "header_row": int(header_row),
        "ignore_columns": mapping.get("ignore_columns")
        if mapping.get("ignore_columns") is not None
        else config.get("ignore_columns") or [],
        "key_columns": mapping.get("key_columns")
        if mapping.get("key_columns") is not None
        else config.get("key_columns") or [],
        "sample_limit": int(config.get("sample_limit", 10)),
    }


def sheet_has_data_below_header(workbook_path, sheet_name, header_row=1, range_address=None) -> bool:
    header_row = 1 if header_row is None else int(header_row)
    if header_row < 0:
        raise ValueError("header_row 必须大于等于 0")
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        rows = trim_empty_edges(normalize_matrix(worksheet_range_values(workbook[sheet_name], range_address)))
    finally:
        workbook.close()
    data_rows = rows[header_row:] if header_row else rows
    return any(not is_empty_row(row) for row in data_rows)


def find_empty_download_sheet_mappings(download_manifest, compare_sources, base_dir=PROJECT_DIR):
    paths_by_name = {
        item.get("name"): item.get("output_path")
        for item in (download_manifest or {}).get("results") or []
        if item.get("name") and item.get("output_path")
    }
    empty_items = []
    for source_index, source in enumerate(compare_sources or [], start=1):
        download_name = source.get("download_name")
        output_path = paths_by_name.get(download_name)
        if not output_path:
            raise RuntimeError(f"下载结果中找不到 compare_sources 对应报表: {download_name}")
        workbook_path = resolve_path(output_path, base_dir)
        for mapping_index, mapping in enumerate(source.get("sheet_mappings") or [], start=1):
            header_row = mapping.get("header_row")
            if header_row is None:
                header_row = 1
            sheet_name = mapping.get("new_sheet_name")
            if not sheet_name:
                raise ValueError(f"compare_sources[{source_index}].sheet_mappings[{mapping_index}] 缺少 new_sheet_name")
            if not sheet_has_data_below_header(workbook_path, sheet_name, header_row, mapping.get("new_range")):
                empty_items.append({
                    "download_name": download_name,
                    "source_report_path": str(workbook_path),
                    "new_sheet_name": sheet_name,
                    "template_sheet_name": mapping.get("template_sheet_name"),
                    "header_row": int(header_row),
                    "source_index": source_index,
                    "mapping_index": mapping_index,
                    "reason": "download_sheet_empty_below_header",
                })
    return empty_items


def resolve_compare_config(config, base_dir=PROJECT_DIR):
    resolved = dict(config)
    for key in ("new_report_path", "template_path", "output_path"):
        if resolved.get(key):
            resolved[key] = str(resolve_path(resolved[key], base_dir))
    return resolved


def validate_config(config):
    missing = [key for key in ("new_report_path", "template_path") if not config.get(key)]
    if missing:
        raise ValueError(f"缺少必要配置: {missing}")


def normalize_compare_engine(engine):
    normalized = str(engine or "openpyxl").strip().lower()
    if normalized not in COMPARE_ENGINES:
        raise ValueError("compare.engine 只支持 openpyxl 或 com")
    return normalized


def compare_sheet_mapping_openpyxl(config, mapping, template_sheet_names, new_workbook, template_workbook):
    new_sheet_name = mapping["new_sheet_name"]
    template_sheet_name = mapping["template_sheet_name"]
    if template_sheet_name not in template_sheet_names:
        return {
            "name": mapping.get("name") or new_sheet_name,
            "new_sheet_name": new_sheet_name,
            "template_sheet_name": template_sheet_name,
            "result": "invalid",
            "message": f"模板缺少对应工作表: {template_sheet_name}",
        }

    options = merge_sheet_config(config, mapping)
    new_stream = stream_sheet_hash_openpyxl(new_workbook, new_sheet_name, mapping.get("new_range"))
    old_stream = stream_sheet_hash_openpyxl(template_workbook, template_sheet_name, mapping.get("template_range"))
    if (
        new_stream["hash"] == old_stream["hash"]
        and new_stream["non_empty_rows"] >= options["header_row"]
        and old_stream["non_empty_rows"] >= options["header_row"]
    ):
        data_rows = max(new_stream["non_empty_rows"] - options["header_row"], 0)
        payload = asdict(CompareResult(
            "same",
            diff_summary={
                "added_rows": 0,
                "removed_rows": 0,
                "changed_rows": 0,
                "content_hash": new_stream["hash"],
                "fast_path": "stream_hash_equal",
            },
            new_row_count=data_rows,
            old_row_count=data_rows,
            message="数据一致",
        ))
    else:
        payload = asdict(compare_tables(
            read_sheet_table_openpyxl(new_workbook, new_sheet_name, mapping.get("new_range")),
            read_sheet_table_openpyxl(template_workbook, template_sheet_name, mapping.get("template_range")),
            header_row=options["header_row"],
            ignore_columns=options["ignore_columns"],
            key_columns=options["key_columns"],
            sample_limit=options["sample_limit"],
        ))
    payload.update({
        "name": mapping.get("name") or new_sheet_name,
        "new_sheet_name": new_sheet_name,
        "template_sheet_name": template_sheet_name,
    })
    return payload


def compare_sheet_mapping_openpyxl_by_path(config, mapping, template_sheet_names, new_path, template_path):
    keep_vba_new = new_path.suffix.lower() == ".xlsm"
    keep_vba_template = template_path.suffix.lower() == ".xlsm"
    new_workbook = load_workbook(new_path, data_only=True, read_only=True, keep_vba=keep_vba_new)
    template_workbook = load_workbook(template_path, data_only=True, read_only=True, keep_vba=keep_vba_template)
    try:
        return compare_sheet_mapping_openpyxl(config, mapping, template_sheet_names, new_workbook, template_workbook)
    finally:
        new_workbook.close()
        template_workbook.close()


def normalize_max_workers(value, item_count):
    try:
        workers = int(value or 1)
    except (TypeError, ValueError):
        workers = 1
    workers = max(1, min(workers, 16))
    if item_count:
        workers = min(workers, item_count)
    return workers


def compare_sheet_mappings(config, mappings, template_sheet_names, read_new_sheet, read_template_sheet):
    sheet_results = []
    for mapping in mappings:
        new_sheet_name = mapping["new_sheet_name"]
        template_sheet_name = mapping["template_sheet_name"]
        if template_sheet_name not in template_sheet_names:
            sheet_results.append({
                "name": mapping.get("name") or new_sheet_name,
                "new_sheet_name": new_sheet_name,
                "template_sheet_name": template_sheet_name,
                "result": "invalid",
                "message": f"模板缺少对应工作表: {template_sheet_name}",
            })
            continue

        options = merge_sheet_config(config, mapping)
        result = compare_tables(
            read_new_sheet(new_sheet_name, mapping.get("new_range")),
            read_template_sheet(template_sheet_name, mapping.get("template_range")),
            header_row=options["header_row"],
            ignore_columns=options["ignore_columns"],
            key_columns=options["key_columns"],
            sample_limit=options["sample_limit"],
        )
        payload = asdict(result)
        payload.update({
            "name": mapping.get("name") or new_sheet_name,
            "new_sheet_name": new_sheet_name,
            "template_sheet_name": template_sheet_name,
        })
        sheet_results.append(payload)
    return sheet_results


def summarize_workbook_compare(sheet_results):
    summary = {
        "same": sum(1 for item in sheet_results if item.get("result") == "same"),
        "changed": sum(1 for item in sheet_results if item.get("result") == "changed"),
        "invalid": sum(1 for item in sheet_results if item.get("result") == "invalid"),
        "total": len(sheet_results),
    }
    if summary["invalid"]:
        result = "invalid"
        message = "存在无法比对的工作表"
    elif summary["changed"]:
        result = "changed"
        message = "至少一个工作表数据有变化"
    else:
        result = "same"
        message = "所有工作表数据一致"
    return WorkbookCompareResult(result=result, sheets=sheet_results, summary=summary, message=message)


def compare_workbook_com(config):
    config = resolve_compare_config(config)
    validate_config(config)
    new_path = Path(config["new_report_path"]).resolve()
    template_path = Path(config["template_path"]).resolve()
    if not new_path.exists():
        raise FileNotFoundError(f"新报表不存在: {new_path}")
    if not template_path.exists():
        raise FileNotFoundError(f"模板文件不存在: {template_path}")

    excel = open_excel(visible=bool(config.get("visible", False)))
    new_workbook = None
    template_workbook = None
    sheet_results = []
    try:
        new_workbook = excel.Workbooks.Open(
            str(new_path),
            UpdateLinks=0,
            ReadOnly=True,
            IgnoreReadOnlyRecommended=True,
        )
        template_workbook = excel.Workbooks.Open(
            str(template_path),
            UpdateLinks=0,
            ReadOnly=True,
            IgnoreReadOnlyRecommended=True,
        )
        template_sheet_names = set(workbook_sheet_names(template_workbook))
        mappings = build_sheet_mappings(config, new_workbook, template_workbook)
        sheet_results = compare_sheet_mappings(
            config,
            mappings,
            template_sheet_names,
            lambda sheet_name, range_address=None: read_sheet_table_com(new_workbook, sheet_name, range_address),
            lambda sheet_name, range_address=None: read_sheet_table_com(template_workbook, sheet_name, range_address),
        )
    finally:
        if new_workbook is not None:
            new_workbook.Close(SaveChanges=False)
        if template_workbook is not None:
            template_workbook.Close(SaveChanges=False)
        excel.Quit()

    return summarize_workbook_compare(sheet_results)


def compare_workbook_openpyxl(config):
    config = resolve_compare_config(config)
    validate_config(config)
    new_path = Path(config["new_report_path"]).resolve()
    template_path = Path(config["template_path"]).resolve()
    if not new_path.exists():
        raise FileNotFoundError(f"新报表不存在: {new_path}")
    if not template_path.exists():
        raise FileNotFoundError(f"模板文件不存在: {template_path}")

    keep_vba_new = new_path.suffix.lower() == ".xlsm"
    keep_vba_template = template_path.suffix.lower() == ".xlsm"
    new_workbook = None
    template_workbook = None
    new_workbook = load_workbook(new_path, data_only=True, read_only=True, keep_vba=keep_vba_new)
    template_workbook = load_workbook(template_path, data_only=True, read_only=True, keep_vba=keep_vba_template)
    try:
        template_sheet_names = set(template_workbook.sheetnames)
        mappings = build_sheet_mappings_from_names(config, new_workbook.sheetnames, template_workbook.sheetnames)
        max_workers = normalize_max_workers(config.get("max_workers"), len(mappings))
        if max_workers > 1:
            new_workbook.close()
            template_workbook.close()
            new_workbook = None
            template_workbook = None
            with ThreadPoolExecutor(max_workers=max_workers) as executor:
                sheet_results = list(executor.map(
                    lambda mapping: compare_sheet_mapping_openpyxl_by_path(
                        config,
                        mapping,
                        template_sheet_names,
                        new_path,
                        template_path,
                    ),
                    mappings,
                ))
        else:
            sheet_results = [
                compare_sheet_mapping_openpyxl(config, mapping, template_sheet_names, new_workbook, template_workbook)
                for mapping in mappings
            ]
    finally:
        if new_workbook is not None:
            new_workbook.close()
        if template_workbook is not None:
            template_workbook.close()

    result = summarize_workbook_compare(sheet_results)
    result.summary["max_workers"] = max_workers
    return result


def compare_workbook(config):
    engine = normalize_compare_engine(config.get("engine"))
    if engine == "com":
        return compare_workbook_com(config)
    try:
        return compare_workbook_openpyxl(config)
    except (InvalidFileException, BadZipFile, OSError):
        return compare_workbook_com(config)


def result_to_dict(result):
    return asdict(result)


def write_result(output_path, result):
    path = Path(output_path).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = result if isinstance(result, dict) else result_to_dict(result)
    with path.open("w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    return path


def compare_report(config, base_dir=PROJECT_DIR):
    config = resolve_compare_config(config, base_dir)
    started = perf_counter()
    result = compare_workbook(config)
    payload = result_to_dict(result)
    payload["engine"] = normalize_compare_engine(config.get("engine"))
    payload["max_workers"] = normalize_max_workers(config.get("max_workers"), payload.get("summary", {}).get("total", 0))
    payload["timings"] = {
        "total_seconds": round(perf_counter() - started, 3),
    }
    if config.get("output_path"):
        write_result(config["output_path"], payload)
    return payload


def compare_report_from_config(config_path, base_dir=PROJECT_DIR):
    config_path = resolve_path(config_path, base_dir)
    config, resolved = load_json(config_path)
    return compare_report(config, base_dir=resolved.parent)
