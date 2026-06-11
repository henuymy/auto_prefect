"""Import dashboard areas from the hierarchy workbook."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from sqlalchemy import Engine, select
from sqlalchemy.orm import Session, sessionmaker

from models.dashboard_area import Area


AREA_LEVELS = {"CITY", "BRANCH", "GRID", "CHANNEL"}
LEVEL_NO = {
    "CITY": 1,
    "BRANCH": 2,
    "GRID": 3,
    "CHANNEL": 4,
}


@dataclass(frozen=True)
class AreaImportRow:
    area_code: str
    area_name: str
    level_type: str
    parent_code: str | None
    last_seen_at: datetime


def read_hierarchy_workbook(path: str | Path) -> list[dict]:
    workbook = load_workbook(Path(path).resolve(), read_only=True, data_only=True)
    try:
        worksheet = workbook["区域层级"]
        iterator = worksheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(iterator)]
        return [
            dict(zip(headers, values))
            for values in iterator
            if any(value is not None for value in values)
        ]
    finally:
        workbook.close()


def normalize_import_rows(raw_rows: Iterable[dict]) -> list[AreaImportRow]:
    rows = list(raw_rows)
    by_identity = {
        (
            str(row.get("level_type") or "").strip().upper(),
            str(row.get("area_code") or "").strip(),
        ): row
        for row in rows
    }
    manager_parent = {
        str(row.get("area_code") or "").strip(): str(row.get("parent_code") or "").strip()
        for row in rows
        if str(row.get("level_type") or "").strip().upper() == "CHANNEL_MANAGER"
    }

    normalized = []
    seen = set()
    for row in rows:
        level_type = str(row.get("level_type") or "").strip().upper()
        if level_type not in AREA_LEVELS:
            continue
        area_code = str(row.get("area_code") or "").strip()
        area_name = str(row.get("area_name") or "").strip()
        if not area_code or not area_name:
            raise ValueError("区域编码和区域名称不能为空")
        identity = (level_type, area_code)
        if identity in seen:
            raise ValueError(f"Excel 中存在重复区域: {level_type}/{area_code}")
        seen.add(identity)

        parent_code = str(row.get("parent_code") or "").strip() or None
        if level_type == "CHANNEL":
            parent_code = manager_parent.get(parent_code or "")
        if level_type == "CITY":
            parent_code = None
        elif not parent_code:
            raise ValueError(f"区域缺少父级编码: {level_type}/{area_code}")

        collected_at = row.get("collected_at")
        if isinstance(collected_at, datetime):
            last_seen_at = collected_at
        elif collected_at:
            last_seen_at = datetime.fromisoformat(str(collected_at))
        else:
            last_seen_at = datetime.now()

        normalized.append(
            AreaImportRow(
                area_code=area_code,
                area_name=area_name,
                level_type=level_type,
                parent_code=parent_code,
                last_seen_at=last_seen_at,
            )
        )

    for row in normalized:
        if row.level_type == "CITY":
            continue
        parent_level = {
            "BRANCH": "CITY",
            "GRID": "BRANCH",
            "CHANNEL": "GRID",
        }[row.level_type]
        if (parent_level, row.parent_code) not in by_identity:
            raise ValueError(
                f"找不到父级区域: {row.level_type}/{row.area_code} -> "
                f"{parent_level}/{row.parent_code}"
            )
    return normalized


def import_areas(engine: Engine, rows: Iterable[AreaImportRow]) -> dict[str, int]:
    ordered_rows = sorted(rows, key=lambda row: LEVEL_NO[row.level_type])
    session_factory = sessionmaker(bind=engine, expire_on_commit=False, class_=Session)
    created = 0
    updated = 0
    with session_factory.begin() as session:
        existing = {
            (area.level_type, area.area_code): area
            for area in session.scalars(select(Area)).all()
        }
        resolved_ids: dict[tuple[str, str], int] = {}

        for row in ordered_rows:
            parent_id = None
            if row.parent_code:
                parent_level = {
                    "BRANCH": "CITY",
                    "GRID": "BRANCH",
                    "CHANNEL": "GRID",
                }[row.level_type]
                parent_identity = (parent_level, row.parent_code)
                parent_id = resolved_ids.get(parent_identity)
                if parent_id is None:
                    parent = existing.get(parent_identity)
                    parent_id = parent.id if parent else None
                if parent_id is None:
                    raise ValueError(
                        f"数据库导入时找不到父级: {row.level_type}/{row.area_code} "
                        f"-> {parent_level}/{row.parent_code}"
                    )

            identity = (row.level_type, row.area_code)
            area = existing.get(identity)
            if area is None:
                area = Area(
                    area_code=row.area_code,
                    area_name=row.area_name,
                    level_type=row.level_type,
                    level_no=LEVEL_NO[row.level_type],
                    parent_id=parent_id,
                    enabled=True,
                    missing_count=0,
                    last_seen_at=row.last_seen_at,
                )
                session.add(area)
                session.flush()
                existing[identity] = area
                created += 1
            else:
                area.area_name = row.area_name
                area.level_no = LEVEL_NO[row.level_type]
                area.parent_id = parent_id
                area.enabled = True
                area.missing_count = 0
                area.last_seen_at = row.last_seen_at
                area.updated_at = datetime.now()
                updated += 1
            resolved_ids[identity] = area.id

    return {
        "total": len(ordered_rows),
        "created": created,
        "updated": updated,
    }
