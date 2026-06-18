"""Build and import metric target values from an Excel workbook."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from sqlalchemy import Engine, select
from sqlalchemy.orm import Session, sessionmaker

from models.dashboard_area import Area
from models.dashboard_indicator import Indicator
from models.dashboard_metric_target import MetricTarget


VALID_PERIOD_TYPES = {"REALTIME", "DAY_ACC", "MONTH"}


@dataclass(frozen=True)
class MetricTargetImportRow:
    period_type: str
    area_code: str
    indicator_code: str
    target_value: Decimal


def read_metric_target_rows(path: str | Path) -> list[dict]:
    """Read raw rows from the Excel sheet '指标目标值'."""
    workbook = load_workbook(Path(path).resolve(), read_only=True, data_only=True)
    try:
        worksheet = workbook["指标目标值"]
        iterator = worksheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(iterator)]
        return [
            dict(zip(headers, values))
            for values in iterator
            if any(value is not None for value in values)
        ]
    finally:
        workbook.close()


def normalize_metric_targets(
    raw_rows: Iterable[dict],
) -> list[MetricTargetImportRow]:
    """Validate and normalize raw Excel rows into import rows."""
    normalized: list[MetricTargetImportRow] = []
    seen: set[tuple[str, str, str]] = set()
    row_index = 0

    for raw_row in raw_rows:
        row_index += 1
        period_type = str(raw_row.get("period_type") or "").strip().upper()
        area_code = str(raw_row.get("area_code") or "").strip()
        indicator_code = str(raw_row.get("indicator_code") or "").strip()
        raw_value = raw_row.get("target_value")

        if not period_type:
            if not area_code and not indicator_code and raw_value is None:
                continue
            raise ValueError(f"第 {row_index} 行: period_type 不能为空")
        if period_type not in VALID_PERIOD_TYPES:
            # Skip template label/note rows, but reject data rows with typos.
            if (
                period_type == "周期类型"
                or (not area_code and not indicator_code and raw_value is None)
            ):
                continue
            raise ValueError(
                f"第 {row_index} 行: period_type 只支持 "
                f"{', '.join(sorted(VALID_PERIOD_TYPES))}: {period_type!r}"
            )
        if not area_code:
            raise ValueError(f"第 {row_index} 行: area_code 不能为空")
        if not indicator_code:
            raise ValueError(f"第 {row_index} 行: indicator_code 不能为空")
        if raw_value is None:
            raise ValueError(f"第 {row_index} 行: target_value 不能为空")

        try:
            target_value = Decimal(str(raw_value))
        except (InvalidOperation, ValueError) as exc:
            raise ValueError(
                f"第 {row_index} 行: target_value 不是有效数值: {raw_value!r}"
            ) from exc

        identity = (period_type, area_code, indicator_code)
        if identity in seen:
            raise ValueError(
                f"第 {row_index} 行: 重复记录 "
                f"period_type={period_type} area_code={area_code} "
                f"indicator_code={indicator_code}"
            )
        seen.add(identity)

        normalized.append(
            MetricTargetImportRow(
                period_type=period_type,
                area_code=area_code,
                indicator_code=indicator_code,
                target_value=target_value,
            )
        )

    return normalized


def import_metric_targets(
    engine: Engine,
    rows: Iterable[MetricTargetImportRow],
) -> dict[str, int]:
    """Upsert metric targets into the dashboard database.

    For each row, if a matching record exists (same period_type + area_id +
    indicator_id), update its target_value; otherwise insert a new row.
    Rows previously enabled but absent from the import are disabled.
    """
    ordered_rows = sorted(rows, key=lambda r: (r.period_type, r.area_code, r.indicator_code))
    session_factory = sessionmaker(bind=engine, expire_on_commit=False, class_=Session)
    created = 0
    updated = 0
    disabled = 0

    with session_factory.begin() as session:
        # Build lookup maps from reference tables
        areas = {
            area.area_code: area.id
            for area in session.scalars(
                select(Area).where(Area.enabled.is_(True))
            ).all()
        }
        indicators = {
            indicator.code: indicator.id
            for indicator in session.scalars(
                select(Indicator).where(Indicator.enabled.is_(True))
            ).all()
        }

        # Pre-load existing metric_target rows
        existing = {
            (mt.period_type, mt.area_id, mt.indicator_id): mt
            for mt in session.scalars(select(MetricTarget)).all()
        }

        for row in ordered_rows:
            area_id = areas.get(row.area_code)
            if area_id is None:
                raise ValueError(
                    f"找不到启用的区域: area_code={row.area_code}"
                )

            indicator_id = indicators.get(row.indicator_code)
            if indicator_id is None:
                raise ValueError(
                    f"找不到启用的指标: indicator_code={row.indicator_code}"
                )

            identity = (row.period_type, area_id, indicator_id)
            target = existing.get(identity)

            if target is None:
                target = MetricTarget(
                    period_type=row.period_type,
                    area_id=area_id,
                    indicator_id=indicator_id,
                    target_value=row.target_value,
                    enabled=True,
                )
                session.add(target)
                session.flush()
                existing[identity] = target
                created += 1
            else:
                target.target_value = row.target_value
                target.enabled = True
                target.updated_at = datetime.now()
                updated += 1

        # Disable rows present in DB but absent from the import
        imported_identities = {
            (row.period_type, areas[row.area_code], indicators[row.indicator_code])
            for row in ordered_rows
        }
        for identity, target in existing.items():
            if identity in imported_identities or not target.enabled:
                continue
            target.enabled = False
            target.updated_at = datetime.now()
            disabled += 1

    return {
        "total": len(ordered_rows),
        "created": created,
        "updated": updated,
        "disabled": disabled,
    }
