"""Admin helpers for Dashboard V2 target plans and target values."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import Engine, and_, func, or_, select
from sqlalchemy.orm import Session

from models.dashboard_v2 import HierarchyNode, IndicatorV2, MetricTargetValue, TargetPlan
from services.dashboard_metrics import MetricValueError, parse_metric_value
from services.dashboard_v2_target_service import (
    TargetPlanError,
    activate_v2_target_plan_in_session,
    clone_v2_target_plan_in_session,
)


VALID_SCENARIOS = {"NORMAL", "PK"}
VALID_TARGET_PERIODS = {"DAY", "MONTH"}
VALID_NODE_TYPES = {"CITY", "BRANCH", "GRID", "CHANNEL_MANAGER", "CHANNEL"}
SOURCE_TARGET_SHEETS = (
    ("区公司级", "BRANCH"),
    ("网格级", "GRID"),
    ("渠道经理级", "CHANNEL_MANAGER"),
    ("渠道级", "CHANNEL"),
)
SOURCE_TARGET_SHEET_TYPES = {
    **dict(SOURCE_TARGET_SHEETS),
    "市公司级": "CITY",
}


@dataclass(frozen=True)
class TargetExcelRow:
    row_number: int
    scenario: str
    period_type: str
    effective_from: date
    node_type: str
    node_code: str
    indicator_code: str
    target_value: Decimal


def _decimal_to_float(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def _date_to_iso(value: date | None) -> str | None:
    return value.isoformat() if value is not None else None


def _time_to_iso(value: object) -> str | None:
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()  # type: ignore[no-any-return]
    return str(value)


def serialize_target_plan(plan: TargetPlan, value_count: int | None = None) -> dict:
    return {
        "id": plan.id,
        "plan_name": plan.plan_name,
        "scenario": plan.scenario,
        "period_type": plan.period_type,
        "effective_from": _date_to_iso(plan.effective_from),
        "effective_to": _date_to_iso(plan.effective_to),
        "priority": plan.priority,
        "version_no": plan.version_no,
        "status": plan.status,
        "is_realtime": bool(plan.is_realtime),
        "supersedes_plan_id": plan.supersedes_plan_id,
        "activated_at": _time_to_iso(plan.activated_at),
        "retired_at": _time_to_iso(plan.retired_at),
        "created_at": _time_to_iso(plan.created_at),
        "updated_at": _time_to_iso(plan.updated_at),
        "value_count": value_count,
    }


def list_target_plans(engine: Engine, status: str | None = None) -> dict:
    normalized_status = str(status or "").strip().upper()
    if normalized_status and normalized_status not in {"DRAFT", "ACTIVE", "RETIRED"}:
        raise ValueError("status 只支持 DRAFT/ACTIVE/RETIRED")
    with Session(engine) as session:
        query = select(TargetPlan)
        if normalized_status:
            query = query.where(TargetPlan.status == normalized_status)
        plans = session.scalars(
            query.order_by(
                TargetPlan.status,
                TargetPlan.scenario,
                TargetPlan.period_type,
                TargetPlan.effective_from.desc(),
                TargetPlan.version_no.desc(),
            )
        ).all()
        counts = dict(
            session.execute(
                select(MetricTargetValue.plan_id, func.count(MetricTargetValue.id))
                .group_by(MetricTargetValue.plan_id)
            ).all()
        )
        return {
            "plans": [
                serialize_target_plan(plan, int(counts.get(plan.id, 0)))
                for plan in plans
            ]
        }


def create_target_plan(
    engine: Engine,
    *,
    plan_name: str,
    scenario: str,
    period_type: str,
    effective_from: date,
    priority: int = 0,
) -> dict:
    name = plan_name.strip()
    normalized_scenario = scenario.strip().upper()
    normalized_period = period_type.strip().upper()
    if not name:
        raise ValueError("plan_name 不能为空")
    if normalized_scenario not in VALID_SCENARIOS:
        raise ValueError("scenario 只支持 NORMAL/PK")
    if normalized_period not in VALID_TARGET_PERIODS:
        raise ValueError("period_type 只支持 DAY/MONTH")

    with Session(engine) as session, session.begin():
        latest_version = session.scalar(
            select(TargetPlan.version_no)
            .where(
                TargetPlan.plan_name == name,
                TargetPlan.scenario == normalized_scenario,
                TargetPlan.period_type == normalized_period,
                TargetPlan.status == "DRAFT",
            )
            .order_by(TargetPlan.version_no.desc())
            .limit(1)
            .with_for_update()
        )
        plan = TargetPlan(
            plan_name=name,
            scenario=normalized_scenario,
            period_type=normalized_period,
            effective_from=effective_from,
            priority=priority,
            version_no=int(latest_version or 0) + 1,
            status="DRAFT",
        )
        session.add(plan)
        session.flush()
        return serialize_target_plan(plan, 0)


def activate_target_plan(engine: Engine, plan_id: int) -> dict:
    with Session(engine) as session, session.begin():
        plan = activate_v2_target_plan_in_session(
            session,
            plan_id=plan_id,
            activated_at=datetime.now(),
        )
        count = session.scalar(
            select(func.count(MetricTargetValue.id)).where(
                MetricTargetValue.plan_id == plan.id
            )
        )
        return serialize_target_plan(plan, int(count or 0))


def set_target_plan_realtime(engine: Engine, plan_id: int) -> dict:
    """Make one editable target the live target for its scenario and period."""
    with Session(engine) as session, session.begin():
        plan = session.scalar(
            select(TargetPlan).where(TargetPlan.id == plan_id).with_for_update()
        )
        if plan is None:
            raise TargetPlanError(f"目标方案不存在: {plan_id}")
        if plan.status != "DRAFT":
            raise TargetPlanError("只有目标草稿可以设为实时目标")
        drafts = session.scalars(
            select(TargetPlan)
            .where(
                TargetPlan.scenario == plan.scenario,
                TargetPlan.period_type == plan.period_type,
                TargetPlan.status == "DRAFT",
            )
            .with_for_update()
        ).all()
        for draft in drafts:
            draft.is_realtime = draft.id == plan.id
        session.flush()
        count = session.scalar(
            select(func.count(MetricTargetValue.id)).where(
                MetricTargetValue.plan_id == plan.id
            )
        )
        return serialize_target_plan(plan, int(count or 0))


def clone_target_plan(
    engine: Engine,
    *,
    source_plan_id: int,
    effective_from: date | None = None,
) -> dict:
    """Copy an immutable plan and all values into a new editable DRAFT."""
    with Session(engine) as session, session.begin():
        source = session.get(TargetPlan, source_plan_id)
        if source is None:
            raise TargetPlanError(f"目标方案不存在: {source_plan_id}")
        clone = clone_v2_target_plan_in_session(
            session,
            source_plan_id=source_plan_id,
            effective_from=effective_from or source.effective_from,
        )
        count = session.scalar(
            select(func.count(MetricTargetValue.id)).where(
                MetricTargetValue.plan_id == clone.id
            )
        )
        return serialize_target_plan(clone, int(count or 0))


def get_target_values(
    engine: Engine,
    *,
    plan_id: int,
    node_type: str | None = None,
    indicator_code: str | None = None,
    search: str | None = None,
    limit: int = 500,
) -> dict:
    normalized_node_type = str(node_type or "").strip().upper()
    normalized_indicator_code = str(indicator_code or "").strip()
    if normalized_node_type and normalized_node_type not in VALID_NODE_TYPES:
        raise ValueError("node_type 非法")

    with Session(engine) as session:
        plan = session.get(TargetPlan, plan_id)
        if plan is None:
            raise TargetPlanError(f"目标方案不存在: {plan_id}")

        node_query = select(HierarchyNode).where(
            HierarchyNode.enabled.is_(True),
            HierarchyNode.metric_enabled.is_(True),
        )
        if normalized_node_type:
            node_query = node_query.where(HierarchyNode.node_type == normalized_node_type)
        keyword = str(search or "").strip()
        if keyword:
            like = f"%{keyword}%"
            node_query = node_query.where(
                or_(HierarchyNode.node_name.like(like), HierarchyNode.node_code.like(like))
            )
        nodes = session.scalars(
            node_query.order_by(
                HierarchyNode.level_no,
                HierarchyNode.sort_order,
                HierarchyNode.node_code,
            ).limit(limit)
        ).all()

        indicator_query = select(IndicatorV2).where(IndicatorV2.enabled.is_(True))
        if normalized_indicator_code:
            indicator_query = indicator_query.where(
                IndicatorV2.code == normalized_indicator_code
            )
        indicators = session.scalars(
            indicator_query.order_by(IndicatorV2.sort_order, IndicatorV2.code)
        ).all()

        if not nodes or not indicators:
            values = {}
        else:
            value_rows = session.scalars(
                select(MetricTargetValue).where(
                    MetricTargetValue.plan_id == plan_id,
                    MetricTargetValue.node_id.in_([node.id for node in nodes]),
                    MetricTargetValue.indicator_id.in_(
                        [indicator.id for indicator in indicators]
                    ),
                )
            ).all()
            values = {
                (row.node_id, row.indicator_id): row.target_value
                for row in value_rows
            }

        rows = []
        for node in nodes:
            for indicator in indicators:
                value = values.get((node.id, indicator.id))
                rows.append(
                    {
                        "node_id": node.id,
                        "node_type": node.node_type,
                        "node_code": node.node_code,
                        "node_name": node.node_name,
                        "indicator_id": indicator.id,
                        "indicator_code": indicator.code,
                        "indicator_name": indicator.name,
                        "target_value": _decimal_to_float(value),
                    }
                )

        return {
            "plan": serialize_target_plan(plan),
            "rows": rows,
            "row_count": len(rows),
            "node_count": len(nodes),
            "indicator_count": len(indicators),
        }


def save_target_values(
    engine: Engine,
    *,
    plan_id: int,
    values: Iterable[dict[str, object]],
) -> dict:
    normalized_values = list(values)
    if not normalized_values:
        return {"plan_id": plan_id, "saved": 0}

    with Session(engine) as session, session.begin():
        plan = session.scalar(
            select(TargetPlan).where(TargetPlan.id == plan_id).with_for_update()
        )
        if plan is None:
            raise TargetPlanError(f"目标方案不存在: {plan_id}")
        if plan.status != "DRAFT":
            raise TargetPlanError("只有 DRAFT 目标方案允许修改目标值")

        node_ids = {int(item["node_id"]) for item in normalized_values}
        indicator_ids = {int(item["indicator_id"]) for item in normalized_values}
        existing_nodes = set(
            session.scalars(
                select(HierarchyNode.id).where(
                    HierarchyNode.id.in_(node_ids),
                    HierarchyNode.enabled.is_(True),
                    HierarchyNode.metric_enabled.is_(True),
                )
            ).all()
        )
        existing_indicators = set(
            session.scalars(
                select(IndicatorV2.id).where(
                    IndicatorV2.id.in_(indicator_ids),
                    IndicatorV2.enabled.is_(True),
                )
            ).all()
        )
        missing_nodes = sorted(node_ids - existing_nodes)
        missing_indicators = sorted(indicator_ids - existing_indicators)
        if missing_nodes:
            raise TargetPlanError(f"目标值节点不存在或不可用: {missing_nodes[:10]}")
        if missing_indicators:
            raise TargetPlanError(f"目标值指标不存在或不可用: {missing_indicators[:10]}")

        normalized: dict[tuple[int, int], Decimal] = {}
        for item in normalized_values:
            node_id = int(item["node_id"])
            indicator_id = int(item["indicator_id"])
            value = parse_metric_value(item.get("target_value"))
            key = (node_id, indicator_id)
            if key in normalized and normalized[key] != value:
                raise TargetPlanError(
                    f"目标值冲突: node_id={node_id}, indicator_id={indicator_id}"
                )
            normalized[key] = value

        existing = {
            (row.node_id, row.indicator_id): row
            for row in session.scalars(
                select(MetricTargetValue).where(
                    MetricTargetValue.plan_id == plan_id,
                    MetricTargetValue.node_id.in_(node_ids),
                    MetricTargetValue.indicator_id.in_(indicator_ids),
                )
            ).all()
        }
        created = 0
        updated = 0
        for (node_id, indicator_id), value in normalized.items():
            target = existing.get((node_id, indicator_id))
            if target is None:
                session.add(
                    MetricTargetValue(
                        plan_id=plan_id,
                        node_id=node_id,
                        indicator_id=indicator_id,
                        target_value=value,
                    )
                )
                created += 1
            else:
                target.target_value = value
                updated += 1
        return {
            "plan_id": plan_id,
            "saved": len(normalized),
            "created": created,
            "updated": updated,
        }


def build_target_template(engine: Engine) -> bytes:
    with Session(engine) as session:
        nodes = session.scalars(
            select(HierarchyNode)
            .where(HierarchyNode.enabled.is_(True), HierarchyNode.metric_enabled.is_(True))
            .order_by(HierarchyNode.level_no, HierarchyNode.sort_order, HierarchyNode.node_code)
        ).all()
        indicators = session.scalars(
            select(IndicatorV2)
            .where(IndicatorV2.enabled.is_(True))
            .order_by(IndicatorV2.sort_order, IndicatorV2.code)
        ).all()

    workbook = Workbook()
    ws = workbook.active
    ws.title = "目标值"
    headers = [
        ("scenario", "场景", 14),
        ("period_type", "周期", 14),
        ("effective_from", "生效日期", 16),
        ("node_type", "区域层级", 18),
        ("node_code", "区域编码", 24),
        ("indicator_code", "指标编码", 24),
        ("target_value", "目标值", 16),
    ]
    header_fill = PatternFill("solid", fgColor="2563EB")
    label_fill = PatternFill("solid", fgColor="DBEAFE")
    for column, (field, label, width) in enumerate(headers, 1):
        ws.cell(row=1, column=column, value=field)
        ws.cell(row=2, column=column, value=label)
        ws.column_dimensions[get_column_letter(column)].width = width
        for row in (1, 2):
            cell = ws.cell(row=row, column=column)
            cell.font = Font(name="Microsoft YaHei", bold=row == 1, color="FFFFFF" if row == 1 else "1E40AF")
            cell.fill = header_fill if row == 1 else label_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
    examples = [
        ("NORMAL", "DAY", date.today().isoformat(), "BRANCH", nodes[0].node_code if nodes else "AQ", indicators[0].code if indicators else "channel_count", 100),
        ("NORMAL", "MONTH", date.today().isoformat(), "BRANCH", nodes[0].node_code if nodes else "AQ", indicators[0].code if indicators else "channel_count", 3000),
    ]
    for row_index, values in enumerate(examples, 3):
        for column, value in enumerate(values, 1):
            ws.cell(row=row_index, column=column, value=value)
    for cell_range, values in {
        "A3:A10000": "NORMAL,PK",
        "B3:B10000": "DAY,MONTH",
        "D3:D10000": "CITY,BRANCH,GRID,CHANNEL_MANAGER,CHANNEL",
    }.items():
        validation = DataValidation(type="list", formula1=f'"{values}"')
        validation.sqref = cell_range
        ws.add_data_validation(validation)
    ws.freeze_panes = "A3"

    ref_nodes = workbook.create_sheet("区域参照")
    node_headers = ["node_type", "node_code", "node_name", "parent_id", "level_no"]
    for column, header in enumerate(node_headers, 1):
        ref_nodes.cell(row=1, column=column, value=header)
        ref_nodes.column_dimensions[get_column_letter(column)].width = 24
    for row_index, node in enumerate(nodes, 2):
        for column, value in enumerate(
            [node.node_type, node.node_code, node.node_name, node.parent_id, node.level_no],
            1,
        ):
            ref_nodes.cell(row=row_index, column=column, value=value)
    ref_nodes.freeze_panes = "A2"
    ref_nodes.auto_filter.ref = f"A1:E{max(len(nodes) + 1, 2)}"

    ref_indicators = workbook.create_sheet("指标参照")
    indicator_headers = ["indicator_code", "indicator_name", "storage_mode", "indicator_type"]
    for column, header in enumerate(indicator_headers, 1):
        ref_indicators.cell(row=1, column=column, value=header)
        ref_indicators.column_dimensions[get_column_letter(column)].width = 24
    for row_index, indicator in enumerate(indicators, 2):
        for column, value in enumerate(
            [indicator.code, indicator.name, indicator.storage_mode, indicator.indicator_type],
            1,
        ):
            ref_indicators.cell(row=row_index, column=column, value=value)
    ref_indicators.freeze_panes = "A2"
    ref_indicators.auto_filter.ref = f"A1:D{max(len(indicators) + 1, 2)}"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_target_plan_export(engine: Engine, *, plan_id: int) -> bytes:
    """Export a plan in the business workbook layout accepted by the importer."""
    with Session(engine) as session:
        plan = session.get(TargetPlan, plan_id)
        if plan is None:
            raise TargetPlanError(f"目标方案不存在: {plan_id}")
        nodes = session.scalars(
            select(HierarchyNode)
            .where(
                HierarchyNode.enabled.is_(True),
                HierarchyNode.metric_enabled.is_(True),
            )
            .order_by(
                HierarchyNode.level_no,
                HierarchyNode.sort_order,
                HierarchyNode.node_code,
            )
        ).all()
        indicators = session.scalars(
            select(IndicatorV2)
            .where(IndicatorV2.enabled.is_(True))
            .order_by(IndicatorV2.sort_order, IndicatorV2.code)
        ).all()
        target_values = session.scalars(
            select(MetricTargetValue).where(MetricTargetValue.plan_id == plan_id)
        ).all()

    nodes_by_id = {node.id: node for node in nodes}
    indicators_by_id = {indicator.id: indicator for indicator in indicators}
    values_by_key = {
        (target.node_id, target.indicator_id): target.target_value
        for target in target_values
        if target.node_id in nodes_by_id and target.indicator_id in indicators_by_id
    }
    used_indicators_by_type: dict[str, set[int]] = {}
    for node_id, indicator_id in values_by_key:
        node = nodes_by_id[node_id]
        used_indicators_by_type.setdefault(node.node_type, set()).add(indicator_id)

    sheet_types = list(SOURCE_TARGET_SHEETS)
    if used_indicators_by_type.get("CITY"):
        sheet_types.insert(0, ("市公司级", "CITY"))
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")

    for sheet_name, node_type in sheet_types:
        worksheet = workbook.create_sheet(sheet_name)
        type_nodes = [node for node in nodes if node.node_type == node_type]
        type_indicator_ids = used_indicators_by_type.get(node_type, set())
        type_indicators = [
            indicator for indicator in indicators if indicator.id in type_indicator_ids
        ]
        headers = ["编码", "名称", *[indicator.name for indicator in type_indicators]]
        for column, header in enumerate(headers, 1):
            cell = worksheet.cell(row=1, column=column, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            worksheet.column_dimensions[get_column_letter(column)].width = (
                24 if column < 3 else 18
            )
        for row_index, node in enumerate(type_nodes, 2):
            worksheet.cell(row=row_index, column=1, value=node.node_code)
            worksheet.cell(row=row_index, column=2, value=node.node_name)
            for column, indicator in enumerate(type_indicators, 3):
                target_value = values_by_key.get((node.id, indicator.id))
                if target_value is not None:
                    worksheet.cell(
                        row=row_index,
                        column=column,
                        value=_decimal_to_float(target_value),
                    )
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(2, len(type_nodes) + 1)}"

    references = workbook.create_sheet("指标参考表")
    for column, header in enumerate(["编码", "名称"], 1):
        cell = references.cell(row=1, column=column, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        references.column_dimensions[get_column_letter(column)].width = 28
    for row_index, indicator in enumerate(indicators, 2):
        references.cell(row=row_index, column=1, value=indicator.code)
        references.cell(row=row_index, column=2, value=indicator.name)
    references.freeze_panes = "A2"
    references.auto_filter.ref = f"A1:B{max(2, len(indicators) + 1)}"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _read_normalized_target_excel_rows(workbook) -> list[TargetExcelRow]:
    worksheet = workbook["目标值"]
    iterator = worksheet.iter_rows(values_only=True)
    try:
        headers = [str(value or "").strip() for value in next(iterator)]
    except StopIteration as exc:
        raise ValueError("目标值 sheet 没有表头") from exc
    required = {
        "scenario",
        "period_type",
        "effective_from",
        "node_type",
        "node_code",
        "indicator_code",
        "target_value",
    }
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(f"目标值 sheet 缺少字段: {', '.join(missing)}")
    rows: list[TargetExcelRow] = []
    for excel_row_number, values in enumerate(iterator, 2):
        raw = dict(zip(headers, values))
        if not any(value is not None for value in raw.values()):
            continue
        scenario = str(raw.get("scenario") or "").strip().upper()
        period_type = str(raw.get("period_type") or "").strip().upper()
        node_type = str(raw.get("node_type") or "").strip().upper()
        node_code = str(raw.get("node_code") or "").strip()
        indicator_code = str(raw.get("indicator_code") or "").strip()
        if scenario not in VALID_SCENARIOS:
            raise ValueError(f"第 {excel_row_number} 行: scenario 只支持 NORMAL/PK")
        if period_type not in VALID_TARGET_PERIODS:
            raise ValueError(f"第 {excel_row_number} 行: period_type 只支持 DAY/MONTH")
        if node_type not in VALID_NODE_TYPES:
            raise ValueError(f"第 {excel_row_number} 行: node_type 非法")
        if not node_code:
            raise ValueError(f"第 {excel_row_number} 行: node_code 不能为空")
        if not indicator_code:
            raise ValueError(f"第 {excel_row_number} 行: indicator_code 不能为空")
        raw_effective_from = raw.get("effective_from")
        if isinstance(raw_effective_from, datetime):
            effective_from = raw_effective_from.date()
        elif isinstance(raw_effective_from, date):
            effective_from = raw_effective_from
        else:
            try:
                effective_from = date.fromisoformat(str(raw_effective_from))
            except ValueError as exc:
                raise ValueError(
                    f"第 {excel_row_number} 行: effective_from 不是有效日期"
                ) from exc
        try:
            target_value = parse_metric_value(raw.get("target_value"))
        except (ValueError, MetricValueError) as exc:
            raise ValueError(
                f"第 {excel_row_number} 行: target_value 不是有效数值"
            ) from exc
        rows.append(
            TargetExcelRow(
                row_number=excel_row_number,
                scenario=scenario,
                period_type=period_type,
                effective_from=effective_from,
                node_type=node_type,
                node_code=node_code,
                indicator_code=indicator_code,
                target_value=target_value,
            )
        )
    return rows


def _first_nonempty_row(worksheet) -> tuple[int, tuple[object, ...]] | None:
    for row_number, values in enumerate(worksheet.iter_rows(values_only=True), 1):
        if any(value is not None and str(value).strip() for value in values):
            return row_number, values
    return None


def _source_indicator_codes(workbook) -> dict[str, str]:
    if "指标参考表" not in workbook.sheetnames:
        raise ValueError("分层目标文件缺少“指标参考表”sheet")
    worksheet = workbook["指标参考表"]
    header = _first_nonempty_row(worksheet)
    if header is None:
        raise ValueError("指标参考表为空")
    header_row, header_values = header
    if len(header_values) < 2 or [str(value or "").strip() for value in header_values[:2]] != ["编码", "名称"]:
        raise ValueError("指标参考表前两列必须为“编码”“名称”")

    codes: dict[str, str] = {}
    for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        code = str(values[0] or "").strip() if values else ""
        name = str(values[1] or "").strip() if len(values) > 1 else ""
        if not code and not name:
            continue
        if not code or not name:
            raise ValueError("指标参考表存在缺少编码或名称的行")
        existing = codes.get(name)
        if existing and existing != code:
            raise ValueError(f"指标参考表名称映射重复: {name}")
        codes[name] = code
        codes[code] = code
    return codes


def _read_source_target_excel_rows(
    workbook,
    *,
    scenario: str,
    period_type: str,
    effective_from: date,
) -> list[TargetExcelRow]:
    indicator_codes = _source_indicator_codes(workbook)
    rows: list[TargetExcelRow] = []
    recognized_sheets = 0
    for sheet_name, node_type in SOURCE_TARGET_SHEET_TYPES.items():
        if sheet_name not in workbook.sheetnames:
            continue
        recognized_sheets += 1
        worksheet = workbook[sheet_name]
        header = _first_nonempty_row(worksheet)
        if header is None:
            continue
        header_row, header_values = header
        header_cells = [str(value or "").strip() for value in header_values]
        if len(header_cells) < 2 or header_cells[:2] != ["编码", "名称"]:
            raise ValueError(f"{sheet_name} 前两列必须为“编码”“名称”")
        target_columns: list[tuple[int, str, str]] = []
        for index, header_name in enumerate(header_cells[2:], 3):
            if not header_name:
                continue
            indicator_code = indicator_codes.get(header_name)
            if not indicator_code:
                raise ValueError(
                    f"{sheet_name} 指标列“{header_name}”未在指标参考表中定义"
                )
            target_columns.append((index, header_name, indicator_code))
        for excel_row_number, values in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            header_row + 1,
        ):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            node_code = str(values[0] or "").strip() if values else ""
            node_name = str(values[1] or "").strip() if len(values) > 1 else ""
            if not node_code:
                raise ValueError(f"{sheet_name} 第 {excel_row_number} 行: 编码不能为空")
            if not node_name:
                raise ValueError(f"{sheet_name} 第 {excel_row_number} 行: 名称不能为空")
            for column, header_name, indicator_code in target_columns:
                raw_value = values[column - 1] if len(values) >= column else None
                if raw_value is None or not str(raw_value).strip():
                    continue
                try:
                    target_value = parse_metric_value(raw_value)
                except (ValueError, MetricValueError) as exc:
                    raise ValueError(
                        f"{sheet_name} 第 {excel_row_number} 行“{header_name}”不是有效数值"
                    ) from exc
                rows.append(
                    TargetExcelRow(
                        row_number=excel_row_number,
                        scenario=scenario,
                        period_type=period_type,
                        effective_from=effective_from,
                        node_type=node_type,
                        node_code=node_code,
                        indicator_code=indicator_code,
                        target_value=target_value,
                    )
                )
    if not recognized_sheets:
        raise ValueError("Excel 缺少“目标值”或区公司级/网格级等目标数据 sheet")
    return rows


def read_target_excel_rows(
    content: bytes,
    *,
    scenario: str = "NORMAL",
    period_type: str = "DAY",
    effective_from: date | None = None,
) -> list[TargetExcelRow]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        if "目标值" in workbook.sheetnames:
            return _read_normalized_target_excel_rows(workbook)
        normalized_scenario = str(scenario or "").strip().upper()
        normalized_period = str(period_type or "").strip().upper()
        if normalized_scenario not in VALID_SCENARIOS:
            raise ValueError("scenario 只支持 NORMAL/PK")
        if normalized_period not in VALID_TARGET_PERIODS:
            raise ValueError("period_type 只支持 DAY/MONTH")
        return _read_source_target_excel_rows(
            workbook,
            scenario=normalized_scenario,
            period_type=normalized_period,
            effective_from=effective_from or date.today(),
        )
    finally:
        workbook.close()


def import_target_template(engine: Engine, *, plan_id: int, content: bytes) -> dict:
    with Session(engine) as session:
        plan = session.get(TargetPlan, plan_id)
        if plan is None:
            raise TargetPlanError(f"目标方案不存在: {plan_id}")
        if plan.status != "DRAFT":
            raise TargetPlanError("只有 DRAFT 目标方案允许导入目标值")
        rows = read_target_excel_rows(
            content,
            scenario=plan.scenario,
            period_type=plan.period_type,
            effective_from=plan.effective_from,
        )
    if not rows:
        raise ValueError("目标值文件没有可导入的数据行")
    identities = {
        (row.node_type, row.node_code, row.indicator_code)
        for row in rows
    }
    if len(identities) != len(rows):
        raise ValueError("目标值文件存在重复的 node_type/node_code/indicator_code")

    with Session(engine) as session, session.begin():
        plan = session.scalar(
            select(TargetPlan).where(TargetPlan.id == plan_id).with_for_update()
        )
        if plan is None:
            raise TargetPlanError(f"目标方案不存在: {plan_id}")
        if plan.status != "DRAFT":
            raise TargetPlanError("只有 DRAFT 目标方案允许导入目标值")
        mismatched = [
            row.row_number
            for row in rows
            if row.scenario != plan.scenario
            or row.period_type != plan.period_type
            or row.effective_from != plan.effective_from
        ]
        if mismatched:
            raise ValueError(
                f"导入行与目标方案场景/周期/生效日期不一致: {mismatched[:10]}"
            )
        nodes = session.scalars(
            select(HierarchyNode).where(
                HierarchyNode.enabled.is_(True),
                HierarchyNode.metric_enabled.is_(True),
                and_(
                    HierarchyNode.node_type.in_({row.node_type for row in rows}),
                    HierarchyNode.node_code.in_({row.node_code for row in rows}),
                ),
            )
        ).all()
        node_by_key = {(node.node_type, node.node_code): node.id for node in nodes}
        indicators = session.scalars(
            select(IndicatorV2).where(
                IndicatorV2.enabled.is_(True),
                IndicatorV2.code.in_({row.indicator_code for row in rows}),
            )
        ).all()
        indicator_by_code = {indicator.code: indicator.id for indicator in indicators}
        missing_nodes = sorted(
            {(row.node_type, row.node_code) for row in rows} - set(node_by_key)
        )
        missing_indicators = sorted(
            {row.indicator_code for row in rows} - set(indicator_by_code)
        )
        if missing_nodes:
            raise ValueError(f"目标值节点不存在或不可用: {missing_nodes[:10]}")
        if missing_indicators:
            raise ValueError(f"目标值指标不存在或不可用: {missing_indicators[:10]}")

        session.query(MetricTargetValue).filter(
            MetricTargetValue.plan_id == plan_id
        ).delete(synchronize_session=False)
        session.add_all(
            MetricTargetValue(
                plan_id=plan_id,
                node_id=node_by_key[(row.node_type, row.node_code)],
                indicator_id=indicator_by_code[row.indicator_code],
                target_value=row.target_value,
            )
            for row in rows
        )
        return {
            "plan": serialize_target_plan(plan, len(rows)),
            "imported": len(rows),
            "created": len(rows),
            "updated": 0,
        }
