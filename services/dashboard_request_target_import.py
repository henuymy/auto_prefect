"""Build and import platform request targets from the hierarchy workbook."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from sqlalchemy import Engine, select
from sqlalchemy.orm import Session, sessionmaker

from models.dashboard_area import Area
from models.dashboard_request_target import RequestTarget


TARGET_TYPES = {"CITY", "BRANCH", "GRID", "CHANNEL_MANAGER"}
TARGET_ORDER = {
    "CITY": 1,
    "BRANCH": 2,
    "GRID": 3,
    "CHANNEL_MANAGER": 4,
}


@dataclass(frozen=True)
class RequestTargetImportRow:
    target_code: str
    target_name: str
    target_type: str
    parent_code: str | None
    area_identity: tuple[str, str] | None
    sort_order: int


def read_hierarchy_rows(path: str | Path) -> list[dict]:
    workbook = load_workbook(Path(path).resolve(), read_only=True, data_only=True)
    try:
        worksheet = workbook["区域层级"]
        iterator = worksheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(iterator)]
        return [
            dict(zip(headers, values))
            for values in iterator
            if any(value is not None for value in values)
        ]
    finally:
        workbook.close()


def normalize_request_targets(
    raw_rows: Iterable[dict],
) -> list[RequestTargetImportRow]:
    normalized = []
    seen = set()
    per_type_order = {target_type: 0 for target_type in TARGET_TYPES}
    for raw_row in raw_rows:
        target_type = str(raw_row.get("level_type") or "").strip().upper()
        if target_type not in TARGET_TYPES:
            continue
        target_code = str(raw_row.get("area_code") or "").strip()
        target_name = str(raw_row.get("area_name") or "").strip()
        parent_code = str(raw_row.get("parent_code") or "").strip() or None
        if not target_code or not target_name:
            raise ValueError(f"请求目标编码和名称不能为空: {raw_row}")
        if target_type == "CITY":
            parent_code = None
        elif not parent_code:
            raise ValueError(f"请求目标缺少父级编码: {target_type}/{target_code}")

        identity = (target_type, target_code)
        if identity in seen:
            raise ValueError(f"请求目标重复: {target_type}/{target_code}")
        seen.add(identity)
        per_type_order[target_type] += 10
        normalized.append(
            RequestTargetImportRow(
                target_code=target_code,
                target_name=target_name,
                target_type=target_type,
                parent_code=parent_code,
                area_identity=(
                    (target_type, target_code)
                    if target_type != "CHANNEL_MANAGER"
                    else None
                ),
                sort_order=per_type_order[target_type],
            )
        )

    identities = {(row.target_type, row.target_code) for row in normalized}
    parent_types = {
        "BRANCH": "CITY",
        "GRID": "BRANCH",
        "CHANNEL_MANAGER": "GRID",
    }
    for row in normalized:
        if not row.parent_code:
            continue
        parent_identity = (parent_types[row.target_type], row.parent_code)
        if parent_identity not in identities:
            raise ValueError(
                f"找不到请求路径父级: {row.target_type}/{row.target_code} "
                f"-> {parent_identity[0]}/{parent_identity[1]}"
            )
    return normalized


def import_request_targets(
    engine: Engine,
    rows: Iterable[RequestTargetImportRow],
) -> dict[str, int]:
    ordered_rows = sorted(rows, key=lambda row: TARGET_ORDER[row.target_type])
    session_factory = sessionmaker(bind=engine, expire_on_commit=False, class_=Session)
    created = 0
    updated = 0
    disabled = 0
    with session_factory.begin() as session:
        areas = {
            (area.level_type, area.area_code): area.id
            for area in session.scalars(select(Area)).all()
        }
        existing = {
            (target.target_type, target.target_code): target
            for target in session.scalars(select(RequestTarget)).all()
        }
        resolved_ids: dict[tuple[str, str], int] = {}
        parent_types = {
            "BRANCH": "CITY",
            "GRID": "BRANCH",
            "CHANNEL_MANAGER": "GRID",
        }

        for row in ordered_rows:
            area_id = None
            if row.area_identity:
                area_id = areas.get(row.area_identity)
                if area_id is None:
                    raise ValueError(
                        f"请求目标找不到 area: "
                        f"{row.area_identity[0]}/{row.area_identity[1]}"
                    )

            parent_target_id = None
            if row.parent_code:
                parent_identity = (parent_types[row.target_type], row.parent_code)
                parent_target_id = resolved_ids.get(parent_identity)
                if parent_target_id is None:
                    parent = existing.get(parent_identity)
                    parent_target_id = parent.id if parent else None
                if parent_target_id is None:
                    raise ValueError(
                        f"数据库导入时找不到请求父级: "
                        f"{row.target_type}/{row.target_code}"
                    )

            identity = (row.target_type, row.target_code)
            target = existing.get(identity)
            if target is None:
                target = RequestTarget(
                    target_code=row.target_code,
                    target_name=row.target_name,
                    target_type=row.target_type,
                    area_id=area_id,
                    parent_target_id=parent_target_id,
                    enabled=True,
                    sort_order=row.sort_order,
                )
                session.add(target)
                session.flush()
                existing[identity] = target
                created += 1
            else:
                target.target_name = row.target_name
                target.area_id = area_id
                target.parent_target_id = parent_target_id
                target.enabled = True
                target.sort_order = row.sort_order
                target.updated_at = datetime.now()
                updated += 1
            resolved_ids[identity] = target.id

        imported_identities = {
            (row.target_type, row.target_code)
            for row in ordered_rows
        }
        for identity, target in existing.items():
            if identity in imported_identities or not target.enabled:
                continue
            target.enabled = False
            target.updated_at = datetime.now()
            disabled += 1

    return {
        "total": len(ordered_rows),
        "created": created,
        "updated": updated,
        "disabled": disabled,
    }
