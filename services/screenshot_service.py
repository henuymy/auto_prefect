"""Build image/text message packages from Excel workbooks."""

from __future__ import annotations

import base64
import hashlib
import json
import shutil
import tempfile
import time
import uuid
from datetime import datetime
from pathlib import Path


from infrastructure.excel_client import get_sheet, open_excel, open_workbook
from services.runtime_paths import is_runtime_relative_path, resolve_runtime_relative_path

PROJECT_DIR = Path(__file__).resolve().parents[1]
XL_TYPE_PDF = 0
XL_QUALITY_STANDARD = 0


def resolve_path(value, base_dir=PROJECT_DIR):
    path = Path(value)
    if is_runtime_relative_path(path):
        return resolve_runtime_relative_path(path)
    if not path.is_absolute():
        path = base_dir / path
    return path


def get_output_paths(config, base_dir=PROJECT_DIR):
    output_config = config.get("output", {})
    runtime_dir = resolve_path(
        output_config.get("runtime_dir", "modules/wecom_sender/output"),
        base_dir,
    )
    image_dir = resolve_path(output_config.get("image_dir", "images"), runtime_dir)
    package_file = resolve_path(output_config.get("package_file", "message_package.json"), runtime_dir)
    preview_file = resolve_path(output_config.get("preview_image_file", "preview.png"), runtime_dir)
    return image_dir, package_file, preview_file


def get_intermediate_dir(config, base_dir=PROJECT_DIR):
    output_config = config.get("output", {})
    runtime_dir = resolve_path(
        output_config.get("runtime_dir", "modules/wecom_sender/output"),
        base_dir,
    )
    return resolve_path(output_config.get("intermediate_dir", "intermediates"), runtime_dir)


def merge_capture_config(defaults, item_config):
    merged = dict(defaults or {})
    merged.update(item_config or {})
    return merged


def with_intermediate_capture(config, capture_config, image_name, base_dir=PROJECT_DIR):
    output_config = config.get("output", {})
    if not output_config.get("keep_intermediate_files", False):
        return capture_config
    capture_config = dict(capture_config or {})
    capture_config["_keep_intermediate_files"] = True
    capture_config["_intermediate_dir"] = str(get_intermediate_dir(config, base_dir))
    capture_config["_intermediate_stem"] = Path(image_name).stem
    return capture_config


def intermediate_path(capture, stage, suffix):
    if not capture.get("_keep_intermediate_files"):
        return None
    directory = Path(capture["_intermediate_dir"])
    directory.mkdir(parents=True, exist_ok=True)
    return directory / f"{capture['_intermediate_stem']}_{stage}{suffix}"


def copy_intermediate_file(source_path, capture, stage, suffix=None):
    target = intermediate_path(capture, stage, suffix or Path(source_path).suffix)
    if not target:
        return None
    shutil.copyfile(source_path, target)
    return str(target.resolve())


def normalize_2d(values):
    if values is None:
        return [[None]]
    if not isinstance(values, tuple):
        return [[values]]
    if values and not isinstance(values[0], tuple):
        return [list(values)]
    return [list(row) for row in values]


def has_value(value):
    return value is not None and str(value).strip() != ""


def used_range_bounds(ws, shrink_empty_edges=True):
    used = ws.UsedRange
    first_row = used.Row
    first_col = used.Column
    last_row = first_row + used.Rows.Count - 1
    last_col = first_col + used.Columns.Count - 1

    if not shrink_empty_edges:
        return first_row, first_col, last_row, last_col

    values = normalize_2d(ws.Range(ws.Cells(first_row, first_col), ws.Cells(last_row, last_col)).Value)
    populated = []
    for row_index, row in enumerate(values):
        for col_index, value in enumerate(row):
            if has_value(value):
                populated.append((first_row + row_index, first_col + col_index))

    if not populated:
        raise ValueError(f"工作表 {ws.Name} 没有可截图的非空单元格")

    rows = [item[0] for item in populated]
    cols = [item[1] for item in populated]
    bounds = [min(rows), min(cols), max(rows), max(cols)]

    # A merged title/group cell only stores its value in the top-left cell.  If
    # empty edges are shrunk solely from values, the remaining columns/rows of
    # that merge are silently cut off. Expand every merge that intersects the
    # current content bounds; repeat because an expanded merge can touch another.
    areas = []
    seen_areas = set()
    for row_index, col_index in populated:
        try:
            cell = ws.Cells(row_index, col_index)
            if not cell.MergeCells:
                continue
            area = cell.MergeArea
            identity = (
                int(area.Row),
                int(area.Column),
                int(area.Rows.Count),
                int(area.Columns.Count),
            )
            if identity not in seen_areas:
                seen_areas.add(identity)
                areas.append(area)
        except Exception:
            # Some Excel-compatible COM implementations do not expose merge
            # metadata. Value-based shrinking remains a safe fallback there.
            continue

    changed = True
    while changed:
        changed = False
        for area in areas:
            area_first_row = int(area.Row)
            area_first_col = int(area.Column)
            area_last_row = area_first_row + int(area.Rows.Count) - 1
            area_last_col = area_first_col + int(area.Columns.Count) - 1
            intersects = not (
                area_last_row < bounds[0]
                or area_first_row > bounds[2]
                or area_last_col < bounds[1]
                or area_first_col > bounds[3]
            )
            if not intersects:
                continue
            expanded = [
                min(bounds[0], area_first_row),
                min(bounds[1], area_first_col),
                max(bounds[2], area_last_row),
                max(bounds[3], area_last_col),
            ]
            if expanded != bounds:
                bounds = expanded
                changed = True

    return tuple(bounds)


def range_from_capture(ws, capture):
    capture = capture or {}
    mode = capture.get("mode", "used_range")

    if mode == "explicit_range":
        address = capture.get("range")
        if not address:
            raise ValueError("capture.mode=explicit_range 时必须提供 capture.range")
        return ws.Range(address)

    if mode == "current_region":
        return ws.Range(capture.get("start_cell", "A1")).CurrentRegion

    if mode != "used_range":
        raise ValueError(f"不支持的截图模式: {mode}")

    first_row, first_col, last_row, last_col = used_range_bounds(
        ws,
        shrink_empty_edges=capture.get("shrink_empty_edges", True),
    )
    padding_rows = int(capture.get("padding_rows", 0) or 0)
    padding_cols = int(capture.get("padding_cols", 0) or 0)
    return ws.Range(
        ws.Cells(max(1, first_row - padding_rows), max(1, first_col - padding_cols)),
        ws.Cells(last_row + padding_rows, last_col + padding_cols),
    )


def refresh_workbook(excel, workbook, timeout_seconds=120):
    workbook.RefreshAll()
    try:
        excel.CalculateUntilAsyncQueriesDone()
    except Exception:
        pass

    deadline = time.time() + timeout_seconds
    while time.time() < deadline:
        refreshing = False
        try:
            for index in range(1, workbook.Connections.Count + 1):
                connection = workbook.Connections(index)
                try:
                    if connection.OLEDBConnection.Refreshing:
                        refreshing = True
                except Exception:
                    pass
                try:
                    if connection.ODBCConnection.Refreshing:
                        refreshing = True
                except Exception:
                    pass
        except Exception:
            break
        if not refreshing:
            return
        time.sleep(1)

    raise TimeoutError(f"刷新查询和连接超时: {timeout_seconds} 秒")


def prepare_excel_for_capture(ws, capture):
    """Normalize and wait for Excel state before touching printer pagination."""
    excel = ws.Application
    timeout_seconds = max(5, int(capture.get("excel_ready_timeout_seconds", 60) or 60))
    poll_seconds = max(0.05, float(capture.get("excel_ready_poll_seconds", 0.2) or 0.2))

    active_printer = str(
        capture.get("excel_active_printer", "Microsoft Print to PDF") or ""
    ).strip()
    if active_printer:
        try:
            current_printer = str(excel.ActivePrinter)
            expected_name = active_printer.split(" on ", 1)[0]
            if expected_name not in current_printer:
                raise RuntimeError(f"当前 Excel 打印机为 {current_printer!r}")
        except Exception as exc:
            raise RuntimeError(
                f"无法为 Excel 设置稳定打印机 {active_printer!r}；拒绝使用系统默认打印机进行分页"
            ) from exc

    # These properties can be left altered after a failed COM call. Set them
    # explicitly for every image rather than trusting application-global state.
    for name, value in (
        ("DisplayAlerts", False),
        ("ScreenUpdating", False),
        ("EnableEvents", False),
        ("CutCopyMode", False),
        ("PrintCommunication", True),
    ):
        try:
            setattr(excel, name, value)
        except Exception:
            pass

    try:
        ws.Parent.Activate()
        ws.Activate()
    except Exception:
        pass

    try:
        excel.CalculateUntilAsyncQueriesDone()
    except Exception:
        pass

    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        try:
            ready = bool(excel.Ready)
        except Exception:
            ready = True
        try:
            calculation_done = int(excel.CalculationState) == 0
        except Exception:
            calculation_done = True
        if ready and calculation_done:
            return
        try:
            import pythoncom  # type: ignore

            pythoncom.PumpWaitingMessages()
        except Exception:
            pass
        time.sleep(poll_seconds)

    raise TimeoutError(f"Excel 截图前等待就绪超时: sheet={ws.Name}, timeout={timeout_seconds}s")


def wait_for_pdf_export(excel, pdf_path, timeout_seconds=30, poll_seconds=0.2):
    """Wait until the asynchronous Excel/PDF driver has finished writing."""
    path = Path(pdf_path)
    deadline = time.monotonic() + max(5, timeout_seconds)
    previous_size = -1
    stable_checks = 0
    while time.monotonic() < deadline:
        try:
            size = path.stat().st_size
        except (FileNotFoundError, OSError):
            size = 0
        try:
            ready = bool(excel.Ready)
        except Exception:
            ready = True
        if size > 0 and size == previous_size and ready:
            stable_checks += 1
            if stable_checks >= 3:
                return
        else:
            stable_checks = 0
        previous_size = size
        time.sleep(max(0.05, poll_seconds))
    raise TimeoutError(f"等待 Excel PDF 导出完成超时: {path}, timeout={timeout_seconds}s")


def export_range_to_pdf(ws, rng, pdf_path, capture):
    pdf_path = Path(pdf_path)
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    page_setup = ws.PageSetup
    restore_values = {}
    page_setup_keys = [
        "PrintArea",
        "Zoom",
        "FitToPagesWide",
        "FitToPagesTall",
        "LeftMargin",
        "RightMargin",
        "TopMargin",
        "BottomMargin",
        "HeaderMargin",
        "FooterMargin",
        "CenterHorizontally",
        "CenterVertically",
        "Orientation",
    ]
    for key in page_setup_keys:
        try:
            restore_values[key] = getattr(page_setup, key)
        except Exception:
            pass

    excel = ws.Application
    try:
        # Manual page breaks survive ordinary FitToPages assignments and can
        # force Excel to emit several PDF pages. The workbook is opened read-only
        # and closed without saving, so resetting them is safe for the template.
        try:
            ws.Activate()
            ws.ResetAllPageBreaks()
        except Exception:
            pass
        # Excel may defer PageSetup changes while talking to the printer driver.
        # Batch the changes and explicitly flush them before exporting so stale
        # pagination from a previous/default printer is not used.
        try:
            excel.PrintCommunication = False
        except Exception:
            pass
        page_setup.PrintArea = rng.Address
        page_setup.Zoom = False
        page_setup.FitToPagesWide = 1
        page_setup.FitToPagesTall = 1
        margin_points = float(capture.get("pdf_margin_points", 0) or 0)
        page_setup.LeftMargin = margin_points
        page_setup.RightMargin = margin_points
        page_setup.TopMargin = margin_points
        page_setup.BottomMargin = margin_points
        page_setup.HeaderMargin = 0
        page_setup.FooterMargin = 0
        page_setup.CenterHorizontally = False
        page_setup.CenterVertically = False
        orientation = str(capture.get("pdf_orientation", "auto") or "auto").lower()
        if orientation == "auto":
            page_setup.Orientation = 2 if float(rng.Width) >= float(rng.Height) else 1
        elif orientation == "portrait":
            page_setup.Orientation = 1
        elif orientation == "landscape":
            page_setup.Orientation = 2
        else:
            raise ValueError("pdf_orientation 仅支持 auto/portrait/landscape")
        try:
            excel.PrintCommunication = True
        except Exception:
            pass
        # Reading the values back forces Excel to finish applying the queued
        # printer settings. Reapply once when a driver ignored the first batch.
        try:
            fit_applied = (
                page_setup.Zoom is False
                and int(page_setup.FitToPagesWide) == 1
                and int(page_setup.FitToPagesTall) == 1
            )
        except Exception:
            fit_applied = True
        if not fit_applied:
            page_setup.Zoom = False
            page_setup.FitToPagesWide = 1
            page_setup.FitToPagesTall = 1
        try:
            excel.CalculateFull()
        except Exception:
            pass
        export_kwargs = {
            "Type": XL_TYPE_PDF,
            "Filename": str(pdf_path.resolve()),
            "Quality": XL_QUALITY_STANDARD,
            "IncludeDocProperties": False,
            "IgnorePrintAreas": False,
            "OpenAfterPublish": False,
        }
        # Some Excel/pywin32 installations do not expose the worksheet
        # dispatch member even though the workbook export API is available.
        # Export the active sheet through its parent workbook in that case.
        try:
            ws.ExportAsFixedFormat(**export_kwargs)
        except AttributeError as exc:
            if "ExportAsFixedFormat" not in str(exc):
                raise
            workbook = ws.Parent
            workbook.ExportAsFixedFormat(**export_kwargs)
        # ExportAsFixedFormat may return before the printer driver finishes
        # consuming PageSetup. Do not restore it until the PDF is stable.
        wait_for_pdf_export(
            excel,
            pdf_path,
            timeout_seconds=int(capture.get("pdf_export_timeout_seconds", 30) or 30),
            poll_seconds=float(capture.get("pdf_export_poll_seconds", 0.2) or 0.2),
        )
    finally:
        try:
            excel.PrintCommunication = False
        except Exception:
            pass
        for key, value in restore_values.items():
            try:
                setattr(page_setup, key, value)
            except Exception:
                pass
        try:
            excel.PrintCommunication = True
        except Exception:
            pass


def crop_png_whitespace(image_path, background=(255, 255, 255), tolerance=8):
    try:
        from PIL import Image, ImageChops
    except ImportError:
        return None

    path = Path(image_path)
    with Image.open(path) as image:
        source = image.convert("RGB")
        bg = Image.new("RGB", source.size, background)
        diff = ImageChops.difference(source, bg)
        if tolerance > 0:
            diff = diff.point(lambda value: 0 if value <= tolerance else 255)
        bbox = diff.getbbox()
        if not bbox:
            return None
        padding = 2
        left = max(0, bbox[0] - padding)
        top = max(0, bbox[1] - padding)
        right = min(source.width, bbox[2] + padding)
        bottom = min(source.height, bbox[3] + padding)
        if (left, top, right, bottom) == (0, 0, source.width, source.height):
            return source.size
        cropped = source.crop((left, top, right, bottom))
        cropped.save(path, format="PNG")
        return cropped.size


def resolve_pdf_dpi(capture):
    configured = capture.get("pdf_dpi", 300)
    return max(96, min(600, int(configured or 300)))


def render_pdf_to_png(pdf_path, output_path, capture):
    try:
        import fitz  # type: ignore
    except ImportError as exc:
        raise RuntimeError("高清 PDF 截图需要 PyMuPDF，请先安装: pip install PyMuPDF") from exc

    dpi = resolve_pdf_dpi(capture)
    matrix = fitz.Matrix(dpi / 72, dpi / 72)
    with fitz.open(str(pdf_path)) as document:
        if document.page_count < 1:
            raise RuntimeError(f"Excel 导出的 PDF 没有页面: {pdf_path}")
        if document.page_count != 1:
            raise RuntimeError(
                f"Excel PDF 导出产生 {document.page_count} 页；发送图片只允许单页"
            )
        pixmap = document.load_page(0).get_pixmap(matrix=matrix, alpha=False)
        pixmap.save(str(Path(output_path).resolve()))
        width, height = pixmap.width, pixmap.height
        copy_intermediate_file(output_path, capture, "02_pdf_render_raw", ".png")

    if capture.get("pdf_crop_whitespace", True):
        cropped_size = crop_png_whitespace(
            output_path,
            tolerance=int(capture.get("pdf_crop_tolerance", 8) or 8),
        )
        if cropped_size:
            width, height = cropped_size
        copy_intermediate_file(output_path, capture, "03_after_crop", ".png")
    return dpi, width, height


def capture_range_to_png_pdf(ws, output_path, capture=None):
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    capture = capture or {}
    prepare_excel_for_capture(ws, capture)
    rng = range_from_capture(ws, capture)

    temp_dir = Path(tempfile.gettempdir()) / "auto_notify_excel_exports"
    temp_dir.mkdir(parents=True, exist_ok=True)
    pdf_path = temp_dir / f"capture_{uuid.uuid4().hex}.pdf"
    try:
        export_range_to_pdf(ws, rng, pdf_path, capture)
        copy_intermediate_file(pdf_path, capture, "01_excel_export", ".pdf")
        dpi, width, height = render_pdf_to_png(pdf_path, output.resolve(), capture)
    finally:
        try:
            pdf_path.unlink(missing_ok=True)
        except Exception:
            pass

    optimize_png(output, capture)
    final_intermediate = copy_intermediate_file(output, capture, "04_final", ".png")

    try:
        address = rng.Address(False, False)
    except TypeError:
        address = str(rng.Address)

    return {
        "path": str(output.resolve()),
        "sheet": ws.Name,
        "range": address,
        "width": width,
        "height": height,
        "engine": "pdf_render",
        "format": "png",
        "pdf_dpi": dpi,
        "optimize_png": bool(capture.get("optimize_png", False)),
        "png_colors": capture.get("png_colors", 256),
        "intermediate_dir": str(Path(capture["_intermediate_dir"]).resolve()) if capture.get("_keep_intermediate_files") else None,
        "final_intermediate": final_intermediate,
    }


def capture_range_to_png(ws, output_path, capture=None):
    capture = capture or {}
    engine = capture.get("engine", "pdf_render")
    if engine != "pdf_render":
        raise ValueError(f"不支持的截图引擎: {engine}；当前仅支持 pdf_render")
    return capture_range_to_png_pdf(ws, output_path, capture)


def optimize_png(image_path, capture):
    if not capture.get("optimize_png", False):
        return

    try:
        from PIL import Image
    except ImportError:
        return

    path = Path(image_path)
    png_colors = int(capture.get("png_colors", 256) or 256)
    png_colors = max(2, min(256, png_colors))
    with Image.open(path) as image:
        source = image.convert("RGB")
        optimized = source.quantize(colors=png_colors, method=Image.Quantize.MEDIANCUT)
        if "transparency" in optimized.info and isinstance(optimized.info["transparency"], tuple):
            del optimized.info["transparency"]
        optimized.save(path, format="PNG", optimize=True, compress_level=9)


def read_text_from_sheet(ws, text_config=None):
    text_config = text_config or {}
    mode = text_config.get("mode", "used_range")

    if mode == "explicit_range":
        address = text_config.get("range")
        if not address:
            raise ValueError("text.mode=explicit_range 时必须提供 text.range")
        rng = ws.Range(address)
    elif mode == "current_region":
        rng = ws.Range(text_config.get("start_cell", "A1")).CurrentRegion
    elif mode == "used_range":
        first_row, first_col, last_row, last_col = used_range_bounds(
            ws,
            shrink_empty_edges=text_config.get("shrink_empty_edges", True),
        )
        rng = ws.Range(ws.Cells(first_row, first_col), ws.Cells(last_row, last_col))
    else:
        raise ValueError(f"不支持的文字读取模式: {mode}")

    values = normalize_2d(rng.Value)
    lines = []
    for row in values:
        cells = [str(value).strip() for value in row if has_value(value)]
        if cells:
            lines.append("\t".join(cells))
    return "\n".join(lines).strip()


def image_payload_from_file(image_path):
    image_bytes = Path(image_path).read_bytes()
    return {
        "base64": base64.b64encode(image_bytes).decode("ascii"),
        "md5": hashlib.md5(image_bytes).hexdigest(),
    }


def save_package_payload_intermediate(capture_result, image_payload):
    intermediate_dir = capture_result.get("intermediate_dir")
    if not intermediate_dir:
        return None
    final_path = Path(capture_result["path"])
    stem = final_path.stem
    target = Path(intermediate_dir) / f"{stem}_05_package_base64_decoded.png"
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(base64.b64decode(image_payload["base64"]))
    return str(target.resolve())


def build_preview_image(package, preview_file):
    image_paths = [
        Path(item["capture"]["path"])
        for item in package.get("items", [])
        if item.get("type") == "image" and item.get("capture", {}).get("path")
    ]
    if not image_paths:
        return None

    try:
        from PIL import Image
    except ImportError:
        return None

    images = [Image.open(path).convert("RGB") for path in image_paths]
    try:
        padding = 24
        background = (255, 255, 255)
        width = max(image.width for image in images)
        height = sum(image.height for image in images) + padding * (len(images) - 1)
        preview = Image.new("RGB", (width, height), background)

        top = 0
        for image in images:
            left = (width - image.width) // 2
            preview.paste(image, (left, top))
            top += image.height + padding

        preview_file.parent.mkdir(parents=True, exist_ok=True)
        preview.save(preview_file)
        return str(preview_file.resolve())
    finally:
        for image in images:
            image.close()


def item_output_name(workbook_index, report_index, item_index, workbook_name, report_name, sheet_name):
    safe_parts = [
        f"{workbook_index:02d}",
        f"{report_index:02d}",
        f"{item_index:02d}",
        workbook_name,
        report_name,
        sheet_name,
    ]
    raw = "_".join(str(part) for part in safe_parts if part)
    for char in '<>:"/\\|?*':
        raw = raw.replace(char, "_")
    return raw[:160] + ".png"


def save_package(package_file, package):
    package_file.parent.mkdir(parents=True, exist_ok=True)
    with package_file.open("w", encoding="utf-8") as f:
        json.dump(package, f, ensure_ascii=False, indent=2)


def save_package_snapshot(config, package, base_dir=PROJECT_DIR):
    output_config = config.get("output", {})
    if not output_config.get("keep_intermediate_files", False):
        return None
    snapshot_path = get_intermediate_dir(config, base_dir) / "message_package_snapshot.json"
    snapshot_path.parent.mkdir(parents=True, exist_ok=True)
    with snapshot_path.open("w", encoding="utf-8") as f:
        json.dump(package, f, ensure_ascii=False, indent=2)
    return str(snapshot_path.resolve())


def get_excel_process_id(excel):
    try:
        import win32process

        _, pid = win32process.GetWindowThreadProcessId(excel.Hwnd)
        return pid
    except Exception:
        return None


def quit_excel(excel, pid=None):
    try:
        excel.Quit()
    except Exception:
        pass
    if not pid:
        return
    time.sleep(0.5)
    try:
        import win32api
        import win32con
        import win32process

        handle = win32api.OpenProcess(win32con.PROCESS_TERMINATE | win32con.PROCESS_QUERY_LIMITED_INFORMATION, False, pid)
        try:
            exit_code = win32process.GetExitCodeProcess(handle)
            if exit_code == 259:  # STILL_ACTIVE
                win32api.TerminateProcess(handle, 1)
        finally:
            handle.Close()
    except Exception:
        pass


def openpyxl_used_bounds(ws):
    min_row = ws.max_row or 1
    min_col = ws.max_column or 1
    max_row = 1
    max_col = 1
    found = False
    for row in ws.iter_rows():
        for cell in row:
            if has_value(cell.value):
                found = True
                min_row = min(min_row, cell.row)
                min_col = min(min_col, cell.column)
                max_row = max(max_row, cell.row)
                max_col = max(max_col, cell.column)
    if not found:
        raise ValueError(f"工作表 {ws.title} 没有非空单元格")
    return min_row, min_col, max_row, max_col


def openpyxl_range_bounds(ws, capture):
    capture = capture or {}
    mode = capture.get("mode", "used_range")
    if mode == "explicit_range":
        from openpyxl.utils.cell import range_boundaries

        address = capture.get("range")
        if not address:
            raise ValueError("capture.mode=explicit_range 时必须提供 capture.range")
        min_col, min_row, max_col, max_row = range_boundaries(address)
        return min_row, min_col, max_row, max_col
    if mode == "current_region":
        # The fallback cannot fully reproduce Excel's CurrentRegion. Use the used range
        # so the message can still be sent instead of failing the flow.
        return openpyxl_used_bounds(ws)
    if mode != "used_range":
        raise ValueError(f"不支持的截图模式: {mode}")
    return openpyxl_used_bounds(ws)


def read_text_from_openpyxl_sheet(ws, text_config=None):
    text_config = text_config or {}
    min_row, min_col, max_row, max_col = openpyxl_range_bounds(ws, text_config)
    lines = []
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True):
        values = [str(value).strip() for value in row if has_value(value)]
        if values:
            lines.append(" ".join(values))
    return "\n".join(lines).strip()


def build_text_message_package(config, base_dir=PROJECT_DIR):
    from openpyxl import load_workbook

    started = time.perf_counter()
    _, package_file, _ = get_output_paths(config, base_dir)
    package = {
        "generated_at": datetime.now().isoformat(),
        "items": [],
    }

    for workbook_index, workbook_config in enumerate(config.get("workbooks", []), start=1):
        workbook_name = workbook_config.get("name", f"workbook-{workbook_index}")
        workbook_path = resolve_path(workbook_config["file"], base_dir)
        workbook = load_workbook(workbook_path, data_only=True)
        try:
            for report_index, report in enumerate(workbook_config.get("reports", []), start=1):
                report_name = report.get("name", f"report-{report_index}")
                for item_index, item in enumerate(report.get("items", []), start=1):
                    item_type = item.get("type")
                    sheet_name = item.get("sheet")
                    if item_type != "text":
                        raise ValueError(f"纯文本消息包不支持 item.type: {item_type}")
                    if not sheet_name:
                        raise ValueError(f"{workbook_name}/{report_name} 的 item 缺少 sheet")
                    if sheet_name not in workbook.sheetnames:
                        raise KeyError(f"找不到工作表 {sheet_name!r}，当前工作表: {workbook.sheetnames}")

                    ws = workbook[sheet_name]
                    package_item = {
                        "workbook": workbook_name,
                        "workbook_file": str(workbook_path),
                        "report": report_name,
                        "item_index": item_index,
                        "type": item_type,
                        "sheet": sheet_name,
                    }
                    text = read_text_from_openpyxl_sheet(ws, item.get("text"))
                    if not text:
                        raise ValueError(f"{workbook_name}/{report_name}/{sheet_name} 未读取到文字")
                    package_item["text"] = text
                    package_item["text_length"] = len(text)
                    package["items"].append(package_item)
                    save_package(package_file, package)
        finally:
            workbook.close()

    package["timings"] = {
        "total_seconds": round(time.perf_counter() - started, 3),
        "engine": "openpyxl_text",
    }
    save_package(package_file, package)
    save_package_snapshot(config, package, base_dir)
    return package, str(package_file)


def switch_default_printer_if_needed(printer_api, target_printer):
    current_printer = str(printer_api.GetDefaultPrinter())
    if current_printer.casefold() == target_printer.casefold():
        return False
    printer_api.SetDefaultPrinter(target_printer)
    confirmed_printer = str(printer_api.GetDefaultPrinter())
    if confirmed_printer.casefold() != target_printer.casefold():
        raise RuntimeError(f"系统默认打印机仍为 {confirmed_printer!r}")
    return True


def open_excel_after_setting_default_printer(printer_api, target_printer, open_excel_fn, *, visible):
    switch_default_printer_if_needed(printer_api, target_printer)
    return open_excel_fn(visible=visible, cleanup_orphaned=True)


def build_message_package_com(config, base_dir=PROJECT_DIR, visible=False):
    started = time.perf_counter()
    image_dir, package_file, preview_file = get_output_paths(config, base_dir)
    default_capture = config.get("capture_defaults", {})
    workbooks = config.get("workbooks", [])
    if not workbooks:
        raise ValueError("配置中缺少 workbooks[]")

    package = {
        "generated_at": datetime.now().isoformat(),
        "items": [],
    }

    # A previous abnormal task may leave an invisible EXCEL.EXE holding stale
    # printer/COM state. The shared lock is acquired before cleanup, so another
    # healthy automation task cannot be killed here.
    active_printer = str(
        default_capture.get("excel_active_printer", "Microsoft Print to PDF") or ""
    ).strip()
    default_printer = active_printer or "Microsoft Print to PDF"
    try:
        import win32print  # type: ignore

        excel = open_excel_after_setting_default_printer(
            win32print,
            default_printer,
            open_excel,
            visible=visible,
        )
    except Exception as exc:
        if "excel" in locals():
            quit_excel(excel, get_excel_process_id(excel))
        raise RuntimeError(f"无法将系统默认打印机切换为 {default_printer!r} 后启动 Excel") from exc
    excel_pid = get_excel_process_id(excel)
    try:
        for workbook_index, workbook_config in enumerate(workbooks, start=1):
            workbook_name = workbook_config.get("name", f"workbook-{workbook_index}")
            workbook_path = resolve_path(workbook_config["file"], base_dir)
            open_config = workbook_config.get("excel_open", {})

            workbook = None
            try:
                workbook = open_workbook(
                    excel,
                    workbook_path,
                    update_links=open_config.get("update_links", False),
                    read_only=True,
                )

                if open_config.get("refresh_before_capture", False):
                    timeout = int(open_config.get("refresh_timeout_seconds", 120) or 120)
                    refresh_workbook(excel, workbook, timeout_seconds=timeout)

                for report_index, report in enumerate(workbook_config.get("reports", []), start=1):
                    report_name = report.get("name", f"report-{report_index}")
                    for item_index, item in enumerate(report.get("items", []), start=1):
                        item_type = item.get("type")
                        sheet_name = item.get("sheet")
                        if item_type not in {"image", "text"}:
                            raise ValueError(f"{workbook_name}/{report_name} 存在不支持的 item.type: {item_type}")
                        if not sheet_name:
                            raise ValueError(f"{workbook_name}/{report_name} 的 item 缺少 sheet")

                        ws = get_sheet(workbook, sheet_name)
                        package_item = {
                            "workbook": workbook_name,
                            "workbook_file": str(workbook_path),
                            "report": report_name,
                            "item_index": item_index,
                            "type": item_type,
                            "sheet": sheet_name,
                        }

                        if item_type == "image":
                            capture_config = merge_capture_config(default_capture, item.get("capture"))
                            image_name = item_output_name(
                                workbook_index,
                                report_index,
                                item_index,
                                workbook_name,
                                report_name,
                                sheet_name,
                            )
                            image_path = image_dir / image_name
                            capture_config = with_intermediate_capture(config, capture_config, image_name, base_dir)
                            capture_result = capture_range_to_png(ws, image_path, capture_config)
                            image_payload = image_payload_from_file(capture_result["path"])
                            package_payload_path = save_package_payload_intermediate(capture_result, image_payload)
                            if package_payload_path:
                                capture_result["package_payload_intermediate"] = package_payload_path
                            package_item["capture"] = capture_result
                            package_item["image"] = image_payload
                        else:
                            text = read_text_from_sheet(ws, item.get("text"))
                            if not text:
                                raise ValueError(f"{workbook_name}/{report_name}/{sheet_name} 未读取到文字")
                            package_item["text"] = text
                            package_item["text_length"] = len(text)

                        package["items"].append(package_item)
                        save_package(package_file, package)
            finally:
                if workbook is not None:
                    try:
                        workbook.Close(SaveChanges=False)
                    except Exception:
                        # Preserve the capture/opening exception. COM can
                        # return a dynamic placeholder when Open failed.
                        pass
    finally:
        quit_excel(excel, excel_pid)

    preview_path = build_preview_image(package, preview_file)
    if preview_path:
        package["preview_image_file"] = preview_path
    package["timings"] = {
        "total_seconds": round(time.perf_counter() - started, 3),
        "engine": "com",
    }
    save_package(package_file, package)
    save_package_snapshot(config, package, base_dir)
    return package, str(package_file)


def package_has_only_text_items(config):
    has_items = False
    for workbook_config in config.get("workbooks", []):
        for report in workbook_config.get("reports", []):
            for item in report.get("items", []):
                has_items = True
                if item.get("type") != "text":
                    return False
    return has_items


def build_message_package(config, base_dir=PROJECT_DIR, visible=False):
    if package_has_only_text_items(config):
        return build_text_message_package(config, base_dir=base_dir)
    # Image generation is strict: only Excel -> single-page PDF -> PNG is
    # accepted. Any rendering failure stops the task before an image is sent.
    return build_message_package_com(config, base_dir=base_dir, visible=visible)
