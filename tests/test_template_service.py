import json
import shutil
import tempfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from services.template_service import resolve_path, update_template


def write_json(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def make_work_dir():
    return Path(tempfile.mkdtemp(prefix="auto_notify_template_test_"))


def create_workbook(path, sheets):
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    for sheet_name, rows in sheets.items():
        sheet = workbook.create_sheet(sheet_name)
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                sheet.cell(row=row_index, column=column_index).value = value
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def test_resolve_path_rebases_logical_runtime_paths(monkeypatch, tmp_path):
    runtime_root = tmp_path / "shared-runtime"
    monkeypatch.setenv("AUTO_NOTIFY_RUNTIME_ROOT", str(runtime_root))

    assert resolve_path("runtime/flow/日报/debug/compare_result.json") == (
        runtime_root / "flow" / "日报" / "debug" / "compare_result.json"
    ).resolve()


def test_update_template_skips_when_compare_same():
    work_dir = make_work_dir()
    try:
        compare_result_path = work_dir / "compare_result.json"
        source_report_path = work_dir / "report.xlsx"
        template_path = work_dir / "template.xlsx"
        manifest_path = work_dir / "update_manifest.json"
        source_report_path.write_text("report", encoding="utf-8")
        template_path.write_text("template", encoding="utf-8")
        write_json(compare_result_path, {"result": "same", "sheets": []})

        manifest = update_template(
            {
                "compare_result_path": str(compare_result_path),
                "source_report_path": str(source_report_path),
                "template_path": str(template_path),
                "manifest_path": str(manifest_path),
            }
        )

        assert manifest["status"] == "skipped"
        assert manifest["reason"] == "compare_result_same"
        assert json.loads(manifest_path.read_text(encoding="utf-8"))["status"] == "skipped"
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_update_template_blocks_when_compare_invalid():
    work_dir = make_work_dir()
    try:
        compare_result_path = work_dir / "compare_result.json"
        source_report_path = work_dir / "report.xlsx"
        template_path = work_dir / "template.xlsx"
        source_report_path.write_text("report", encoding="utf-8")
        template_path.write_text("template", encoding="utf-8")
        write_json(compare_result_path, {"result": "invalid", "sheets": []})

        with pytest.raises(RuntimeError, match="比对结果为 invalid"):
            update_template(
                {
                    "compare_result_path": str(compare_result_path),
                    "source_report_path": str(source_report_path),
                    "template_path": str(template_path),
                    "manifest_path": str(work_dir / "update_manifest.json"),
                }
            )
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_update_template_allows_same_when_configured(monkeypatch):
    work_dir = make_work_dir()
    try:
        compare_result_path = work_dir / "compare_result.json"
        source_report_path = work_dir / "report.xlsx"
        template_path = work_dir / "template.xlsx"
        manifest_path = work_dir / "update_manifest.json"
        source_report_path.write_text("report", encoding="utf-8")
        template_path.write_text("template", encoding="utf-8")
        write_json(
            compare_result_path,
            {
                "result": "same",
                "sheets": [
                    {
                        "name": "通报",
                        "new_sheet_name": "通报",
                        "template_sheet_name": "通报",
                        "result": "same",
                        "source_report_path": str(source_report_path),
                    }
                ],
            },
        )

        def fake_copy_multi(source_report_path, template_path, output_path, sheet_results, write_sheets, visible=False):
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_text("updated", encoding="utf-8")
            return [{"name": "通报", "compare_result": "same"}]

        monkeypatch.setattr("services.template_service.update_template_copy_multi", fake_copy_multi)

        manifest = update_template(
            {
                "compare_result_path": str(compare_result_path),
                "source_report_path": str(source_report_path),
                "template_path": str(template_path),
                "manifest_path": str(manifest_path),
                "engine": "com_copy",
                "allow_same_update": True,
                "write_sheets": "all_compared",
            }
        )

        assert manifest["status"] == "updated"
        assert manifest["updated_sheets"][0]["compare_result"] == "same"
        assert Path(manifest["output_path"]).exists()
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_update_template_defaults_to_hybrid_and_only_writes_changed(monkeypatch):
    work_dir = make_work_dir()
    try:
        compare_result_path = work_dir / "compare_result.json"
        source_report_path = work_dir / "report.xlsx"
        template_path = work_dir / "template.xlsx"
        manifest_path = work_dir / "update_manifest.json"
        create_workbook(
            source_report_path,
            {
                "源变更": [["名称", "值"], ["A", 1]],
                "源相同": [["名称", "值"], ["B", 2]],
            },
        )
        create_workbook(
            template_path,
            {
                "模板变更": [["旧名称", "旧值"], ["old", 9], ["残留", 8]],
                "模板相同": [["不要动"]],
            },
        )
        workbook = load_workbook(template_path)
        workbook["模板变更"]["C1"] = "=SUM(B2:B10)"
        workbook.save(template_path)
        workbook.close()
        write_json(
            compare_result_path,
            {
                "result": "changed",
                "sheets": [
                    {
                        "name": "变更",
                        "new_sheet_name": "源变更",
                        "template_sheet_name": "模板变更",
                        "result": "changed",
                    },
                    {
                        "name": "相同",
                        "new_sheet_name": "源相同",
                        "template_sheet_name": "模板相同",
                        "result": "same",
                    },
                ],
            },
        )
        monkeypatch.setattr("services.template_service.refresh_workbook_with_excel", lambda *args, **kwargs: None)

        manifest = update_template(
            {
                "compare_result_path": str(compare_result_path),
                "source_report_path": str(source_report_path),
                "template_path": str(template_path),
                "manifest_path": str(manifest_path),
            }
        )

        assert manifest["engine"] == "hybrid"
        assert [item["name"] for item in manifest["updated_sheets"]] == ["变更"]
        output = load_workbook(manifest["output_path"], data_only=False)
        assert output["模板变更"]["A1"].value == "名称"
        assert output["模板变更"]["B2"].value == 1
        assert output["模板变更"]["A3"].value is None
        assert output["模板变更"]["B3"].value is None
        assert output["模板变更"]["C1"].value == "=SUM(B2:B10)"
        assert output["模板相同"]["A1"].value == "不要动"
        output.close()
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_update_template_hybrid_all_compared_writes_same(monkeypatch):
    work_dir = make_work_dir()
    try:
        compare_result_path = work_dir / "compare_result.json"
        source_report_path = work_dir / "report.xlsx"
        template_path = work_dir / "template.xlsx"
        manifest_path = work_dir / "update_manifest.json"
        create_workbook(source_report_path, {"源": [["新值"]]})
        create_workbook(template_path, {"模板": [["旧值"]]})
        write_json(
            compare_result_path,
            {
                "result": "same",
                "sheets": [
                    {
                        "name": "通报",
                        "new_sheet_name": "源",
                        "template_sheet_name": "模板",
                        "result": "same",
                    }
                ],
            },
        )
        monkeypatch.setattr("services.template_service.refresh_workbook_with_excel", lambda *args, **kwargs: None)

        manifest = update_template(
            {
                "compare_result_path": str(compare_result_path),
                "source_report_path": str(source_report_path),
                "template_path": str(template_path),
                "manifest_path": str(manifest_path),
                "allow_same_update": True,
                "write_sheets": "all_compared",
            }
        )

        assert manifest["engine"] == "hybrid"
        assert manifest["updated_sheets"][0]["compare_result"] == "same"
        output = load_workbook(manifest["output_path"])
        assert output["模板"]["A1"].value == "新值"
        output.close()
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_update_template_com_copy_uses_legacy_copy_engine(monkeypatch):
    work_dir = make_work_dir()
    try:
        compare_result_path = work_dir / "compare_result.json"
        source_report_path = work_dir / "report.xlsx"
        template_path = work_dir / "template.xlsx"
        manifest_path = work_dir / "update_manifest.json"
        source_report_path.write_text("report", encoding="utf-8")
        template_path.write_text("template", encoding="utf-8")
        write_json(
            compare_result_path,
            {
                "result": "changed",
                "sheets": [
                    {
                        "name": "通报",
                        "new_sheet_name": "源",
                        "template_sheet_name": "模板",
                        "result": "changed",
                    }
                ],
            },
        )
        called = {}

        def fake_copy_multi(source_report_path, template_path, output_path, sheet_results, write_sheets, visible=False):
            called["write_sheets"] = write_sheets
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_text("updated", encoding="utf-8")
            return [{"name": "通报", "compare_result": "changed"}]

        monkeypatch.setattr("services.template_service.update_template_copy_multi", fake_copy_multi)

        manifest = update_template(
            {
                "compare_result_path": str(compare_result_path),
                "source_report_path": str(source_report_path),
                "template_path": str(template_path),
                "manifest_path": str(manifest_path),
                "engine": "com_copy",
            }
        )

        assert called["write_sheets"] == "changed"
        assert manifest["engine"] == "com_copy"
        assert manifest["updated_sheets"][0]["compare_result"] == "changed"
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)
