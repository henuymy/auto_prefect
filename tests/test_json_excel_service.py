import json
import shutil
import time
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from services.json_excel_service import (
    drilldown_json_to_excel,
    get_by_path,
    json_response_to_excel,
)


def make_work_dir():
    work_dir = Path("runtime/test_work") / uuid4().hex
    work_dir.mkdir(parents=True, exist_ok=True)
    return work_dir


def test_get_by_path_reads_result_table_data():
    payload = {"result": {"tableData": [{"areaName": "中原区"}]}}

    assert get_by_path(payload, "result.tableData") == [{"areaName": "中原区"}]


def test_json_response_to_excel_writes_configured_columns():
    work_dir = make_work_dir()
    output_path = work_dir / "detail.xlsx"
    payload = {
        "reCode": "0000",
        "result": {
            "tableData": [
                {
                    "areaName": "樊晓玉&西流湖网格",
                    "areaCode": "13598027262&AQ726",
                    "sgs_ajvwdz": "0",
                }
            ]
        },
    }

    try:
        result = json_response_to_excel(
            payload,
            output_path,
            {
                "data_path": "result.tableData",
                "sheet_name": "地市作战明细",
                "columns": [
                    {"field": "areaName", "header": "名称"},
                    {"field": "areaCode", "header": "编码"},
                    {"field": "sgs_ajvwdz", "header": "爱家亲情网(V网版)"},
                ],
            },
        )

        workbook = load_workbook(output_path)
        worksheet = workbook["地市作战明细"]
        assert result["rows"] == 1
        assert [cell.value for cell in worksheet[1]] == ["名称", "编码", "爱家亲情网(V网版)"]
        assert [cell.value for cell in worksheet[2]] == ["樊晓玉&西流湖网格", "13598027262&AQ726", "0"]
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_drilldown_json_to_excel_uses_area_code_for_next_request():
    work_dir = make_work_dir()
    output_path = work_dir / "drilldown.xlsx"
    initial_payload = {
        "indCodes": "sgs_ajvwdz",
        "areaId": "AQ726",
        "diyCodes": "sgs_ajvwdz",
        "queryDate": "20260511",
    }
    responses = {
        "AQ726": {
            "reCode": "0000",
            "result": {
                "tableData": [
                    {"areaName": "西流湖网格", "areaCode": "13598027262&AQ726", "sgs_ajvwdz": "0"}
                ]
            },
        },
        "13598027262&AQ726": {
            "reCode": "0000",
            "result": {
                "tableData": [
                    {"areaName": "郑州市中原诺亚通讯商店", "areaCode": "AQ008E", "sgs_ajvwdz": "0"}
                ]
            },
        },
        "AQ008E": {
            "reCode": "0000",
            "result": {"tableData": []},
        },
    }
    requested = []

    def fetch_next(payload):
        requested.append(payload["areaId"])
        return json.loads(json.dumps(responses[payload["areaId"]], ensure_ascii=False))

    try:
        result = drilldown_json_to_excel(
            initial_payload,
            fetch_next,
            output_path,
            {
                "data_path": "result.tableData",
                "request_area_field": "areaId",
                "next_area_field": "areaCode",
                "levels": ["网格", "渠道/门店", "人员"],
            },
            {
                "sheet_name": "地市作战明细",
                "columns": [
                    {"field": "__level_name", "header": "层级"},
                    {"field": "__request_area_id", "header": "请求areaId"},
                    {"field": "areaName", "header": "名称"},
                    {"field": "areaCode", "header": "编码"},
                ],
            },
            initial_response_json=responses["AQ726"],
        )

        workbook = load_workbook(output_path)
        rows = list(workbook["地市作战明细"].iter_rows(values_only=True))
        assert requested == ["13598027262&AQ726", "AQ008E"]
        assert result["rows"] == 2
        assert rows[1] == ("网格", "AQ726", "西流湖网格", "13598027262&AQ726")
        assert rows[2] == ("渠道/门店", "13598027262&AQ726", "郑州市中原诺亚通讯商店", "AQ008E")
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_drilldown_json_to_excel_fetches_same_level_in_parallel():
    work_dir = make_work_dir()
    output_path = work_dir / "drilldown_parallel.xlsx"
    initial_payload = {"areaId": "ROOT"}
    responses = {
        "ROOT": {
            "reCode": "0000",
            "result": {
                "tableData": [
                    {"areaName": "A", "areaCode": "A"},
                    {"areaName": "B", "areaCode": "B"},
                ]
            },
        },
        "A": {"reCode": "0000", "result": {"tableData": [{"areaName": "A1", "areaCode": ""}]}},
        "B": {"reCode": "0000", "result": {"tableData": [{"areaName": "B1", "areaCode": ""}]}},
    }
    requested = []

    def fetch_next(payload):
        requested.append(payload["areaId"])
        time.sleep(0.2)
        return json.loads(json.dumps(responses[payload["areaId"]], ensure_ascii=False))

    try:
        started = time.perf_counter()
        result = drilldown_json_to_excel(
            initial_payload,
            fetch_next,
            output_path,
            {
                "data_path": "result.tableData",
                "request_area_field": "areaId",
                "next_area_field": "areaCode",
                "levels": ["区县", "网格"],
                "max_workers": 2,
            },
            {
                "sheet_name": "地市作战明细",
                "columns": [
                    {"field": "__level_name", "header": "层级"},
                    {"field": "areaName", "header": "名称"},
                ],
            },
            initial_response_json=responses["ROOT"],
        )
        elapsed = time.perf_counter() - started

        workbook = load_workbook(output_path)
        rows = list(workbook["地市作战明细"].iter_rows(values_only=True))
        assert set(requested) == {"A", "B"}
        assert elapsed < 0.35
        assert result["requests"] == 2
        assert result["max_workers"] == 2
        assert rows[1:] == [("区县", "A"), ("区县", "B"), ("网格", "A1"), ("网格", "B1")]
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)
