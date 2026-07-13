import shutil
import tempfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from backend.services import starter_template


def make_work_dir():
    return Path(tempfile.mkdtemp(prefix="auto_notify_starter_template_test_"))


def create_workbook(path, sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, rows in sheets.items():
        sheet = workbook.create_sheet(sheet_name)
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                sheet.cell(row=row_index, column=column_index).value = value
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def test_merge_downloads_to_template_copies_values_and_safely_names_sheets():
    work_dir = make_work_dir()
    try:
        source_one = work_dir / "one.xlsx"
        source_two = work_dir / "two.xlsx"
        create_workbook(
            source_one,
            {
                "特别长的原始sheet名称会导致目标sheet超长": [["名称", "值"], ["A", 1]],
                "重复": [["B", 2]],
            },
        )
        create_workbook(source_two, {"重复": [["C", 3]]})
        output_path = work_dir / "starter.xlsx"

        merged = starter_template._merge_downloads_to_template(
            {
                "results": [
                    {"name": "下载项名称也很长很长", "stage": "report_analysis", "output_path": str(source_one)},
                    {"name": "下载项名称也很长很长", "stage": "report_analysis", "output_path": str(source_two)},
                ]
            },
            output_path,
        )

        output = load_workbook(output_path)
        assert output.sheetnames[0] == "通报"
        assert len(output.sheetnames) == 4
        assert len(set(output.sheetnames)) == len(output.sheetnames)
        assert all(len(name) <= 31 for name in output.sheetnames)
        assert output[merged[0]["template_sheet_name"]]["A2"].value == "A"
        assert output[merged[2]["template_sheet_name"]]["B1"].value == 3
        output.close()
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_merge_downloads_to_template_requires_output_path():
    work_dir = make_work_dir()
    try:
        with pytest.raises(RuntimeError, match="缺少 output_path"):
            starter_template._merge_downloads_to_template({"results": [{"name": "下载"}]}, work_dir / "starter.xlsx")
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_generate_starter_template_uses_request_snapshot(monkeypatch):
    work_dir = make_work_dir()
    try:
        source = work_dir / "source.xlsx"
        create_workbook(source, {"明细": [["名称"], ["未保存配置数据"]]})
        monkeypatch.setattr(starter_template, "PROJECT_ROOT", work_dir)
        monkeypatch.setattr(starter_template, "RUNTIME_DIR", work_dir / "runtime")
        monkeypatch.setattr(starter_template, "TEMPLATES_DIR", work_dir / "templates")
        monkeypatch.setattr(
            starter_template,
            "_build_download_config",
            lambda downloads, _run_dir: {"reports": downloads},
        )
        monkeypatch.setattr(starter_template, "_prepare_required_session", lambda stages: {"status": "reused", "validation": {"required": stages}})

        def fake_download_reports(config, base_dir, dry_run=False, debug=False):
            assert config["reports"][0]["name"] == "页面未保存下载项"
            return {
                "results": [
                    {
                        "name": config["reports"][0]["name"],
                        "stage": config["reports"][0]["stage"],
                        "output_path": str(source),
                    }
                ]
            }

        monkeypatch.setattr(starter_template, "download_reports", fake_download_reports)

        result = starter_template.generate_starter_template(
            {
                "id": "draft",
                "name": "未保存配置",
                "downloads": [
                    {
                        "name": "页面未保存下载项",
                        "stage": "custom_stage",
                        "method": "POST",
                        "url": "https://example.invalid/export",
                        "body_type": "json",
                        "response_mode": "file",
                    }
                ],
            }
        )

        assert result["status"] == "success"
        assert result["required_stages"] == ["custom_stage"]
        output = load_workbook(starter_template.PROJECT_ROOT / result["template_path"])
        assert output.sheetnames[0] == "通报"
        assert output[result["merged_sheets"][0]["template_sheet_name"]]["A2"].value == "未保存配置数据"
        output.close()
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_generate_starter_template_returns_logical_shared_runtime_manifest_path(monkeypatch):
    work_dir = make_work_dir()
    try:
        project_dir = work_dir / "project"
        runtime_dir = work_dir / "shared-runtime"
        source = work_dir / "source.xlsx"
        create_workbook(source, {"明细": [["名称"], ["共享目录数据"]]})
        monkeypatch.setenv("AUTO_NOTIFY_RUNTIME_ROOT", str(runtime_dir))
        monkeypatch.setattr(starter_template, "PROJECT_ROOT", project_dir)
        monkeypatch.setattr(starter_template, "RUNTIME_DIR", runtime_dir / "starter_templates")
        monkeypatch.setattr(starter_template, "TEMPLATES_DIR", project_dir / "templates")
        monkeypatch.setattr(
            starter_template,
            "_build_download_config",
            lambda downloads, _run_dir: {"reports": downloads},
        )
        monkeypatch.setattr(
            starter_template,
            "_prepare_required_session",
            lambda stages: {"status": "reused", "validation": {"required": stages}},
        )
        monkeypatch.setattr(
            starter_template,
            "download_reports",
            lambda config, base_dir, dry_run=False, debug=False: {
                "results": [
                    {
                        "name": config["reports"][0]["name"],
                        "stage": config["reports"][0]["stage"],
                        "output_path": str(source),
                    }
                ]
            },
        )

        result = starter_template.generate_starter_template(
            {
                "name": "共享目录模板",
                "downloads": [
                    {
                        "name": "下载",
                        "stage": "city_ops",
                        "method": "POST",
                        "url": "https://example.invalid/export",
                        "body_type": "json",
                        "response_mode": "file",
                    }
                ],
            }
        )

        assert result["status"] == "success"
        assert result["download_manifest_path"].startswith(
            "runtime/starter_templates/共享目录模板/"
        )
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_generate_starter_template_relogs_once_when_session_expired(monkeypatch):
    work_dir = make_work_dir()
    try:
        source = work_dir / "source.xlsx"
        create_workbook(source, {"明细": [["名称"], ["刷新后数据"]]})
        monkeypatch.setattr(starter_template, "PROJECT_ROOT", work_dir)
        monkeypatch.setattr(starter_template, "RUNTIME_DIR", work_dir / "runtime")
        monkeypatch.setattr(starter_template, "TEMPLATES_DIR", work_dir / "templates")
        monkeypatch.setattr(
            starter_template,
            "_build_download_config",
            lambda downloads, _run_dir: {"reports": downloads},
        )
        session_calls = []

        def fake_prepare(stages, force_refresh=False):
            session_calls.append(force_refresh)
            return {"status": "refreshed" if force_refresh else "reused"}

        download_calls = {"count": 0}

        def fake_download_reports(config, base_dir, dry_run=False, debug=False):
            download_calls["count"] += 1
            if download_calls["count"] == 1:
                raise RuntimeError("JSON 接口返回 session 已过期: reCode=1101, reMsg=单点登录超时，请登录后重新跳转")
            return {
                "results": [
                    {
                        "name": config["reports"][0]["name"],
                        "stage": config["reports"][0]["stage"],
                        "output_path": str(source),
                    }
                ]
            }

        monkeypatch.setattr(starter_template, "_prepare_required_session_with_refresh", fake_prepare)
        monkeypatch.setattr(starter_template, "download_reports", fake_download_reports)

        result = starter_template.generate_starter_template(
            {
                "id": "session",
                "name": "会话刷新",
                "downloads": [
                    {
                        "name": "下载",
                        "stage": "smart_ops",
                        "method": "POST",
                        "url": "https://example.invalid/export",
                        "body_type": "json",
                        "response_mode": "file",
                    }
                ],
            }
        )

        assert result["status"] == "success"
        assert session_calls == [False, True]
        assert download_calls["count"] == 2
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)
