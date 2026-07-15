import json
import shutil
import tempfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from services.compare_service import (
    compare_report,
    compare_tables,
    find_empty_download_sheet_mappings,
    resolve_path,
    sheet_has_data_below_header,
)


def make_work_dir():
    return Path(tempfile.mkdtemp(prefix="auto_notify_compare_test_"))


def create_workbook(path, sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, rows in sheets.items():
        sheet = workbook.create_sheet(sheet_name)
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                sheet.cell(row=row_index, column=column_index).value = value
    workbook.save(path)
    workbook.close()


def test_resolve_path_rebases_root_relative_flow_paths(monkeypatch, tmp_path):
    runtime_root = tmp_path / "shared-runtime"
    monkeypatch.setenv("AUTO_NOTIFY_RUNTIME_ROOT", str(runtime_root))

    assert resolve_path("flow/日报/tmp/compare_results/result.json") == (
        runtime_root / "flow" / "日报" / "tmp" / "compare_results" / "result.json"
    ).resolve()


def test_resolve_path_keeps_project_relative_templates(tmp_path):
    assert resolve_path("templates/日报.xlsx", base_dir=tmp_path) == (
        tmp_path / "templates" / "日报.xlsx"
    ).resolve()


def test_resolve_path_rejects_legacy_runtime_prefix(monkeypatch, tmp_path):
    monkeypatch.setenv("AUTO_NOTIFY_RUNTIME_ROOT", str(tmp_path / "shared-runtime"))

    with pytest.raises(ValueError, match="不得以 runtime/"):
        resolve_path("runtime/flow/日报/tmp/result.json")


def test_compare_tables_same():
    table = [
        ["name", "count"],
        ["A", 1],
        ["B", "2"],
    ]

    result = compare_tables(table, table)

    assert result.result == "same"
    assert result.new_row_count == 2
    assert result.old_row_count == 2


def test_compare_tables_supports_header_row_zero():
    table = [
        ["A", 1],
        ["B", 2],
    ]

    result = compare_tables(table, table, header_row=0)

    assert result.result == "same"
    assert result.new_row_count == 2
    assert result.old_row_count == 2


def test_compare_tables_changed_without_key():
    old_table = [
        ["name", "count"],
        ["A", 1],
    ]
    new_table = [
        ["name", "count"],
        ["A", 2],
        ["B", 1],
    ]

    result = compare_tables(new_table, old_table)

    assert result.result == "changed"
    assert result.diff_summary["added_rows"] == 1
    assert result.diff_summary["changed_rows"] == 2


def test_compare_tables_invalid_header_mismatch():
    old_table = [
        ["name", "count"],
        ["A", 1],
    ]
    new_table = [
        ["name", "total"],
        ["A", 1],
    ]

    result = compare_tables(new_table, old_table)

    assert result.result == "invalid"
    assert result.message == "新报表和模板旧数据字段不一致"


def test_compare_report_uses_openpyxl_engine_by_default():
    work_dir = make_work_dir()
    try:
        new_report_path = work_dir / "new.xlsx"
        template_path = work_dir / "template.xlsx"
        output_path = work_dir / "compare_result.json"
        create_workbook(new_report_path, {"明细": [["name", "count"], ["A", 2]]})
        create_workbook(template_path, {"明细": [["name", "count"], ["A", 1]]})

        result = compare_report(
            {
                "new_report_path": str(new_report_path),
                "template_path": str(template_path),
                "output_path": str(output_path),
                "sheet_mappings": ["明细"],
            }
        )

        assert result["result"] == "changed"
        assert result["summary"]["changed"] == 1
        assert json.loads(output_path.read_text(encoding="utf-8"))["result"] == "changed"
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_compare_report_stream_hash_fast_path_for_same_sheet():
    work_dir = make_work_dir()
    try:
        new_report_path = work_dir / "new.xlsx"
        template_path = work_dir / "template.xlsx"
        rows = [["name", "count"], ["A", 1], ["B", 2]]
        create_workbook(new_report_path, {"明细": rows})
        create_workbook(template_path, {"明细": rows})

        result = compare_report(
            {
                "new_report_path": str(new_report_path),
                "template_path": str(template_path),
                "sheet_mappings": ["明细"],
            }
        )

        assert result["result"] == "same"
        assert result["sheets"][0]["diff_summary"]["fast_path"] == "stream_hash_equal"
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_compare_report_openpyxl_parallel_multiple_sheets():
    work_dir = make_work_dir()
    try:
        new_report_path = work_dir / "new.xlsx"
        template_path = work_dir / "template.xlsx"
        create_workbook(
            new_report_path,
            {
                "一": [["name", "count"], ["A", 1]],
                "二": [["name", "count"], ["B", 3]],
                "三": [["name", "count"], ["C", 5]],
            },
        )
        create_workbook(
            template_path,
            {
                "一": [["name", "count"], ["A", 1]],
                "二": [["name", "count"], ["B", 2]],
                "三": [["name", "count"], ["C", 5]],
            },
        )

        result = compare_report(
            {
                "new_report_path": str(new_report_path),
                "template_path": str(template_path),
                "sheet_mappings": ["一", "二", "三"],
                "max_workers": 2,
            }
        )

        assert result["result"] == "changed"
        assert result["summary"]["same"] == 2
        assert result["summary"]["changed"] == 1
        assert result["summary"]["max_workers"] == 2
        assert result["max_workers"] == 2
        assert [item["name"] for item in result["sheets"]] == ["一", "二", "三"]
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_sheet_has_data_below_header_uses_header_row():
    work_dir = make_work_dir()
    try:
        workbook_path = work_dir / "new.xlsx"
        create_workbook(workbook_path, {"明细": [["name", "count"], [None, None]]})

        assert sheet_has_data_below_header(workbook_path, "明细", header_row=1) is False
        assert sheet_has_data_below_header(workbook_path, "明细", header_row=0) is True
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_find_empty_download_sheet_mappings_reports_empty_data_area():
    work_dir = make_work_dir()
    try:
        workbook_path = work_dir / "new.xlsx"
        create_workbook(workbook_path, {"明细": [["name", "count"]]})
        manifest = {"results": [{"name": "下载A", "output_path": str(workbook_path)}]}
        compare_sources = [
            {
                "download_name": "下载A",
                "sheet_mappings": [
                    {"new_sheet_name": "明细", "template_sheet_name": "模板明细", "header_row": 1}
                ],
            }
        ]

        empty = find_empty_download_sheet_mappings(manifest, compare_sources)

        assert empty == [
            {
                "download_name": "下载A",
                "source_report_path": str(workbook_path.resolve()),
                "new_sheet_name": "明细",
                "template_sheet_name": "模板明细",
                "header_row": 1,
                "source_index": 1,
                "mapping_index": 1,
                "reason": "download_sheet_empty_below_header",
            }
        ]
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)
