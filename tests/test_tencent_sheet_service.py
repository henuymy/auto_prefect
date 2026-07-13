import json
import shutil
import tempfile
from pathlib import Path

import requests
from openpyxl import load_workbook

from services.tencent_sheet_service import (
    cell_to_value,
    download_tencent_sheet_report,
    extract_encoded_id,
    resolve_sheet_range,
    sheet_id_from_doc_url,
    split_a1_range,
)


def make_work_dir():
    return Path(tempfile.mkdtemp(prefix="auto_notify_tencent_sheet_test_"))


def make_response(payload):
    response = requests.Response()
    response.status_code = 200
    response._content = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    response.headers["Content-Type"] = "application/json"
    return response


def make_response_with_status(status_code, payload):
    response = make_response(payload)
    response.status_code = status_code
    return response


def write_json(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def test_extract_doc_url_parts():
    url = "https://docs.qq.com/sheet/DY1h4R1Rmd0FwWFhF?tab=000002"

    assert extract_encoded_id(url) == "DY1h4R1Rmd0FwWFhF"
    assert sheet_id_from_doc_url(url) == "000002"


def test_split_a1_range_respects_cell_limit():
    chunks = split_a1_range(
        "A1:AH5000",
        {
            "max_rows_per_request": 1000,
            "max_columns_per_request": 200,
            "max_cells_per_request": 10000,
        },
    )

    assert chunks[0] == "A1:AH294"
    assert chunks[-1] == "A4999:AH5000"
    assert len(chunks) == 18


def test_resolve_sheet_range_can_auto_detect_or_clamp_to_sheet_bounds():
    sheet_property = {"rowTotal": 200, "columnTotal": 26}

    assert resolve_sheet_range("auto", sheet_property) == ("A1:Z200", True)
    assert resolve_sheet_range("", sheet_property) == ("A1:Z200", True)
    assert resolve_sheet_range("A1:Z1000", sheet_property) == ("A1:Z200", True)
    assert resolve_sheet_range("A1:B20", sheet_property) == ("A1:B20", False)


def test_cell_to_value_handles_common_types():
    assert cell_to_value({"cellValue": {"text": "abc"}}) == "abc"
    assert cell_to_value({"cellValue": {"number": "12.0"}}) == 12
    assert cell_to_value({"cellValue": {"time": {"year": 2026, "month": 6, "day": 20, "hour": 9, "minute": 5, "second": 0}}}) == "2026-06-20 09:05:00"
    assert cell_to_value({"cellValue": {"time": {"year": 1899, "month": 12, "day": 30, "hour": 0, "minute": 0, "second": 0}}}) == ""
    assert cell_to_value({"cellValue": {"location": {"name": "科兴科学园"}}}) == "科兴科学园"
    assert cell_to_value({"cellValue": {"link": {"text": "腾讯文档", "url": "https://docs.qq.com"}}}) == "腾讯文档"
    assert cell_to_value({"cellValue": {"select": {"multiple": False, "value": ["a"], "options": [{"id": "a", "text": "选项A"}]}}}) == "选项A"
    assert cell_to_value({"cellValue": {"multiple": False, "value": ["a"], "options": [{"id": "a", "text": "选项A"}]}}) == "选项A"
    assert cell_to_value({"cellValue": {"text": '{"multiple":false,"value":["a"],"options":[{"id":"a","text":"选项A"}]}'}}) == "选项A"


def test_download_tencent_sheet_report_writes_workbook():
    work_dir = make_work_dir()
    try:
        config_path = work_dir / "tencent_docs.json"
        local_path = work_dir / "tencent_docs.local.json"
        write_json(
            config_path,
            {
                "credentials": {"client_id": "", "access_token": "", "open_id": ""},
                "request_timeout_seconds": 10,
                "chunk_request_interval_seconds": 0,
                "retry": {"attempts": 1, "backoff_seconds": 0, "max_backoff_seconds": 0},
                "read_limits": {"max_rows_per_request": 1000, "max_columns_per_request": 200, "max_cells_per_request": 10000},
            },
        )
        write_json(
            local_path,
            {
                "credentials": {"client_id": "cid", "access_token": "token", "open_id": "openid"},
            },
        )
        calls = []

        class FakeSession:
            def get(self, url, **kwargs):
                calls.append(url)
                if "util/converter" in url:
                    return make_response({"ret": 0, "data": {"fileID": "300000000$ABC"}})
                if url.endswith("/openapi/spreadsheet/v3/files/300000000$ABC"):
                    return make_response({"ret": 0, "data": {"properties": [{"sheetId": "000002", "title": "日报", "rowTotal": 2, "columnTotal": 2}]}})
                if "/000002/A1:B2" in url:
                    return make_response(
                        {
                            "ret": 0,
                            "data": {
                                "gridData": {
                                    "startRow": 0,
                                    "startColumn": 0,
                                    "rows": [
                                        {"values": [{"cellValue": {"text": "区域"}}, {"cellValue": {"text": "值"}}]},
                                        {"values": [{"cellValue": {"text": "中原"}}, {"cellValue": {"number": "3"}}]},
                                    ],
                                }
                            },
                        }
                    )
                raise AssertionError(f"unexpected url: {url}")

        result = download_tencent_sheet_report(
            {
                "name": "腾讯日报",
                "source": "tencent_sheet",
                "doc_url": "https://docs.qq.com/sheet/DY1h4R1Rmd0FwWFhF?tab=000002",
                "tencent_config_path": str(config_path),
                "sheets": [{"sheet_id": "000002", "range": "A1:B2", "output_sheet_name": "日报"}],
            },
            work_dir / "downloads",
            session=FakeSession(),
        )

        workbook = load_workbook(result["output_path"], data_only=True)
        try:
            sheet = workbook["日报"]
            assert sheet["A1"].value == "区域"
            assert sheet["B2"].value == 3
        finally:
            workbook.close()
        assert result["source"] == "tencent_sheet"
        assert len(calls) == 3
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_download_tencent_sheet_report_stops_when_later_chunk_exceeds_sheet():
    work_dir = make_work_dir()
    try:
        config_path = work_dir / "tencent_docs.json"
        write_json(
            config_path,
            {
                "credentials": {"client_id": "cid", "access_token": "token", "open_id": "openid"},
                "request_timeout_seconds": 10,
                "chunk_request_interval_seconds": 0,
                "retry": {"attempts": 1, "backoff_seconds": 0, "max_backoff_seconds": 0},
                "read_limits": {"max_rows_per_request": 1000, "max_columns_per_request": 200, "max_cells_per_request": 4},
            },
        )
        calls = []

        class FakeSession:
            def get(self, url, **kwargs):
                calls.append(url)
                if "util/converter" in url:
                    return make_response({"ret": 0, "data": {"fileID": "300000000$ABC"}})
                if url.endswith("/openapi/spreadsheet/v3/files/300000000$ABC"):
                    return make_response({"ret": 0, "data": {"properties": [{"sheetId": "000002", "title": "日报", "rowTotal": 2, "columnTotal": 2}]}})
                if "/000002/A1:B2" in url:
                    return make_response(
                        {
                            "ret": 0,
                            "data": {
                                "gridData": {
                                    "startRow": 0,
                                    "startColumn": 0,
                                    "rows": [
                                        {"values": [{"cellValue": {"text": "区域"}}, {"cellValue": {"text": "值"}}]},
                                        {"values": [{"cellValue": {"text": "中原"}}, {"cellValue": {"number": "3"}}]},
                                    ],
                                }
                            },
                        }
                    )
                raise AssertionError(f"unexpected url: {url}")

        result = download_tencent_sheet_report(
            {
                "name": "腾讯日报",
                "source": "tencent_sheet",
                "doc_url": "https://docs.qq.com/sheet/DY1h4R1Rmd0FwWFhF?tab=000002",
                "tencent_config_path": str(config_path),
                "sheets": [{"sheet_id": "000002", "range": "A1:B4", "output_sheet_name": "日报"}],
            },
            work_dir / "downloads",
            session=FakeSession(),
        )

        workbook = load_workbook(result["output_path"], data_only=True)
        try:
            sheet = workbook["日报"]
            assert sheet["A1"].value == "区域"
            assert sheet["B2"].value == 3
        finally:
            workbook.close()
        assert result["sheets"][0]["configured_range"] == "A1:B4"
        assert result["sheets"][0]["range"] == "A1:B2"
        assert result["sheets"][0]["range_auto_adjusted"] is True
        assert result["sheets"][0]["planned_chunks"] == 1
        assert result["sheets"][0]["fetched_chunks"] == 1
        assert len(calls) == 3
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_download_tencent_sheet_report_stops_later_invalid_without_bounds():
    work_dir = make_work_dir()
    try:
        config_path = work_dir / "tencent_docs.json"
        write_json(
            config_path,
            {
                "credentials": {"client_id": "cid", "access_token": "token", "open_id": "openid"},
                "request_timeout_seconds": 10,
                "chunk_request_interval_seconds": 0,
                "retry": {"attempts": 1, "backoff_seconds": 0, "max_backoff_seconds": 0},
                "read_limits": {"max_rows_per_request": 1000, "max_columns_per_request": 200, "max_cells_per_request": 4},
            },
        )
        calls = []

        class FakeSession:
            def get(self, url, **kwargs):
                calls.append(url)
                if "util/converter" in url:
                    return make_response({"ret": 0, "data": {"fileID": "300000000$ABC"}})
                if url.endswith("/openapi/spreadsheet/v3/files/300000000$ABC"):
                    return make_response({"ret": 0, "data": {"properties": [{"sheetId": "000002", "title": "日报"}]}})
                if "/000002/A1:B2" in url:
                    return make_response(
                        {
                            "ret": 0,
                            "data": {
                                "gridData": {
                                    "startRow": 0,
                                    "startColumn": 0,
                                    "rows": [
                                        {"values": [{"cellValue": {"text": "区域"}}, {"cellValue": {"text": "值"}}]},
                                        {"values": [{"cellValue": {"text": "中原"}}, {"cellValue": {"number": "3"}}]},
                                    ],
                                }
                            },
                        }
                    )
                if "/000002/A3:B4" in url:
                    return make_response({"code": 400001, "message": "invalid param error: 'range' invalid"})
                raise AssertionError(f"unexpected url: {url}")

        result = download_tencent_sheet_report(
            {
                "name": "腾讯日报",
                "source": "tencent_sheet",
                "doc_url": "https://docs.qq.com/sheet/DY1h4R1Rmd0FwWFhF?tab=000002",
                "tencent_config_path": str(config_path),
                "sheets": [{"sheet_id": "000002", "range": "A1:B4", "output_sheet_name": "日报"}],
            },
            work_dir / "downloads",
            session=FakeSession(),
        )

        assert result["sheets"][0]["planned_chunks"] == 2
        assert result["sheets"][0]["fetched_chunks"] == 1
        assert result["sheets"][0]["stopped_reason"] == "range_invalid_after_data"
        assert len(calls) == 4
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def test_download_tencent_sheet_report_retries_retryable_api_errors():
    work_dir = make_work_dir()
    try:
        config_path = work_dir / "tencent_docs.json"
        write_json(
            config_path,
            {
                "credentials": {"client_id": "cid", "access_token": "token", "open_id": "openid"},
                "request_timeout_seconds": 10,
                "chunk_request_interval_seconds": 0,
                "retry": {"attempts": 2, "backoff_seconds": 0, "max_backoff_seconds": 0},
                "read_limits": {"max_rows_per_request": 1000, "max_columns_per_request": 200, "max_cells_per_request": 10000},
            },
        )
        range_calls = 0

        class FakeSession:
            def get(self, url, **kwargs):
                nonlocal range_calls
                if "util/converter" in url:
                    return make_response({"ret": 0, "data": {"fileID": "300000000$ABC"}})
                if url.endswith("/openapi/spreadsheet/v3/files/300000000$ABC"):
                    return make_response({"ret": 0, "data": {"properties": [{"sheetId": "000002", "title": "日报", "rowTotal": 2, "columnTotal": 2}]}})
                if "/000002/A1:B2" in url:
                    range_calls += 1
                    if range_calls == 1:
                        return make_response_with_status(503, {"message": "busy"})
                    return make_response(
                        {
                            "ret": 0,
                            "data": {
                                "gridData": {
                                    "startRow": 0,
                                    "startColumn": 0,
                                    "rows": [
                                        {"values": [{"cellValue": {"text": "区域"}}, {"cellValue": {"text": "值"}}]},
                                        {"values": [{"cellValue": {"text": "中原"}}, {"cellValue": {"number": "3"}}]},
                                    ],
                                }
                            },
                        }
                    )
                raise AssertionError(f"unexpected url: {url}")

        result = download_tencent_sheet_report(
            {
                "name": "腾讯日报",
                "source": "tencent_sheet",
                "doc_url": "https://docs.qq.com/sheet/DY1h4R1Rmd0FwWFhF?tab=000002",
                "tencent_config_path": str(config_path),
                "sheets": [{"sheet_id": "000002", "range": "A1:B2", "output_sheet_name": "日报"}],
            },
            work_dir / "downloads",
            session=FakeSession(),
        )

        assert range_calls == 2
        assert result["sheets"][0]["fetched_chunks"] == 1
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)
