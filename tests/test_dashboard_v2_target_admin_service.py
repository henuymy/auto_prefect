from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from services.dashboard_v2_target_admin_service import (
    activate_target_plan,
    build_target_plan_export,
    build_target_template,
    clone_target_plan,
    create_target_plan,
    get_target_values,
    import_target_template,
    list_target_plans,
    save_target_values,
    set_target_plan_realtime,
)
from services.dashboard_v2_target_service import resolve_v2_target_plan
from models.dashboard_v2 import TargetPlan
from sqlalchemy import select
from sqlalchemy.orm import Session
from services.dashboard_v2_query_service import (
    get_current_wide_table,
    get_historical_with_changes,
)
from tests.test_dashboard_v2_query_service import _engine


def _xlsx(rows: list[tuple[object, ...]]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "目标值"
    worksheet.append([
        "scenario",
        "period_type",
        "effective_from",
        "node_type",
        "node_code",
        "indicator_code",
        "target_value",
    ])
    for row in rows:
        worksheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _source_layout_xlsx() -> bytes:
    workbook = Workbook()
    branch = workbook.active
    branch.title = "区公司级"
    branch.append(["编码", "名称", "渠道数量"])
    branch.append(["AQ", "中原区", 120])
    grid = workbook.create_sheet("网格级")
    grid.append(["编码", "名称", "渠道数量"])
    grid.append(["AQ701", "须水网格", 60])
    manager = workbook.create_sheet("渠道经理级")
    manager.append(["编码", "名称"])
    manager.append(["M001&AQ701", "张经理"])
    channel = workbook.create_sheet("渠道级")
    channel.append(["编码", "名称"])
    channel.append(["C001", "测试渠道"])
    indicator = workbook.create_sheet("指标参考表")
    indicator.append(["编码", "名称"])
    indicator.append(["channel_count", "渠道数量"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_create_edit_activate_target_plan_and_current_uses_target():
    engine = _engine()

    plan = create_target_plan(
        engine,
        plan_name="界面日目标",
        scenario="NORMAL",
        period_type="DAY",
        effective_from=date(2026, 6, 15),
        priority=20,
    )
    assert plan["status"] == "DRAFT"

    saved = save_target_values(
        engine,
        plan_id=plan["id"],
        values=[
            {"node_id": 2, "indicator_id": 1, "target_value": "111"},
            {"node_id": 3, "indicator_id": 1, "target_value": 55},
        ],
    )
    assert saved == {"plan_id": plan["id"], "saved": 2, "created": 2, "updated": 0}

    grid = get_target_values(engine, plan_id=plan["id"], node_type="BRANCH")
    assert grid["row_count"] == 1
    assert grid["rows"][0]["target_value"] == 111.0

    published = activate_target_plan(engine, plan["id"])
    assert published["status"] == "ACTIVE"
    assert published["id"] != plan["id"]

    with Session(engine) as session:
        assert session.get(TargetPlan, plan["id"]).status == "DRAFT"

    realtime = set_target_plan_realtime(engine, plan["id"])
    assert realtime["is_realtime"] is True

    current = get_current_wide_table(
        engine,
        node_type="BRANCH",
        indicator_codes=["channel_count"],
    )
    row = current["rows"][0]
    assert row["metrics"]["channel_count"] == 100
    assert row["targets"]["channel_count"] == 111

    save_target_values(engine, plan_id=plan["id"], values=[
        {"node_id": 2, "indicator_id": 1, "target_value": 222},
        {"node_id": 3, "indicator_id": 1, "target_value": 55},
    ])
    current_after_edit = get_current_wide_table(
        engine,
        node_type="BRANCH",
        indicator_codes=["channel_count"],
    )
    historical = get_historical_with_changes(
        engine,
        as_of=datetime(2026, 6, 30, 10, 10, 0),
        node_type="BRANCH",
        indicator_codes=["channel_count"],
    )
    assert current_after_edit["rows"][0]["targets"]["channel_count"] == 222
    assert historical["rows"][0]["targets"]["channel_count"] == 111


def test_template_download_contains_expected_sheets():
    workbook = load_workbook(BytesIO(build_target_template(_engine())), read_only=True)
    try:
        assert workbook.sheetnames == ["目标值", "区域参照", "指标参照"]
        assert workbook["目标值"]["A1"].value == "scenario"
        assert workbook["区域参照"]["A1"].value == "node_type"
        assert workbook["指标参照"]["A1"].value == "indicator_code"
    finally:
        workbook.close()


def test_import_template_replaces_draft_values():
    engine = _engine()
    plan = create_target_plan(
        engine,
        plan_name="导入日目标",
        scenario="NORMAL",
        period_type="DAY",
        effective_from=date(2026, 7, 1),
    )

    result = import_target_template(
        engine,
        plan_id=plan["id"],
        content=_xlsx([
            ("NORMAL", "DAY", "2026-07-01", "BRANCH", "AQ", "channel_count", 120),
            ("NORMAL", "DAY", "2026-07-01", "GRID", "AQ701", "channel_count", 60),
        ]),
    )

    assert result["imported"] == 2
    values = get_target_values(engine, plan_id=plan["id"], indicator_code="channel_count")
    assert {row["target_value"] for row in values["rows"] if row["target_value"]} == {120.0, 60.0}


def test_import_template_rejects_invalid_rows():
    engine = _engine()
    plan = create_target_plan(
        engine,
        plan_name="错误导入",
        scenario="NORMAL",
        period_type="DAY",
        effective_from=date(2026, 7, 1),
    )

    with pytest.raises(ValueError, match="period_type 只支持 DAY/MONTH"):
        import_target_template(
            engine,
            plan_id=plan["id"],
            content=_xlsx([
                ("NORMAL", "DAY_ACC", "2026-07-01", "BRANCH", "AQ", "channel_count", 120),
            ]),
        )


def test_import_source_layout_uses_selected_plan_metadata():
    engine = _engine()
    plan = create_target_plan(
        engine,
        plan_name="分层日目标",
        scenario="PK",
        period_type="DAY",
        effective_from=date(2026, 7, 1),
    )

    result = import_target_template(
        engine,
        plan_id=plan["id"],
        content=_source_layout_xlsx(),
    )

    assert result["imported"] == 2
    assert result["plan"]["scenario"] == "PK"
    values = get_target_values(engine, plan_id=plan["id"], indicator_code="channel_count")
    assert {
        (row["node_type"], row["node_code"], row["target_value"])
        for row in values["rows"]
        if row["target_value"] is not None
    } == {
        ("BRANCH", "AQ", 120.0),
        ("GRID", "AQ701", 60.0),
    }


def test_plan_export_round_trips_through_source_layout_import():
    engine = _engine()
    source = create_target_plan(
        engine,
        plan_name="待导出日目标",
        scenario="NORMAL",
        period_type="DAY",
        effective_from=date(2026, 7, 1),
    )
    save_target_values(engine, plan_id=source["id"], values=[
        {"node_id": 2, "indicator_id": 1, "target_value": 120},
        {"node_id": 3, "indicator_id": 1, "target_value": 60},
    ])

    content = build_target_plan_export(engine, plan_id=source["id"])
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["区公司级", "网格级", "渠道经理级", "渠道级", "指标参考表"]
        assert list(workbook["区公司级"].values)[:2] == [
            ("编码", "名称", "渠道数量"),
            ("AQ", "中原区", 120),
        ]
    finally:
        workbook.close()

    target = create_target_plan(
        engine,
        plan_name="回传日目标",
        scenario="NORMAL",
        period_type="DAY",
        effective_from=date(2026, 7, 1),
    )
    result = import_target_template(engine, plan_id=target["id"], content=content)
    assert result["imported"] == 2


def test_list_target_plans_returns_value_counts():
    engine = _engine()
    plans = list_target_plans(engine)
    assert plans["plans"]
    assert all("value_count" in plan for plan in plans["plans"])


def test_same_effective_date_retires_previous_and_keeps_newer_version():
    engine = _engine()
    first = create_target_plan(
        engine, plan_name="同月方案", scenario="NORMAL", period_type="DAY",
        effective_from=date(2026, 7, 1),
    )
    save_target_values(engine, plan_id=first["id"], values=[
        {"node_id": 2, "indicator_id": 1, "target_value": 100},
    ])
    first_published = activate_target_plan(engine, first["id"])

    second = clone_target_plan(engine, source_plan_id=first["id"])
    assert second["status"] == "DRAFT"
    assert second["effective_from"] == "2026-07-01"
    assert second["value_count"] == 1
    second_published = activate_target_plan(engine, second["id"])

    with Session(engine) as session:
        statuses = dict(session.execute(select(TargetPlan.id, TargetPlan.status)).all())
    assert statuses[first["id"]] == "DRAFT"
    assert statuses[second["id"]] == "DRAFT"
    assert statuses[first_published["id"]] == "RETIRED"
    assert statuses[second_published["id"]] == "ACTIVE"
    assert first_published["version_no"] == 1
    assert second_published["version_no"] == 2

    republished = activate_target_plan(engine, first["id"])
    assert republished["status"] == "ACTIVE"
    assert republished["version_no"] == 3


def test_active_plan_applies_from_its_effective_date():
    engine = _engine()
    plan = create_target_plan(
        engine, plan_name="月内口径", scenario="NORMAL", period_type="DAY",
        effective_from=date(2026, 7, 15),
    )
    save_target_values(engine, plan_id=plan["id"], values=[
        {"node_id": 2, "indicator_id": 1, "target_value": 100},
    ])
    published = activate_target_plan(engine, plan["id"])

    with Session(engine) as session:
        assert resolve_v2_target_plan(
            session, business_date=date(2026, 7, 1), period_type="DAY"
        ).id == 1
        assert resolve_v2_target_plan(
            session, business_date=date(2026, 7, 14), period_type="DAY"
        ).id == 1
        assert resolve_v2_target_plan(
            session, business_date=date(2026, 7, 15), period_type="DAY"
        ).id == published["id"]
        assert resolve_v2_target_plan(
            session, business_date=date(2026, 7, 31), period_type="DAY"
        ).id == published["id"]
        assert resolve_v2_target_plan(
            session, business_date=date(2026, 6, 30), period_type="DAY"
        ).id == 1
