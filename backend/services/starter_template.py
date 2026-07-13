from __future__ import annotations

import copy
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook, load_workbook

from services.method_service import download_reports
from services.runtime_paths import display_path, resolve_runtime_path
from services.session_manager import prepare_session
from utils.config_loader import load_json_with_local_override
from utils.date_placeholders import resolve_dynamic_structure


PROJECT_ROOT = Path(__file__).resolve().parents[2]
TEMPLATES_DIR = PROJECT_ROOT / "templates"
RUNTIME_DIR = resolve_runtime_path("runtime/starter_templates")
assert RUNTIME_DIR is not None
ALLOWED_EXCEL_SUFFIXES = {".xlsx", ".xlsm"}


def _safe_name(value: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|]+', "_", str(value or "")).strip()
    return cleaned or "未命名配置"


def _read_json(path: Path) -> dict[str, Any]:
    payload, _ = load_json_with_local_override(path)
    return payload


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _resolve_path(path_value: str | Path, base_dir: Path = PROJECT_ROOT) -> Path:
    path = Path(path_value)
    return path if path.is_absolute() else base_dir / path


def _enabled_downloads(config: dict[str, Any]) -> list[dict[str, Any]]:
    return [
        _normalize_download_auth(item)
        for item in config.get("downloads") or []
        if item.get("enabled", True) is not False
    ]


def _normalize_download_auth(item: dict[str, Any]) -> dict[str, Any]:
    next_item = copy.deepcopy(item)
    if next_item.get("source") == "tencent_sheet":
        return next_item
    preset = str(next_item.get("auth_preset") or "")
    if preset == "报表分析 Ssr-token":
        next_item["stage"] = "report_analysis"
        next_item["headers_from_cookies"] = {"Ssr-token": "ssr-token"}
        next_item.pop("headers_from_session_storage", None)
    elif preset == "智慧运营 User-Info":
        next_item["stage"] = "smart_ops"
        next_item["headers_from_session_storage"] = {"User-Info": "zhyyptInfo.accessToken"}
        next_item.pop("headers_from_cookies", None)
    elif preset == "地市平台 Uaptoken":
        next_item["stage"] = "city_ops"
        next_item["headers_from_session_storage"] = {"Uaptoken": "uapToken"}
        next_item.pop("headers_from_cookies", None)
    return next_item


def _required_stages(downloads: list[dict[str, Any]]) -> list[str]:
    stages = []
    seen = set()
    for item in downloads:
        if item.get("source") == "tencent_sheet":
            continue
        stage = str(item.get("stage") or "").strip()
        if stage and stage not in seen:
            stages.append(stage)
            seen.add(stage)
    return stages


def _validate_downloads(downloads: list[dict[str, Any]]) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    if not downloads:
        return [{"path": "/downloads", "message": "至少需要一个启用的抓取项"}]

    required_fields = ("name", "stage", "method", "url", "body_type", "response_mode")
    for index, item in enumerate(downloads):
        if item.get("source") == "tencent_sheet":
            if not item.get("name"):
                issues.append({"path": f"/downloads/{index}/name", "message": "抓取项缺少 name"})
            if not (item.get("doc_url") or item.get("file_id")):
                issues.append({"path": f"/downloads/{index}/doc_url", "message": "腾讯文档必须填写 doc_url 或 file_id"})
            sheets = item.get("sheets") or []
            if not sheets:
                issues.append({"path": f"/downloads/{index}/sheets", "message": "腾讯文档至少需要一个 Sheet 范围"})
            for sheet_index, sheet in enumerate(sheets):
                if not (sheet.get("sheet_id") or sheet.get("sheet_name")):
                    issues.append({"path": f"/downloads/{index}/sheets/{sheet_index}/sheet_id", "message": "Sheet 必须填写 sheet_id 或 Sheet 名称"})
            continue
        for field in required_fields:
            if not item.get(field):
                issues.append({"path": f"/downloads/{index}/{field}", "message": f"抓取项缺少 {field}"})
        response_mode = item.get("response_mode")
        if response_mode in {"json_to_excel", "json_drilldown_to_excel"}:
            columns = ((item.get("excel") or {}).get("columns") or [])
            if not columns:
                issues.append({"path": f"/downloads/{index}/excel/columns", "message": "JSON 转 Excel 必须配置 excel.columns"})
        if response_mode == "json_drilldown_to_excel":
            drilldown = item.get("drilldown") or {}
            for field in ("data_path", "request_area_field", "next_area_field"):
                if not drilldown.get(field):
                    issues.append({"path": f"/downloads/{index}/drilldown/{field}", "message": f"级联下钻缺少 {field}"})
    return issues


def _resolve_dynamic_value(value: Any, now: datetime | None = None) -> Any:
    return resolve_dynamic_structure(value, now=now)


def _build_download_config(downloads: list[dict[str, Any]], run_dir: Path) -> dict[str, Any]:
    base = _read_json(PROJECT_ROOT / "config" / "modules" / "report_downloader.json")
    base.pop("report_defaults", None)
    base["output_dir"] = str(run_dir / "downloads")
    base["manifest_path"] = str(run_dir / "download_manifest.json")
    base["reports"] = [_resolve_dynamic_value(item) for item in downloads]
    return base


def _prepare_required_session(stages: list[str]) -> dict[str, Any]:
    return _prepare_required_session_with_refresh(stages, force_refresh=False)


def _prepare_required_session_with_refresh(stages: list[str], force_refresh: bool) -> dict[str, Any]:
    autologin_path = PROJECT_ROOT / "config" / "modules" / "autologin.json"
    config = _read_json(autologin_path)
    config["required_stages"] = stages
    return prepare_session(config, base_dir=PROJECT_ROOT, force_refresh=force_refresh)


def _session_status_message(session_result: dict[str, Any]) -> str:
    status = session_result.get("status")
    if status in {"reused", "reused_after_wait"}:
        return "session 探活通过，复用已有自动登录会话"
    if status == "refreshed":
        return "session 探活失败或不可用，已关闭旧自动登录浏览器并重新登录，登录浏览器已保留"
    if status == "invalid":
        return f"session 不可用: {session_result.get('reason')}"
    return f"session 状态: {status or 'unknown'}"


def _is_session_expired_error(exc: RuntimeError) -> bool:
    # Keep this aligned with flows.notify_single_flow.run_download_with_session_retry.
    # Only a clear session-expired download error should force a browser re-login.
    return "session 已过期" in str(exc)


def _download_reports_with_session_retry(download_config: dict[str, Any], stages: list[str]) -> dict[str, Any]:
    try:
        return download_reports(download_config, base_dir=PROJECT_ROOT, dry_run=False, debug=False)
    except RuntimeError as exc:
        if not _is_session_expired_error(exc):
            raise
        _prepare_required_session_with_refresh(stages, force_refresh=True)
        return download_reports(download_config, base_dir=PROJECT_ROOT, dry_run=False, debug=False)


def _copy_values(source_sheet, target_sheet) -> None:
    for row_index, row in enumerate(source_sheet.iter_rows(), start=1):
        for column_index, cell in enumerate(row, start=1):
            target_sheet.cell(row=row_index, column=column_index).value = cell.value


def _safe_sheet_name(raw_name: str, used_names: set[str]) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", "_", raw_name).strip("'").strip() or "Sheet"
    cleaned = cleaned[:31]
    candidate = cleaned
    index = 2
    while candidate in used_names:
        suffix = f"_{index}"
        candidate = f"{cleaned[:31 - len(suffix)]}{suffix}"
        index += 1
    used_names.add(candidate)
    return candidate


def _merge_downloads_to_template(download_manifest: dict[str, Any], output_path: Path) -> list[dict[str, str]]:
    workbook = Workbook()
    report_sheet = workbook.active
    report_sheet.title = "通报"
    used_sheet_names = {"通报"}
    merged_sheets: list[dict[str, str]] = []

    for result_index, result in enumerate(download_manifest.get("results") or [], start=1):
        output_path_value = result.get("output_path")
        if not output_path_value:
            raise RuntimeError(f"下载结果缺少 output_path: {result.get('name') or result_index}")
        source_path = _resolve_path(output_path_value)
        if not source_path.exists():
            raise RuntimeError(f"下载结果文件不存在: {source_path}")
        if source_path.suffix.lower() not in ALLOWED_EXCEL_SUFFIXES:
            raise RuntimeError(f"下载结果不是可读取的 xlsx/xlsm: {source_path}")

        try:
            source_workbook = load_workbook(source_path, data_only=False, read_only=True)
        except Exception as exc:
            raise RuntimeError(f"下载结果 Excel 无法读取: {source_path}: {exc}") from exc
        try:
            stage = str(result.get("stage") or result.get("source") or "source")
            download_name = str(result.get("name") or f"下载{result_index}")
            for source_sheet in source_workbook.worksheets:
                sheet_name = _safe_sheet_name(
                    f"{download_name}_{source_sheet.title}",
                    used_sheet_names,
                )
                target_sheet = workbook.create_sheet(sheet_name)
                _copy_values(source_sheet, target_sheet)
                merged_sheets.append(
                    {
                        "stage": stage,
                        "download_name": download_name,
                        "source_sheet_name": source_sheet.title,
                        "template_sheet_name": sheet_name,
                        "source_path": str(source_path),
                    }
                )
        finally:
            source_workbook.close()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return merged_sheets


def _emit_progress(progress: Callable[[str, str], None] | None, title: str, message: str) -> None:
    if progress:
        progress(title, message)


def generate_starter_template(config: dict[str, Any], progress: Callable[[str, str], None] | None = None) -> dict[str, Any]:
    _emit_progress(progress, "校验配置", "正在检查当前页面配置快照")
    downloads = _enabled_downloads(config)
    issues = _validate_downloads(downloads)
    if issues:
        return {"status": "failed", "issues": issues}

    report_name = _safe_name(config.get("name") or config.get("id") or "未命名配置")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = RUNTIME_DIR / report_name / timestamp
    template_filename = f"{report_name}_新手模板_{timestamp}.xlsx"
    template_path = TEMPLATES_DIR / template_filename

    stages = _required_stages(downloads)
    _emit_progress(progress, "探活/登录", f"正在探活本次下载所需 stage: {', '.join(stages) or '无'}")
    if stages:
        session_result = _prepare_required_session(stages)
        _emit_progress(progress, "探活/登录", _session_status_message(session_result))
        if session_result.get("status") == "invalid":
            raise RuntimeError(f"会话不可用: {session_result.get('reason')}")

    download_config = _build_download_config(downloads, run_dir)
    _write_json(run_dir / "download_config.json", download_config)
    _emit_progress(progress, "下载数据", f"正在下载 {len(downloads)} 个抓取项")
    download_manifest = _download_reports_with_session_retry(download_config, stages)
    _emit_progress(progress, "合并模板", "正在把下载数据写入新手模板")
    merged_sheets = _merge_downloads_to_template(download_manifest, template_path)
    relative_path = template_path.relative_to(PROJECT_ROOT).as_posix()

    manifest = {
        "status": "success",
        "template_path": relative_path,
        "filename": template_filename,
        "required_stages": stages,
        "download_manifest_path": display_path(
            run_dir / "download_manifest.json",
            project_dir=PROJECT_ROOT,
        ),
        "merged_sheets": merged_sheets,
        "generated_at": datetime.now().isoformat(),
    }
    _write_json(run_dir / "starter_template_manifest.json", manifest)
    _emit_progress(progress, "生成完成", f"已生成 {relative_path}")
    return manifest
